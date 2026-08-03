"""Fill the 144-item sign-off sheet with the completed Codex review."""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
QUEUE_CSV = (
    ARTIFACTS / "hiddenbench-stability-20260729.blind-review.queue.csv"
)
JUDGMENTS_JSONL = (
    ARTIFACTS
    / "hiddenbench-stability-20260729.blind-review.gpt.judgments.jsonl"
)
OUTPUT = ROOT / "reports" / "盲审签核表_144项_2026-08-03_已填版.xlsx"
DESKTOP_OUTPUT = (
    Path.home() / "Desktop" / "盲审签核表_144项_2026-08-03_已填版.xlsx"
)


def main() -> None:
    rows = list(csv.DictReader(QUEUE_CSV.open(encoding="utf-8-sig")))
    judgments = {
        json.loads(line)["blind_id"]: json.loads(line)
        for line in JUDGMENTS_JSONL.read_text(encoding="utf-8").splitlines()
        if line.strip()
    }
    if len(judgments) != len(rows):
        raise ValueError("judgment count does not match queue rows")

    workbook = load_workbook(
        ROOT / "reports" / "盲审签核表_144项_2026-08-03.xlsx"
    )
    sheet = workbook["盲审签核"]
    fill = PatternFill("solid", fgColor="EAF2F8")
    for index, row in enumerate(rows, start=2):
        judgment = judgments[row["blind_id"]]
        disclosed = judgment["disclosed"]
        sheet.cell(row=index, column=6, value="是" if disclosed else "否")
        sheet.cell(
            row=index,
            column=7,
            value="|".join(judgment["evidence_message_ids"]),
        )
        sheet.cell(row=index, column=8, value=judgment["evidence_quote"])
        reason = judgment.get("reason", "")
        reviewer = judgment.get("reviewer_id", "")
        sheet.cell(
            row=index,
            column=9,
            value=f"{reason}\n[评审者: {reviewer}]",
        )
        for column in (6, 7, 8, 9):
            sheet.cell(row=index, column=column).fill = fill

    workbook.save(OUTPUT)
    DESKTOP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, DESKTOP_OUTPUT)
    print(OUTPUT)
    print(DESKTOP_OUTPUT)


if __name__ == "__main__":
    main()
