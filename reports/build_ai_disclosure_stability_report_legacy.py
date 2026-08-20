from __future__ import annotations

import importlib.util
from pathlib import Path

from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "build_ai_disclosure_stability_report.py"
REPO_OUTPUT = (
    HERE
    / "previous"
    / "AI披露率与四题重复稳定性实验报告_2026-07-30.docx"
)
DESKTOP_OUTPUT = (
    Path.home()
    / "Desktop"
    / "之前版本"
    / "AI披露率与四题重复稳定性实验报告_2026-07-30.docx"
)


def _load_builder():
    spec = importlib.util.spec_from_file_location("ai_report_builder", SOURCE)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load report builder: {SOURCE}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _remove_confirmation_section(document) -> None:
    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "4. 确认性盲审的执行结果与诚信边界"
    )
    body = document._body._element
    heading_element = heading._p
    previous = heading_element.getprevious()
    if (
        previous is not None
        and previous.tag == qn("w:p")
        and previous.xpath(".//w:br[@w:type='page']")
    ):
        body.remove(previous)

    cursor = heading_element
    while cursor is not None and cursor.tag != qn("w:sectPr"):
        next_cursor = cursor.getnext()
        body.remove(cursor)
        cursor = next_cursor


def main() -> None:
    builder = _load_builder()
    original_audit_method = builder._add_ai_audit_method
    original_reproducibility = builder._add_reproducibility

    def add_legacy_audit_method(document, analysis):
        original_audit_method(document, analysis)
        _remove_confirmation_section(document)
        document.add_paragraph(
            "确认性版本应优先人工盲审全部84项AI—规则分歧，并从其余236项一致"
            "判断中分层抽取“已披露”和“未披露”样本，报告人工一致率、AI精确率、"
            "召回率及修订后的披露率。该工作尚未完成，不能在本报告中写成已验证。"
        )

    def add_legacy_reproducibility(document, analysis):
        original_reproducibility(document, analysis)
        conclusion = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("最终结论")
            and "84项分歧全量复核" in paragraph.text
        )
        conclusion._element.getparent().remove(conclusion._element)
        builder._add_callout(
            document,
            "最终结论",
            "本报告已完成Agent级信息包来源披露审计、四题两种机制各10次重复、"
            "结果分类和探索性统计。仍待完成的确认性工作包括：原子事实级审计、"
            "公共可见率、独立人工盲审、提前停止对照、动态额度分配和更多随机任务。",
        )

    builder._add_ai_audit_method = add_legacy_audit_method
    builder._add_reproducibility = add_legacy_reproducibility
    builder.OUTPUT = REPO_OUTPUT
    builder.DESKTOP_OUTPUT = DESKTOP_OUTPUT
    builder.main()


if __name__ == "__main__":
    main()
