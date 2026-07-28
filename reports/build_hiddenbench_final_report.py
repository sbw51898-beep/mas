from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "HiddenBench_MAF复现实验补充报告_2026-07-28.md"
OUTPUT = ROOT / "reports" / "给彭老师的HiddenBench_MAF复现实验最终报告_2026-07-28.docx"

INK = RGBColor(31, 45, 61)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(92, 101, 112)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
WHITE = RGBColor(255, 255, 255)
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    western = "Consolas" if mono else "Calibri"
    east_asia = "Microsoft YaHei"
    run.font.name = western
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), western)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), western)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, value, fld_end])
    tail = paragraph.add_run(" 页")
    set_font(tail, size=9, color=MUTED)


def create_decimal_numbering(doc, start=1):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start_node = OxmlElement("w:start")
    start_node.set(qn("w:val"), str(start))
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    lvl.extend([start_node, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_numbered_item(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    add_inline_markdown(p, text)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("MAS 知识治理｜HiddenBench 复现实验")
    set_font(run, size=9, bold=True, color=MUTED)
    p.paragraph_format.space_after = Pt(0)

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_title_page(doc):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    r = kicker.add_run("阶段复现实验最终报告")
    set_font(r, size=11, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("HiddenBench 在 Microsoft Agent Framework 上的复现")
    set_font(r, size=25, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    r = subtitle.add_run("分布式信息、多智能体讨论与错误共识的可审计实验")
    set_font(r, size=13, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    widths = [2160, 7200]
    set_table_geometry(table, widths)
    metadata = [
        ("提交对象", "彭老师"),
        ("研究对象", "HiddenBench + MAST"),
        ("实现框架", "Microsoft Agent Framework"),
        ("报告日期", "2026 年 7 月 28 日"),
    ]
    for idx, (label, value) in enumerate(metadata):
        set_cell_shading(table.cell(idx, 0), LIGHT_FILL)
        p1 = table.cell(idx, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        set_font(r1, size=10.5, bold=True, color=DARK_BLUE)
        p2 = table.cell(idx, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(value)
        set_font(r2, size=10.5, color=INK)

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [TABLE_WIDTH_DXA])
    set_cell_shading(callout.cell(0, 0), CALLOUT_FILL)
    p = callout.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "阶段结论：十题单种子筛选中，完整信息条件平均正确率为 0.950，"
        "分散信息讨论前为 0.200，讨论后为 0.625；讨论总体有帮助，"
        "但同时出现 3 个全体一致却错误的案例。"
    )
    set_font(r, size=10.5, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("代码与实验档案：github.com/sbw51898-beep/mas")
    set_font(r, size=9.5, color=MUTED)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_inline_markdown(paragraph, text, size=11):
    # Preserve lightweight emphasis and inline-code semantics without carrying Markdown marks.
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size - 0.5, color=DARK_BLUE, mono=True)
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True, color=INK)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size, italic=True, color=INK)
        else:
            run = paragraph.add_run(token)
            set_font(run, size=size, color=INK)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_cell_shading(table.cell(0, 0), CALLOUT_FILL)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, size=9, color=INK, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def choose_widths(rows):
    cols = len(rows[0])
    if cols == 2:
        first_header = rows[0][0].lower()
        if "agent" in first_header:
            return [1800, 7560]
        return [2400, 6960]
    if cols == 6:
        return [600, 3380, 1120, 1120, 1120, 2020]
    return [TABLE_WIDTH_DXA // cols] * cols


def add_table(doc, rows):
    widths = choose_widths(rows)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            if i == 0:
                set_cell_shading(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            if j >= 2 and len(row) >= 5:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(p, value, size=9.2 if len(row) >= 5 else 10)
            for run in p.runs:
                if i == 0:
                    run.bold = True
                    run.font.color.rgb = DARK_BLUE
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_blockquote(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_cell_shading(table.cell(0, 0), "E8EEF5")
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_markdown(p, text, size=10.5)
    for run in p.runs:
        run.italic = True
        run.font.color.rgb = DARK_BLUE


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_body_from_markdown(doc, markdown):
    lines = markdown.splitlines()
    # The title is already represented on the cover.
    if lines and lines[0].startswith("# "):
        lines = lines[1:]

    i = 0
    code = False
    code_lines = []
    paragraph_parts = []

    def flush_paragraph():
        nonlocal paragraph_parts
        if paragraph_parts:
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            add_inline_markdown(p, " ".join(part.strip() for part in paragraph_parts))
            paragraph_parts = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if code:
                add_code_block(doc, code_lines)
                code_lines = []
                code = False
            else:
                code = True
            i += 1
            continue
        if code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markdown(p, stripped[3:], size=16)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = BLUE
            i += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, stripped[4:], size=13)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = BLUE
            i += 1
            continue
        if stripped.startswith("#### "):
            flush_paragraph()
            p = doc.add_paragraph(style="Heading 3")
            add_inline_markdown(p, stripped[5:], size=12)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = DARK_BLUE
            i += 1
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            quote_parts = [stripped[2:]]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#", "-", "|", "```")):
                quote_parts.append(lines[i].strip())
                i += 1
            add_blockquote(doc, " ".join(quote_parts))
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            numbered_items = []
            first_number = int(re.match(r"^(\d+)\.", stripped).group(1))
            while i < len(lines):
                candidate = lines[i].strip()
                if not re.match(r"^\d+\.\s+", candidate):
                    break
                numbered_items.append(re.sub(r"^\d+\.\s+", "", candidate))
                i += 1
            num_id = create_decimal_numbering(doc, start=first_number)
            for item in numbered_items:
                add_numbered_item(doc, item, num_id)
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, stripped[2:])
            i += 1
            continue
        paragraph_parts.append(stripped)
        i += 1

    flush_paragraph()


def audit(doc):
    section = doc.sections[0]
    assert round(section.page_width.inches, 3) == 8.5
    assert round(section.page_height.inches, 3) == 11
    assert all(
        round(value.inches, 3) == 1.0
        for value in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        )
    )
    assert round(section.header_distance.inches, 3) == 0.492
    assert round(section.footer_distance.inches, 3) == 0.492
    assert len(doc.tables) >= 6
    for table in doc.tables:
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        assert tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa"
        assert int(tbl_w.get(qn("w:w"))) == TABLE_WIDTH_DXA
        tbl_ind = table._tbl.tblPr.find(qn("w:tblInd"))
        assert tbl_ind is not None and int(tbl_ind.get(qn("w:w"))) == TABLE_INDENT_DXA


def main():
    if not SOURCE.exists():
        raise SystemExit(f"Source report not found: {SOURCE}")
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_body_from_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    audit(doc)
    doc.core_properties.title = "HiddenBench 在 Microsoft Agent Framework 上的阶段复现实验最终报告"
    doc.core_properties.subject = "多智能体分布式信息整合复现实验"
    doc.core_properties.author = "MAS 知识治理项目组"
    doc.core_properties.keywords = "HiddenBench, Microsoft Agent Framework, MAST, MAS, DeepSeek"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    sys.exit(main())
