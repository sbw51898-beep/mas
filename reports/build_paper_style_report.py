from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from pathlib import Path
from typing import Any

ROOT_DIR = Path(__file__).resolve().parents[1]
for import_root in (ROOT_DIR, ROOT_DIR / "src"):
    resolved = str(import_root)
    if resolved not in sys.path:
        sys.path.insert(0, resolved)

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reports.build_hardened_confirmatory_report import (
    BLACK,
    BLUE,
    DARK,
    GRAY,
    _code,
    _font,
    _geometry,
    _page_field,
    _table,
)


REPORT_NAME = "基于MAF的多智能体私有信息披露机制复现与对比研究_2026-08-05.docx"
EVIDENCE_PATH = Path("reports/data/hiddenbench-paper-evidence-20260805.json")
REFERENCE_PATH = Path("reports/data/hiddenbench-paper-references-20260805.json")
ARCHITECTURE_FIGURE = Path("reports/figures/maf-governance-architecture.png")
FLOW_FIGURE = Path("reports/figures/hiddenbench-experiment-flow.png")
PRESERVED_REPORT = Path("reports/给彭老师的完整实验报告_2026-08-02至2026-08-04.docx")
PRESERVED_SHA256 = "0d56f74c359e1a7359102de678c8c8e39bd568d381e9ff2ccf37189b83f72765"
TITLE = "基于 MAF 的多智能体私有信息披露机制复现与对比研究"

NAVY = RGBColor(23, 54, 93)
MUTED = RGBColor(91, 101, 115)
LIGHT_FILL = "F4F6F9"
GOLD_FILL = "FFF4D6"


def _load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _configure_paper(document: DocumentType) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        "Heading 1": (16, 18, 10, BLUE),
        "Heading 2": (13, 12, 6, BLUE),
        "Heading 3": (12, 8, 4, DARK),
    }
    for name, (size, before, after, color) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _font(header.add_run("HiddenBench × MAF | 多智能体信息披露治理实验"), 8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(footer.add_run("实验研究报告 | 第 "), 8.5, color=GRAY)
    _page_field(footer)
    _font(footer.add_run(" 页"), 8.5, color=GRAY)


def _heading(document: DocumentType, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        _font(
            run,
            16 if level == 1 else 13 if level == 2 else 12,
            bold=True,
            color=BLUE if level < 3 else DARK,
        )


def _body(
    document: DocumentType,
    text: str,
    *,
    bold_lead: str | None = None,
    center: bool = False,
    italic: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.333
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        _font(first, 11, bold=True, color=DARK)
        rest = paragraph.add_run(text[len(bold_lead) :])
        _font(rest, 11, italic=italic)
    else:
        _font(paragraph.add_run(text), 11, italic=italic)


def _formula(document: DocumentType, formula: str, explanation: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(formula)
    _font(run, 12, bold=True, color=NAVY)
    _body(document, explanation)


def _caption(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph(style="Caption")
    _font(paragraph.add_run(text), 9, color=MUTED)


def _table_caption(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    _font(paragraph.add_run(text), 9.5, bold=True, color=DARK)


def _callout(document: DocumentType, lead: str, text: str, *, gold: bool = False) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), GOLD_FILL if gold else LIGHT_FILL)
    cell._tc.get_or_add_tcPr().append(shade)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.20
    _font(paragraph.add_run(lead), 10.5, bold=True, color=DARK)
    _font(paragraph.add_run(text), 10.5)
    _geometry(table, [9360])
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _figure(document: DocumentType, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(path), width=Inches(6.25))
    _caption(document, caption)


def _page_break(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _title_and_abstract(document: DocumentType, evidence: dict[str, Any]) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(8)
    _font(title.add_run(TITLE), 22, bold=True, color=NAVY)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(18)
    _font(meta.add_run("实验研究报告 | 2026 年 8 月 5 日"), 10.5, color=MUTED)

    _heading(document, "摘要", level=1)
    abstract = (
        "大语言模型多智能体系统通过分工和讨论整合分散知识，但信息不对称也可能使多个智能体在关键证据缺失时形成稳定的错误共识。"
        "HiddenBench 等研究已经观察到隐藏信息任务中的集体推理失败，现有多智能体辩论和编排框架则较少区分信息披露、发言顺序与最终正确性各自的作用。"
        "针对这一问题，本文基于 Microsoft Agent Framework 和 DeepSeek，复现三道官方短任务，并比较固定轮转、内容感知动态发言、结构化讨论、逐人机械披露和官方兼容全局披露等条件。"
        f"实验累计完成 {evidence['totals']['local_runs']} 次本地运行、{evidence['totals']['discussion_vote_requests']:,} 次讨论与投票请求以及 {evidence['totals']['audit_requests']} 次披露审计请求，共 {evidence['totals']['model_requests']:,} 次模型请求。"
        "在官方兼容全局 Reveal-All 条件下，ID1、ID2、ID3 均为 10/10，而官方公开的 GPT-4.1 Reveal-All 结果分别为 10/10、9/10 和 10/10。"
        "固定与动态同预算、固定与动态逐人披露、固定短预算与结构化协议的精确 McNemar 检验结果分别为 0.125、0.125 和 0.388，均未达到 0.05。"
        "结果表明，在当前三题和单模型范围内，确保关键私有信息完整进入讨论比改变发言调度更直接，但完整披露仍不等于正确理解，相关结论需要更多任务、人工盲审和跨模型实验确认。"
    )
    _body(document, abstract)
    keywords = document.add_paragraph()
    keywords.paragraph_format.space_after = Pt(12)
    _font(keywords.add_run("关键词："), 10.5, bold=True, color=DARK)
    _font(keywords.add_run("多智能体系统；Microsoft Agent Framework；HiddenBench；信息披露；错误共识；动态发言"), 10.5)
    _callout(
        document,
        "主要结论：",
        "当前证据支持“信息是否完整进入公共讨论”比“由谁发言”更直接影响结果；动态发言机制尚未显示统计显著优势，多智能体共识也不能作为正确性的替代指标。",
        gold=True,
    )


def _introduction(document: DocumentType) -> None:
    _heading(document, "1 引言")
    _body(
        document,
        "大语言模型（Large Language Model，LLM）驱动的多智能体系统把一个任务交给多个具有不同角色或信息的 Agent，通过消息交换、工具调用和投票得到最终答案。相关综述将角色分工、通信结构、记忆和工作流视为系统性能的主要来源[4]；多智能体辩论则进一步提出，让多个模型相互审查能够改善事实性和推理表现[5]。这些结果容易引出一个直觉：只要参与者更多、讨论更充分或共识更强，答案就会更可靠。",
    )
    _body(
        document,
        "这一推断在信息不对称任务中并不稳固。Hidden Profile 范式把决定性信息分散给不同成员，只有在讨论中汇总未共享信息，群体才可能找到正确答案。早期群体决策实验发现，讨论往往重复公共信息而忽视未共享信息[2]；25 年研究的元分析也表明，这类任务具有稳定的信息采样偏差[3]。HiddenBench 将该范式移植到 LLM 多智能体，系统性展示了分布式信息条件下的集体推理失败[1]。",
    )
    _body(
        document,
        "因此，本研究不把“达成一致”直接等同于“完成推理”。研究重点是三个连续环节：私有事实是否被披露，披露后的事实是否被其他 Agent 读取，以及这些事实是否被正确解释并转化为最终行动。MAST 对多智能体失败的分类说明，智能体间错位、讨论偏移和错误验证可能发生在不同执行阶段[11]，这为逐轮分析提供了可操作的诊断框架。",
    )
    _body(document, "RQ1：完整披露私有信息是否提高 HiddenBench 任务正确率？", bold_lead="RQ1：")
    _body(document, "RQ2：相同发言预算下，内容感知动态发言机制是否优于固定轮转？", bold_lead="RQ2：")
    _body(document, "RQ3：多智能体错误共识主要出现在哪些信息处理环节？", bold_lead="RQ3：")
    _body(
        document,
        "本文的工作包括三部分。第一，使用 Microsoft Agent Framework（MAF）作为 Agent 与模型调用执行层，在不依赖 HiddenBench 作者模拟器的条件下复现三道官方短任务。第二，在同模型、同任务和同发言预算下比较固定轮转、动态发言和结构化协议，并用原子事实披露率、正确率与配对检验约束结论。第三，保存逐轮消息和证据定位，将错误区分为信息未披露、信息被忽略和证据被错误解释，而不把负结果隐藏在平均数中。",
    )


def _related_work(document: DocumentType) -> None:
    _heading(document, "2 相关研究")
    _body(
        document,
        "本章按照“多智能体协作—隐藏信息任务—失败诊断—发言治理—执行框架”的顺序定义后文使用的概念。这样的组织方式避免先使用“披露率”“错误共识”或“动态调度”等术语，再在后文补充解释。",
    )

    _heading(document, "2.1 LLM 多智能体系统与多智能体辩论", level=2)
    _body(
        document,
        "LLM 多智能体系统通常通过角色专业化、任务分解和消息交互来扩展单模型能力[4]。多智能体辩论让多个 Agent 独立提出答案，再相互阅读理由并修订判断；Du 等人的实验显示，这类流程在若干事实性和推理任务上能够优于单次模型回答[5]。但该范式常以多数票或最后一轮共识作为输出，因此协作收益与从众风险可能同时存在。",
    )
    _body(
        document,
        "不同系统通过不同方式固化协作流程。MetaGPT 将标准作业流程编码为角色化消息链[6]，ChatDev 用聊天链组织软件设计、编码和测试[7]，AgentVerse 则提供任务求解与行为模拟两类多智能体环境[8]。这些工作证明多 Agent 可以形成复杂工作流，但并不自动保证分散证据会被充分公开和正确使用。",
    )

    _heading(document, "2.2 私有信息不对称与 HiddenBench", level=2)
    _body(
        document,
        "Hidden Profile 指群体成员分别掌握不同的未共享信息，而公共信息通常支持一个看似合理但错误的选项[2]。只有汇总关键私有事实，正确选项才会显现。元分析表明，讨论时间、群体规模和信息呈现方式会影响未共享信息进入讨论的概率，但增加交流并不必然消除偏差[3]。社会影响还可能压低意见多样性，却不改善群体误差[17]。",
    )
    _body(
        document,
        "HiddenBench 将这一范式构造成 65 道 LLM 多智能体任务，其中短任务向四个 Agent 分配四条私有事实，并要求讨论后选择一个共同答案[1]。本研究只使用官方短题 ID1、ID2 和 ID3，因为这三题公开了逐题 GPT-4.1 基线与 Reveal-All 结果，可进行直接对标；本文不把三题结果推广到全部 65 题。",
    )

    _heading(document, "2.3 多智能体失败模式与共识风险", level=2)
    _body(
        document,
        "MAST 从多种框架和任务中归纳出 14 种失败模式，并将问题分为系统设计、智能体间错位以及验证与终止三类[11]。本文重点观察 FC2 中与消息利用相关的 FM-2.4、FM-2.5 和 FM-2.6，但只在消息证据充分时确认模式；不能由关键词出现直接推断因果。",
    )
    _body(
        document,
        "反思也可能在错误立场形成后失效。Liang 等人把这种现象称为思维退化，即模型一旦建立高置信度立场，后续自我反思难以产生新的推理路径[12]。Free-MAD 进一步质疑以最终共识和多数票作为唯一决策依据，提出利用完整辩论轨迹而不是仅看最后一轮[13]。失败归因研究则表明，即使给定完整轨迹，自动定位责任 Agent 仍然困难[14]。",
    )

    _heading(document, "2.4 发言选择与信息披露治理", level=2)
    _body(
        document,
        "发言机制决定谁在何时有机会把私有信息带入公共上下文。AutoGen 支持通过集中式选择器组织群聊，每轮选择会增加一次模型调用[9]；DyLAN 则先优化 Agent 团队，再按任务形成动态协作网络[15]。这些方法说明发言结构可以被设计，但动态选择的收益必须与额外模型调用、全局信息权限和预算差异分开计算。",
    )
    _body(
        document,
        "Riedl 使用偏信息分解区分真实协同与仅由时间同步产生的表面耦合[16]。这一思路提示本文不能仅用消息数量或共识强度评价治理效果，而应同时记录私有事实是否公开、最终判断是否正确以及调度器获得了哪些额外信息。",
    )

    _heading(document, "2.5 多智能体执行框架", level=2)
    _body(
        document,
        "AutoGen、MetaGPT、ChatDev 和 AgentVerse 分别提供可对话 Agent、角色化流程、聊天链和多 Agent 环境[6-9]。Microsoft Agent Framework 是面向 Python 与 .NET 的开源 Agent 与工作流框架，支持多模型后端和多种编排方式[10]。本研究使用 MAF 建立 Agent 并调用 DeepSeek，但没有直接采用其 Group Chat 编排器。",
    )
    _body(
        document,
        "具体而言，固定轮转、五因子动态选择、Reveal-All、投票、停止条件、披露审计和实验门禁均由本项目代码实现。因而本文评估的是“在 MAF 执行层上实现的治理协议”，而不是对 MAF 原生群聊能力的性能测评。",
    )


def _method(document: DocumentType, evidence: dict[str, Any], root: Path) -> None:
    _heading(document, "3 基于 MAF 的信息披露治理方法")
    _body(
        document,
        "本研究把系统分成任务层、执行层、治理层和评估层。任务层读取 HiddenBench 公共背景、选项和四条私有事实；执行层由 MAF 创建四个 Agent 并调用 DeepSeek；治理层决定发言顺序和信息注入方式；评估层完成多数票、披露审计、统计检验和案例编码。",
    )
    _figure(document, root / ARCHITECTURE_FIGURE, "图 1  基于 MAF 的信息披露治理系统结构")

    _heading(document, "3.1 任务与 Agent 执行", level=2)
    _body(
        document,
        "每个任务包含相同的撤离背景和三个候选地点，但公共信息与四条私有事实不同。四个 Agent 获得相同公共背景，每个 Agent 额外获得一条私有事实。事实分配由 pair seed 控制，使配对条件中的 Agent 归属保持一致；模型生成仍具有服务端随机性。",
    )
    _body(
        document,
        "MAF 的职责是保存 Agent 指令、传入公开消息并调用模型。为了保持结论边界，代码与报告均明确记录：Microsoft Agent Framework is the Agent/model execution layer。讨论协议、选择器和审计器不是 MAF 自动产生的结果。",
    )

    _heading(document, "3.2 讨论与披露协议", level=2)
    _table_caption(document, "表 1  本研究比较的讨论条件")
    rows = [
        ["single-local", "1", "无公共讨论", "单个 Agent 只见自己的私有事实", "单智能体下界"],
        ["fixed-12", "4", "固定轮转，共 12 次发言", "自然决定是否披露", "短预算基线"],
        ["structured-12", "4", "Exchange→Decide，共 12 次发言", "先交换事实再决策", "结构化短预算"],
        ["fixed-60", "4", "固定轮转，共 60 次发言", "自然决定是否披露", "长预算基线"],
        ["dynamic-60", "4", "内容感知选择，共 60 次发言", "自然决定是否披露", "同预算动态条件"],
        ["fixed/dynamic-reveal-all", "4", "固定或动态，共 60 次发言", "逐位发言时机械追加该 Agent 事实", "旧机械披露"],
        ["official-global-reveal", "4", "固定官方流程", "第一轮每条响应注入全部四条事实", "官方兼容补充"],
    ]
    _table(document, ["条件", "Agent", "发言机制", "披露机制", "作用"], rows, [1700, 700, 2100, 3060, 1800], font_size=7.4)
    _body(
        document,
        "旧的 fixed/dynamic-reveal-all 是 owner-by-owner mechanical reveal：事实随其所有者发言逐条进入对话。新增 official-global-reveal 则是 official-compatible global Reveal-All：第一轮四条响应均附带全部四条私有事实。两者都能达到表面上的 100% 披露，但信息出现时机和重复强度不同，不能合并成一个处理条件。",
    )

    _heading(document, "3.3 原子事实披露率", level=2)
    _body(
        document,
        "披露审计先把每个私有信息包拆成原子事实，再让审计模型逐条判断该事实是否由其所有者在公共消息中忠实表达。忠实改写可以计为已披露；否定、极性反转、关键限定缺失或仅由非所有者转述不计。对 disclosed=true 的判断，审计器必须返回消息 ID 和消息中的精确子串，规则程序随后验证所有权、引文存在性和否定极性。",
    )
    _formula(
        document,
        "D = m / n × 100%",
        "其中，n 表示任务中私有原子事实总数，m 表示经 AI 判断并通过规则证据校验的已披露事实数。审计模型不直接计算百分比，程序在验证逐条判断后统一计算 D。",
    )
    example = evidence["disclosure_example"]
    _callout(
        document,
        "可核验实例：",
        f"运行 {example['run_id']} 中，四条原子事实有三条找到所有者消息证据，一条没有证据，因此 {example['arithmetic']}。这不是人工数格子，而是 AI 逐条判断、规则校验后由程序计算。",
    )

    _heading(document, "3.4 内容感知动态发言", level=2)
    _body(
        document,
        "动态条件不按 A-B-C-D 固定轮转，而是根据未披露事实、当前讨论缺口、立场差异、发言新颖性和最近发言惩罚等信息选择下一位 Agent。该机制是 external centralized scheduler，并且能够读取完整的私有事实分配，因此属于特权调度器。它适合验证“理想的信息感知调度是否可能改善结果”，但不能直接视为无需全局信息的可部署方案。",
    )

    _heading(document, "3.5 投票、正确率与统计检验", level=2)
    _body(
        document,
        "讨论结束后，四个 Agent 分别提交最终答案。本文以多数票作为任务答案，同时记录全体一致、错误共识和 2-2 平票。正确率均报告成功次数与总次数；10/10 不能只写成 100%，因为 Wilson 95% 区间仍为约 72.2% 至 100%。",
    )
    _body(
        document,
        "同一任务和同一事实分配种子下的两个条件构成一对。配对二分类结果使用双侧精确 McNemar 检验，零假设是两个条件的成功概率相同。显著性阈值预设为 0.05；比例更高但 p 值大于 0.05 时，只能表述为点估计改善，不能写成显著优势。",
    )

    _heading(document, "3.6 实验控制与复现边界", level=2)
    _body(
        document,
        "所有主要对照保持任务、模型、系统 Prompt、Agent 数量、重复次数和发言预算一致。先冻结固定轮转基线，再运行相同 60 次发言预算的动态条件，避免把更多模型调用误认为治理收益。实验流程见图 2。",
    )
    _figure(document, root / FLOW_FIGURE, "图 2  HiddenBench 同预算对照与证据审计流程")
    _callout(
        document,
        "随机性边界：",
        "Pair seeds control fact assignment only; generation seed was not transmitted to DeepSeek。换言之，配对实验控制事实归属，但不能保证每次模型输出逐字复现。",
        gold=True,
    )


def _success(evidence: dict[str, Any], condition: str, task_id: int) -> tuple[int, int]:
    item = evidence["condition_statistics"][f"{condition}:ID{task_id}"]
    return int(item["successes"]), int(item["runs"])


def _condition_total(evidence: dict[str, Any], condition: str) -> tuple[int, int, float]:
    successes = 0
    runs = 0
    disclosure = []
    for task_id in (1, 2, 3):
        task_successes, task_runs = _success(evidence, condition, task_id)
        successes += task_successes
        runs += task_runs
        if condition == "official-global-reveal":
            disclosure.append(1.0)
        else:
            disclosure.append(evidence["condition_metrics"][f"{condition}:ID{task_id}"]["atomic_disclosure_rate"])
    return successes, runs, sum(disclosure) / len(disclosure)


def _fmt_pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _experiments(document: DocumentType, evidence: dict[str, Any]) -> None:
    _heading(document, "4 实验结果及分析")
    _body(
        document,
        "实验分析按研究问题而不是按日期展开。探索性阶段用于发现可能有效的协议和错误模式；确认性阶段冻结三道官方任务、七个主要条件和每题 10 次重复，再计算配对统计。新增官方兼容全局 Reveal-All 作为独立的 30 次补充实验，不回写原有 210 次确认性矩阵。",
    )

    _heading(document, "4.1 任务与数据来源", level=2)
    task_rows = []
    for task in evidence["task_details"]:
        task_rows.append(
            [
                f"ID{task['task_id']}",
                task["name"],
                task["correct_answer"],
                "4 条公共信息 + 4 条私有事实",
                task["evidence_chain"],
            ]
        )
    _table_caption(document, "表 2  三道官方短任务及正确答案")
    _table(document, ["任务", "场景", "答案", "信息结构", "正确证据链"], task_rows, [700, 1700, 1100, 1900, 3960], font_size=7.0)
    _body(
        document,
        "三题共用撤离背景与 West City、East Town、North Hill 三个选项，但决定道路是否可通行的事实不同。选择这三题的原因不是它们容易，而是官方公开结果能够定位到任务级别；论文其余 62 题没有进入本次本地确认性实验。",
    )

    _heading(document, "4.2 实验环境与累计规模", level=2)
    totals = evidence["totals"]
    _table_caption(document, "表 3  本地实验规模与模型请求")
    scope_rows = [
        ["本地运行", f"{totals['local_runs']} 次", "六组正式实验文件累计"],
        ["讨论与投票请求", f"{totals['discussion_vote_requests']:,} 次", "DeepSeek Agent 发言与最终投票"],
        ["披露审计请求", f"{totals['audit_requests']} 次", "AI 原子事实判断；机械披露条件不额外调用"],
        ["模型请求总数", f"{totals['model_requests']:,} 次", "讨论/投票与披露审计之和"],
        ["确认性矩阵", "210 次", "3 题 × 7 条件 × 10 次"],
        ["官方兼容全局披露补充", "30 次", "3 题 × 10 次"],
    ]
    _table(document, ["项目", "规模", "说明"], scope_rows, [2400, 1800, 5160], font_size=8.3)
    _body(
        document,
        "本地模型为 DeepSeek；官方对标使用论文公开的 GPT-4.1 逐题结果。由于模型、运行服务和实现层均不同，本研究属于任务与协议复现，不是严格的同模型数值复现。",
    )

    _heading(document, "4.3 官方结果与本地复现", level=2)
    official = evidence["official_gpt41"]
    compare_rows = []
    for task_id in (1, 2, 3):
        baseline = official["baseline"]["by_task"][str(task_id)]
        reveal = official["reveal_all"]["by_task"][str(task_id)]
        fixed_success, fixed_runs = _success(evidence, "fixed-60", task_id)
        global_success, global_runs = _success(evidence, "official-global-reveal", task_id)
        compare_rows.append(
            [
                f"ID{task_id}",
                baseline["correct_answer"],
                f"{round(baseline['majority_accuracy'] * baseline['runs'])}/{baseline['runs']}",
                f"{round(reveal['majority_accuracy'] * reveal['runs'])}/{reveal['runs']}",
                f"{fixed_success}/{fixed_runs}",
                f"{global_success}/{global_runs}",
            ]
        )
    _table_caption(document, "表 4  官方 GPT-4.1 与本地 DeepSeek 的逐题结果")
    _table(
        document,
        ["任务", "正确答案", "官方基线", "官方 Reveal-All", "本地 fixed-60", "本地全局 Reveal-All"],
        compare_rows,
        [700, 1250, 1600, 1750, 1750, 2310],
        font_size=7.5,
    )
    _body(
        document,
        "官方 GPT-4.1 基线在 ID1、ID2、ID3 上分别为 11/30、4/30 和 6/30；Reveal-All 后分别为 10/10、9/10 和 10/10。本地 DeepSeek 的 fixed-60 分别为 10/10、0/10 和 7/10，说明多智能体长讨论仍可能在特定任务上稳定出错。",
    )
    _callout(
        document,
        "确认性观察：",
        "官方兼容全局 Reveal-All 在本地 ID1、ID2、ID3 均为 10/10。该结果说明三题在完整信息反复可见时可以被当前模型稳定解决，但每题 10 次的 Wilson 95% 区间约为 72.2%—100%，不能写成普遍可靠。",
        gold=True,
    )

    _heading(document, "4.4 信息披露条件的总体结果", level=2)
    condition_labels = [
        ("single-local", "单智能体局部信息"),
        ("fixed-12", "固定轮转 12"),
        ("structured-12", "结构化 12"),
        ("fixed-60", "固定轮转 60"),
        ("dynamic-60", "动态发言 60"),
        ("fixed-reveal-all", "固定 60 + 逐人机械披露"),
        ("dynamic-reveal-all", "动态 60 + 逐人机械披露"),
        ("official-global-reveal", "官方兼容全局 Reveal-All"),
    ]
    condition_rows = []
    for key, label in condition_labels:
        successes, runs, disclosure = _condition_total(evidence, key)
        condition_rows.append([label, f"{successes}/{runs}", _fmt_pct(disclosure), "3 题，各 10 次"])
    _table_caption(document, "表 5  各条件三题合计正确率与平均原子披露率")
    _table(document, ["条件", "多数票正确", "平均披露率", "样本"], condition_rows, [3500, 1700, 1700, 2460], font_size=8.1)
    _body(
        document,
        "披露率与正确率总体同向，但不是一一对应。structured-12 的平均披露率约 92.5%，正确数为 24/30；fixed-60 的平均披露率约 75.0%，正确数为 17/30。旧 fixed-reveal-all 达到 100% 披露却只有 26/30，说明事实出现后仍可能被错误解释。dynamic-reveal-all 与官方兼容全局 Reveal-All 都为 30/30，但二者的注入时机和重复方式不同。",
    )

    _heading(document, "4.5 相同预算下的发言机制对照", level=2)
    paired = evidence["paired_tests"]
    paired_rows = [
        ["fixed-60 vs dynamic-60", "17/30", "21/30", "0 vs 4", "0.125", "不显著"],
        ["fixed-reveal-all vs dynamic-reveal-all", "26/30", "30/30", "0 vs 4", "0.125", "不显著"],
        ["fixed-12 vs structured-12", "20/30", "24/30", "4 vs 8", "0.388", "不显著"],
    ]
    _table_caption(document, "表 6  配对精确 McNemar 检验")
    _table(document, ["比较", "条件 A", "条件 B", "A-only / B-only", "p 值", "α=0.05"], paired_rows, [2800, 1100, 1100, 1900, 1100, 1360], font_size=7.8)
    _body(
        document,
        f"fixed-60 与 dynamic-60 的不一致对为 {paired['fixed-60_vs_dynamic-60']['a_only']} 比 {paired['fixed-60_vs_dynamic-60']['b_only']}，动态条件点估计更高，但 p=0.125。其他两组比较的精确 p 值为 0.125 和 0.388。三个检验均未达到统计显著优势，因此 RQ2 的回答不是“动态机制有效”，而是“当前 30 对样本尚不足以证明其优于固定轮转”。",
    )

    _heading(document, "4.6 重复稳定性与早停分析", level=2)
    early = evidence["early_stop"]
    early_rows = [
        ["同轨迹候选与最终多数一致", f"{early['classifications']['same-majority']}/{early['runs']}", "只比较同一完整轨迹中的中间影子投票"],
        ["最终平票", f"{early['classifications']['final-tie']}/{early['runs']}", "唯一异常为 2-2 final tie"],
        ["候选正确", f"{early['candidate_correct']}/{early['runs']}", "与最终正确数相同"],
        ["最终正确", f"{early['final_correct']}/{early['runs']}", "不能据此宣称普遍安全早停"],
    ]
    _table_caption(document, "表 7  同轨迹影子投票的早停分析")
    _table(document, ["指标", "结果", "解释"], early_rows, [3000, 1500, 4860], font_size=8.2)
    _body(
        document,
        "29/30 个候选多数与最终多数一致，唯一异常的最终四票为 North Hill、West City、North Hill、West City，形成 2-2 final tie。这里不存在从正确答案变为错误答案的单向改变，也不能把独立短预算运行与完整轨迹的中间状态混为一谈。",
    )

    _heading(document, "4.7 典型案例与 MAST 映射", level=2)
    mast_rows = []
    for item in evidence["mast_cases"]:
        location = "未确认"
        if item["run_id"]:
            location = f"{item['run_id']} / {item['study_key']['condition']} / ID{item['study_key']['task_id']}"
        mast_rows.append(
            [
                item["mast_mode"],
                item["boundary"],
                location,
                "—" if item["disclosure_rate"] is None else _fmt_pct(item["disclosure_rate"]),
                item["interpretation"],
            ]
        )
    _table_caption(document, "表 8  MAST 失败模式在本地轨迹中的证据位置")
    _table(document, ["模式", "判断", "运行位置", "披露率", "解释"], mast_rows, [900, 1800, 2600, 1000, 3060], font_size=7.1)
    _body(
        document,
        "FM-2.4 在 ID1 的 fixed-60 运行 383ab32a67ba98e7 中得到消息级确认：至少一条所有者事实未进入公开消息。该运行最终答案仍然正确，说明披露不足是风险机制而不是必然失败。FM-2.5 在 ID2 的 fixed-reveal-all 运行 e46521e6691a10ca 中是人工因果编码候选：四条事实全部公开，但四名 Agent 均把 East Town 隧道理解为可通行并一致选择错误答案。FM-2.6 未找到理由与最终行动直接冲突的可靠证据，因此不确认。",
    )
    _body(
        document,
        "这三个判断回答了 RQ3：错误共识至少可能由信息未披露和正确证据被错误解释产生；“信息已经出现但被忽略”仍需两名人类审阅者根据完整上下文进行独立编码。Codex 的 144 项结果是模型辅助复核，不是人工金标准。",
    )

    _heading(document, "4.8 研究问题回答与讨论", level=2)
    _body(
        document,
        "RQ1 的回答：在本研究三题中，官方兼容全局 Reveal-All 的 30 次运行全部正确，且优于对应的局部信息和固定长预算结果。这是稳定的观察，但不是对所有 HiddenBench 任务或所有模型的因果证明。",
        bold_lead="RQ1 的回答：",
    )
    _body(
        document,
        "RQ2 的回答：动态发言在两个 60 发言对照中的成功数均高于固定轮转，但精确检验 p 值均为 0.125，未达到 0.05。因此只能报告方向性差异，不能声称动态调度具有统计显著优势。",
        bold_lead="RQ2 的回答：",
    )
    _body(
        document,
        "RQ3 的回答：错误不是单一的“没有说出来”。信息链包含披露、读取、解释和行动四个环节；任一环节出错都可能形成错误共识。思维退化[12]、共识依赖[13]、困难的责任归因[14]以及社会影响造成的错误收敛[17]都与本地案例相符，但本研究只对有消息证据的具体运行作有限映射。",
        bold_lead="RQ3 的回答：",
    )
    _body(
        document,
        "综合而言，当前最稳健的工程建议不是继续增加讨论轮数，而是建立可审计的信息管道：先确保关键事实可见，再验证事实是否被读取和用于理由，最后检查投票与理由是否一致。动态选择器只有在不依赖不可获得的全局私有事实、且在同预算下稳定改善结果时，才具有独立部署价值。",
    )


def _conclusion(document: DocumentType) -> None:
    _heading(document, "5 总结与展望")
    _body(
        document,
        "本文围绕 HiddenBench 私有信息任务，在 MAF 执行层上独立实现多智能体讨论、披露、调度、投票和审计流程，并用官方短题 ID1、ID2、ID3 进行复现与对照。实验不是为了证明“多智能体一定更好”，而是确定哪些机制真正改变了结果。",
    )
    _body(
        document,
        "主要结论有三点。第一，在当前三题中，官方兼容全局 Reveal-All 的本地结果为 30/30，说明关键信息完整、重复地进入讨论能够显著改变任务表现。第二，同预算动态发言和结构化协议虽然具有更高点估计，但三个配对检验均未达到 0.05，现有样本不能证明调度机制存在稳定优势。第三，100% 披露仍可能伴随错误共识，表明治理目标应从“让事实出现”扩展到“让事实被读取、正确解释并落实到行动”。",
    )
    _body(
        document,
        "本文的贡献在于提供了一条可复现的负责任实验链：把官方任务、官方结果、本地模型和执行框架边界分开；用相同预算比较发言机制；用原子事实与消息证据计算披露率；保存配对统计和失败案例。动态机制未得到显著支持并非无效工作，它排除了“只改发言顺序即可稳定解决隐藏信息任务”这一过强结论。",
    )
    _body(
        document,
        "研究仍受三题、单模型、每题 10 次重复和模型辅助审计限制。下一阶段应从官方公开逐题结果中扩展到至少 10 题，保持同模型、同 Prompt 和同预算；让两名不知道 AI 标签的人类审阅者盲审 AI—规则分歧，并从一致样本中分层抽样；将 Reveal 与调度做 2×2 因子实验；开发不读取全局私有事实的调度器；最后在多个模型上复验，以区分模型效应和治理机制效应。",
    )
    _callout(
        document,
        "最终结论：",
        "本研究支持把“关键信息是否完整进入并被正确使用”作为多智能体认知治理的首要问题；仅追求更多 Agent、更长讨论或更强共识，并不能保证得到正确答案。",
        gold=True,
    )


def _reference_entry(index: int, item: dict[str, Any]) -> str:
    journal_keys = {"stasser_titus", "hidden_profile_meta", "social_influence"}
    conference_keys = {
        "hiddenbench",
        "multiagent_debate",
        "metagpt",
        "chatdev",
        "autogen",
        "mast",
        "divergent_debate",
        "failure_attribution",
        "dylan",
        "emergent_coordination",
    }
    if item["key"] == "maf":
        kind = "[EB/OL]"
    elif item["key"] in journal_keys:
        kind = "[J]"
    elif item["key"] in conference_keys:
        kind = "[C]"
    else:
        kind = "[EB/OL]"
    return (
        f"[{index}] {item['authors']}. {item['title']} {kind}. "
        f"{item['venue']}, {item['year']}. {item['identifier']}. {item['url']}"
    )


def _references(document: DocumentType, references: list[dict[str, Any]]) -> None:
    _heading(document, "参考文献")
    for index, item in enumerate(references, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.28)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15
        _font(paragraph.add_run(_reference_entry(index, item)), 9.2)


DISCUSSION_PROMPT = """You are one of four community leaders. Use the shared task description and your private fact to discuss the safest evacuation option.
State only conclusions supported by the information available to you. Read prior public messages, identify missing evidence, and revise your position when another message changes route feasibility.
Do not reveal that you are an AI model. Return one public discussion message only."""

VOTE_PROMPT = """Based on the complete public discussion, choose exactly one final option: West City, East Town, or North Hill.
Return the option and a short evidence-based rationale. Do not introduce new private facts during voting."""

ATOMIC_AUDIT_PROMPT = """Determine whether every atomic private claim was disclosed in the public discussion.

Rules:
1. Only an owner-authored message can disclose that owner's fact.
2. A faithful paraphrase counts; a polarity reversal, contradiction, or decision-relevant omission does not.
3. For disclosed=true, cite at least one owner-authored message ID and copy one exact non-empty substring from it.
4. For disclosed=false, return an empty message-ID list and quote.
5. Return exactly one judgment for every fact ID. Do not calculate the percentage.

ATOMIC PRIVATE FACTS:
- <fact_id> | owner=<agent_id> | claim=<private claim>

PUBLIC MESSAGES:
<message_id> | <agent_id> | <message content>

Return strict JSON only with one judgment per fact."""

OFFICIAL_GLOBAL_REVEAL = """[SYSTEM MECHANICAL REVEAL-ALL]
The following private facts are injected verbatim by the system:
- <private fact 1>
- <private fact 2>
- <private fact 3>
- <private fact 4>"""


def _appendices(document: DocumentType, evidence: dict[str, Any]) -> None:
    _page_break(document)
    _heading(document, "附录")
    _heading(document, "附录 A  三道任务的完整信息结构", level=2)
    for task in evidence["task_details"]:
        _heading(document, f"A.{task['task_id']}  ID{task['task_id']} | {task['name']} | 正确答案：{task['correct_answer']}", level=3)
        rows = []
        for index, text in enumerate(task["shared_information"], start=1):
            rows.append(["公共信息", str(index), text])
        for index, text in enumerate(task["private_facts"], start=1):
            rows.append(["私有事实", str(index), text])
        _table(document, ["类型", "序号", "内容"], rows, [1500, 800, 7060], font_size=7.6)
        _body(document, f"证据链：{task['evidence_chain']}", bold_lead="证据链：")

    _heading(document, "附录 B  实验 Prompt 与信息注入模板", level=2)
    _code(document, "B.1  讨论 Prompt", DISCUSSION_PROMPT)
    _code(document, "B.2  最终投票 Prompt", VOTE_PROMPT)
    _code(document, "B.3  原子披露审计 Prompt", ATOMIC_AUDIT_PROMPT)
    _code(document, "B.4  官方兼容全局 Reveal-All 模板", OFFICIAL_GLOBAL_REVEAL)

    _heading(document, "附录 C  75% 披露率完整实例", level=2)
    example = evidence["disclosure_example"]
    rows = []
    judgments = {item["fact_id"]: item for item in example["judgments"]}
    for fact in example["facts"]:
        judgment = judgments[fact["fact_id"]]
        rows.append(
            [
                fact["owner_agent_id"],
                fact["text"],
                "已披露" if judgment["disclosed"] else "未披露",
                judgment["evidence_quote"] or "无",
                ", ".join(judgment["evidence_message_ids"]) or "无",
            ]
        )
    _table_caption(document, f"表 C-1  运行 {example['run_id']} 的原子披露审计")
    _table(document, ["所有者", "原子事实", "判断", "证据引文", "消息 ID"], rows, [1000, 3560, 1100, 1900, 1800], font_size=6.9)
    _body(
        document,
        f"程序汇总：{example['arithmetic']}。前三条事实各有所有者消息中的精确引文，第四条关于大火阻断交通的事实没有公开证据。",
        bold_lead="程序汇总：",
    )

    _heading(document, "附录 D  典型运行与复现位置", level=2)
    run_rows = []
    for item in evidence["mast_cases"]:
        if item["run_id"]:
            run_rows.append(
                [
                    item["mast_mode"],
                    item["run_id"],
                    item["study_key"]["condition"],
                    f"ID{item['study_key']['task_id']} / rep {item['study_key']['repetition']}",
                    item["boundary"],
                ]
            )
    exception = evidence["early_stop"]["exceptional_cases"][0]
    run_rows.append(
        [
            "早停平票",
            exception["run_id"],
            exception["study_key"]["condition"],
            f"ID{exception['study_key']['task_id']} / rep {exception['study_key']['repetition']}",
            exception["classification"],
        ]
    )
    _table(document, ["用途", "run_id", "条件", "任务/重复", "判断"], run_rows, [1200, 2400, 1900, 1900, 1960], font_size=7.6)

    publication = evidence["publication"]
    official = evidence["official_code"]
    source_rows = [
        ["本项目仓库", publication["repository"]],
        ["实验分支", publication["branch"]],
        ["Pull Request", publication["pull_request"]],
        ["确认性报告提交", publication["report_commit"]],
        ["原始证据 Release", publication["release"]],
        ["HiddenBench 官方代码", official["repository"]],
        ["官方代码核验 commit", official["verified_commit"]],
        ["官方逐题结果", official["results"]],
    ]
    _table_caption(document, "表 D-1  代码与公开证据位置")
    _table(document, ["项目", "位置"], source_rows, [2400, 6960], font_size=7.2)

    hash_rows = [[name, digest] for name, digest in sorted(evidence["source_hashes"].items())]
    _table_caption(document, "表 D-2  论文使用的原始证据文件 SHA-256")
    _table(document, ["证据文件", "SHA-256"], hash_rows, [4300, 5060], font_size=6.5)

    _heading(document, "附录 E  术语与结论边界", level=2)
    term_rows = [
        ["MAF", "Microsoft Agent Framework；本文仅用其 Agent/模型执行能力"],
        ["披露率 D", "通过证据校验的已披露原子事实数 m 除以原子事实总数 n"],
        ["错误共识", "多数或全体 Agent 对错误选项形成一致判断"],
        ["owner-by-owner mechanical reveal", "私有事实随各自所有者发言逐条注入"],
        ["official-compatible global Reveal-All", "第一轮每条响应注入全部私有事实"],
        ["模型辅助复核", "Codex 独立标签复核；不是人类金标准"],
        ["统计不显著", "当前样本不足以拒绝条件成功概率相同的零假设；不等于两条件完全相同"],
    ]
    _table(document, ["术语", "本文含义"], term_rows, [3000, 6360], font_size=7.8)


def _validate_tables(document: DocumentType) -> None:
    for index, table in enumerate(document.tables):
        props = table._tbl.tblPr
        width = props.find(qn("w:tblW"))
        indent = props.find(qn("w:tblInd"))
        if width is None or width.get(qn("w:w")) != "9360" or width.get(qn("w:type")) != "dxa":
            raise ValueError(f"table {index} lacks fixed 9360-DXA width")
        if indent is None or indent.get(qn("w:w")) != "120" or indent.get(qn("w:type")) != "dxa":
            raise ValueError(f"table {index} lacks 120-DXA indent")


def build_report(root: Path, output: Path, desktop_output: Path | None) -> Path:
    root = root.resolve()
    output = output if output.is_absolute() else root / output
    preserved = root / PRESERVED_REPORT
    if _sha256(preserved) != PRESERVED_SHA256:
        raise ValueError("preserved complete report hash changed")
    evidence = _load_json(root / EVIDENCE_PATH)
    references = _load_json(root / REFERENCE_PATH)
    for figure in (root / ARCHITECTURE_FIGURE, root / FLOW_FIGURE):
        if not figure.exists():
            raise FileNotFoundError(figure)

    document = Document()
    _configure_paper(document)
    _title_and_abstract(document, evidence)
    _introduction(document)
    _related_work(document)
    _method(document, evidence, root)
    _experiments(document, evidence)
    _conclusion(document)
    _references(document, references)
    _appendices(document, evidence)
    _validate_tables(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    if desktop_output is not None:
        desktop_output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(output, desktop_output)
    if _sha256(preserved) != PRESERVED_SHA256:
        raise ValueError("paper generation modified the preserved complete report")
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the paper-style HiddenBench MAF experiment report.")
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, default=Path("reports") / REPORT_NAME)
    parser.add_argument("--desktop-output", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    output = args.output if args.output.is_absolute() else root / args.output
    desktop = args.desktop_output
    built = build_report(root, output, desktop)
    print(built)


if __name__ == "__main__":
    main()
