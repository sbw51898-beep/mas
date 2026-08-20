from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from lxml import etree

from reports.build_paper_style_report import (
    REFERENCE_PATH,
    _font,
    _reference_entry,
    build_report as build_base_report,
)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKGREL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W_NS, "r": R_NS, "rel": PKGREL_NS, "ct": CT_NS}

REL_TYPE_FOOTNOTES = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
)
CT_FOOTNOTES = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)
CITATION_RE = re.compile(r"\[(\d+)(?:-(\d+))?\]")
MARKER_RE = re.compile(r"\[\[FN(\d{2})\]\]")

# This is also an integrity check: the numbering must be driven by first use in
# the paper, rather than by the historical order of the JSON bibliography.
EXPECTED_OLD_ID_ORDER = [4, 5, 2, 3, 1, 11, 6, 7, 8, 17, 12, 13, 14, 9, 15, 16, 10]


def _xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )


def _expand_citation(match: re.Match[str]) -> list[int]:
    start = int(match.group(1))
    end = int(match.group(2) or start)
    if end < start:
        raise ValueError(f"invalid citation range: {match.group(0)}")
    return list(range(start, end + 1))


def _replace_body_citations(document: Document) -> list[int]:
    seen: dict[int, int] = {}
    first_use_order: list[int] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "参考文献":
            break
        for run in paragraph.runs:
            if not CITATION_RE.search(run.text):
                continue

            def replacement(match: re.Match[str]) -> str:
                markers: list[str] = []
                for old_id in _expand_citation(match):
                    if old_id not in seen:
                        new_id = len(first_use_order) + 1
                        seen[old_id] = new_id
                        first_use_order.append(old_id)
                        markers.append(f"[[FN{new_id:02d}]]")
                return "".join(markers)

            run.text = CITATION_RE.sub(replacement, run.text)

    if first_use_order != EXPECTED_OLD_ID_ORDER:
        raise ValueError(
            "unexpected first-citation order: "
            f"expected {EXPECTED_OLD_ID_ORDER}, got {first_use_order}"
        )
    return first_use_order


def _rewrite_bibliography(
    document: Document, references: list[dict[str, Any]], first_use_order: list[int]
) -> None:
    reference_heading: int | None = None
    appendix_heading: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text == "参考文献":
            reference_heading = index
        elif reference_heading is not None and text == "附录":
            appendix_heading = index
            break

    if reference_heading is None or appendix_heading is None:
        raise ValueError("could not locate bibliography boundaries")

    bibliography = [
        paragraph
        for paragraph in document.paragraphs[reference_heading + 1 : appendix_heading]
        if paragraph.text.strip()
    ]
    if len(bibliography) != len(first_use_order):
        raise ValueError(
            f"expected {len(first_use_order)} bibliography entries, got {len(bibliography)}"
        )

    for new_id, (paragraph, old_id) in enumerate(
        zip(bibliography, first_use_order), start=1
    ):
        paragraph.clear()
        paragraph.paragraph_format.keep_together = True
        _font(paragraph.add_run(_reference_entry(new_id, references[old_id - 1])), 9.2)


def _short_note(item: dict[str, Any]) -> str:
    authors = item["authors"]
    if authors == "Microsoft":
        author_label = "Microsoft"
    else:
        first_author = authors.split(",", 1)[0].split(" and ", 1)[0]
        surname = first_author.rsplit(" ", 1)[-1]
        author_label = f"{surname} et al."
    return (
        f'{author_label}, “{item["title"]},” {item["year"]}, '
        f'{item["identifier"]}.'
    )


def _next_rid(rels_root: etree._Element) -> str:
    highest = 0
    for relationship in rels_root.findall(f"{{{PKGREL_NS}}}Relationship"):
        match = re.fullmatch(r"rId(\d+)", relationship.get("Id") or "")
        if match:
            highest = max(highest, int(match.group(1)))
    return f"rId{highest + 1}"


def _ensure_relationship(rels_root: etree._Element) -> None:
    for relationship in rels_root.findall(f"{{{PKGREL_NS}}}Relationship"):
        if relationship.get("Type") == REL_TYPE_FOOTNOTES:
            relationship.set("Target", "footnotes.xml")
            return
    relationship = etree.SubElement(
        rels_root, f"{{{PKGREL_NS}}}Relationship"
    )
    relationship.set("Id", _next_rid(rels_root))
    relationship.set("Type", REL_TYPE_FOOTNOTES)
    relationship.set("Target", "footnotes.xml")


def _ensure_content_type(content_types: etree._Element) -> None:
    part_name = "/word/footnotes.xml"
    for override in content_types.findall(f"{{{CT_NS}}}Override"):
        if override.get("PartName") == part_name:
            override.set("ContentType", CT_FOOTNOTES)
            return
    override = etree.SubElement(content_types, f"{{{CT_NS}}}Override")
    override.set("PartName", part_name)
    override.set("ContentType", CT_FOOTNOTES)


def _ensure_footnote_settings(settings: etree._Element) -> None:
    footnote_pr = settings.find("w:footnotePr", namespaces=NS)
    if footnote_pr is None:
        footnote_pr = etree.SubElement(settings, f"{{{W_NS}}}footnotePr")
    for child in list(footnote_pr):
        if child.tag in {f"{{{W_NS}}}numFmt", f"{{{W_NS}}}numStart"}:
            footnote_pr.remove(child)
    num_fmt = etree.SubElement(footnote_pr, f"{{{W_NS}}}numFmt")
    num_fmt.set(f"{{{W_NS}}}val", "decimal")
    num_start = etree.SubElement(footnote_pr, f"{{{W_NS}}}numStart")
    num_start.set(f"{{{W_NS}}}val", "1")


def _ensure_styles(styles: etree._Element) -> None:
    existing = {
        style.get(f"{{{W_NS}}}styleId")
        for style in styles.findall("w:style", namespaces=NS)
    }
    if "FootnoteReference" not in existing:
        style = etree.SubElement(styles, f"{{{W_NS}}}style")
        style.set(f"{{{W_NS}}}type", "character")
        style.set(f"{{{W_NS}}}styleId", "FootnoteReference")
        style.set(f"{{{W_NS}}}customStyle", "0")
        name = etree.SubElement(style, f"{{{W_NS}}}name")
        name.set(f"{{{W_NS}}}val", "footnote reference")
        based_on = etree.SubElement(style, f"{{{W_NS}}}basedOn")
        based_on.set(f"{{{W_NS}}}val", "DefaultParagraphFont")
        ui_priority = etree.SubElement(style, f"{{{W_NS}}}uiPriority")
        ui_priority.set(f"{{{W_NS}}}val", "99")
        semi_hidden = etree.SubElement(style, f"{{{W_NS}}}semiHidden")
        unhide = etree.SubElement(style, f"{{{W_NS}}}unhideWhenUsed")
        _ = semi_hidden, unhide
        r_pr = etree.SubElement(style, f"{{{W_NS}}}rPr")
        vert_align = etree.SubElement(r_pr, f"{{{W_NS}}}vertAlign")
        vert_align.set(f"{{{W_NS}}}val", "superscript")

    if "FootnoteText" not in existing:
        style = etree.SubElement(styles, f"{{{W_NS}}}style")
        style.set(f"{{{W_NS}}}type", "paragraph")
        style.set(f"{{{W_NS}}}styleId", "FootnoteText")
        style.set(f"{{{W_NS}}}customStyle", "0")
        name = etree.SubElement(style, f"{{{W_NS}}}name")
        name.set(f"{{{W_NS}}}val", "footnote text")
        based_on = etree.SubElement(style, f"{{{W_NS}}}basedOn")
        based_on.set(f"{{{W_NS}}}val", "Normal")
        ui_priority = etree.SubElement(style, f"{{{W_NS}}}uiPriority")
        ui_priority.set(f"{{{W_NS}}}val", "99")
        semi_hidden = etree.SubElement(style, f"{{{W_NS}}}semiHidden")
        unhide = etree.SubElement(style, f"{{{W_NS}}}unhideWhenUsed")
        _ = semi_hidden, unhide
        p_pr = etree.SubElement(style, f"{{{W_NS}}}pPr")
        spacing = etree.SubElement(p_pr, f"{{{W_NS}}}spacing")
        spacing.set(f"{{{W_NS}}}after", "0")
        spacing.set(f"{{{W_NS}}}line", "240")
        spacing.set(f"{{{W_NS}}}lineRule", "auto")
        r_pr = etree.SubElement(style, f"{{{W_NS}}}rPr")
        size = etree.SubElement(r_pr, f"{{{W_NS}}}sz")
        size.set(f"{{{W_NS}}}val", "18")
        size_cs = etree.SubElement(r_pr, f"{{{W_NS}}}szCs")
        size_cs.set(f"{{{W_NS}}}val", "18")


def _style_run(run: etree._Element, style_id: str) -> None:
    r_pr = run.find("w:rPr", namespaces=NS)
    if r_pr is None:
        r_pr = etree.Element(f"{{{W_NS}}}rPr")
        run.insert(0, r_pr)
    r_style = r_pr.find("w:rStyle", namespaces=NS)
    if r_style is None:
        r_style = etree.SubElement(r_pr, f"{{{W_NS}}}rStyle")
    r_style.set(f"{{{W_NS}}}val", style_id)


def _insert_reference_at_marker(
    document_xml: etree._Element, marker: str, note_id: int
) -> None:
    matches = [
        node
        for node in document_xml.xpath(".//w:t", namespaces=NS)
        if marker in (node.text or "")
    ]
    if len(matches) != 1:
        raise ValueError(f"marker {marker} found {len(matches)} times")

    text_node = matches[0]
    before, after = (text_node.text or "").split(marker, 1)
    run = text_node.getparent()
    if run is None or run.tag != f"{{{W_NS}}}r":
        raise ValueError(f"marker {marker} is not in a normal text run")
    parent = run.getparent()
    if parent is None:
        raise ValueError(f"marker {marker} has no run parent")

    run_index = parent.index(run)
    original_r_pr = run.find("w:rPr", namespaces=NS)
    text_node.text = before

    reference_run = etree.Element(f"{{{W_NS}}}r")
    _style_run(reference_run, "FootnoteReference")
    reference = etree.SubElement(reference_run, f"{{{W_NS}}}footnoteReference")
    reference.set(f"{{{W_NS}}}id", str(note_id))
    parent.insert(run_index + 1, reference_run)

    if after:
        after_run = etree.Element(f"{{{W_NS}}}r")
        if original_r_pr is not None:
            after_run.append(copy.deepcopy(original_r_pr))
        after_text = etree.SubElement(after_run, f"{{{W_NS}}}t")
        if after[0].isspace() or after[-1].isspace():
            after_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        after_text.text = after
        parent.insert(run_index + 2, after_run)

    if not before and len(run) == (1 if original_r_pr is not None else 0):
        parent.remove(run)


def _make_footnotes(note_texts: list[str]) -> etree._Element:
    root = etree.Element(f"{{{W_NS}}}footnotes", nsmap={"w": W_NS, "r": R_NS})

    def separator(note_id: int, tag: str) -> None:
        note = etree.SubElement(root, f"{{{W_NS}}}footnote")
        note.set(f"{{{W_NS}}}id", str(note_id))
        note.set(f"{{{W_NS}}}type", tag)
        paragraph = etree.SubElement(note, f"{{{W_NS}}}p")
        run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        etree.SubElement(run, f"{{{W_NS}}}{tag}")

    separator(-1, "separator")
    separator(0, "continuationSeparator")

    for note_id, note_text in enumerate(note_texts, start=1):
        note = etree.SubElement(root, f"{{{W_NS}}}footnote")
        note.set(f"{{{W_NS}}}id", str(note_id))
        paragraph = etree.SubElement(note, f"{{{W_NS}}}p")
        p_pr = etree.SubElement(paragraph, f"{{{W_NS}}}pPr")
        p_style = etree.SubElement(p_pr, f"{{{W_NS}}}pStyle")
        p_style.set(f"{{{W_NS}}}val", "FootnoteText")

        number_run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        _style_run(number_run, "FootnoteReference")
        etree.SubElement(number_run, f"{{{W_NS}}}footnoteRef")

        text_run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        text = etree.SubElement(text_run, f"{{{W_NS}}}t")
        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = " " + note_text
    return root


def _patch_true_footnotes(
    source: Path, output: Path, note_texts: list[str]
) -> None:
    with zipfile.ZipFile(source, "r") as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        relationships = etree.fromstring(
            archive.read("word/_rels/document.xml.rels")
        )
        content_types = etree.fromstring(archive.read("[Content_Types].xml"))
        settings = etree.fromstring(archive.read("word/settings.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))

        # Descending order prevents a later marker in the same original run from
        # being moved ahead of an earlier reference while that run is split.
        for note_id in range(len(note_texts), 0, -1):
            _insert_reference_at_marker(document_xml, f"[[FN{note_id:02d}]]", note_id)

        remaining_markers = "".join(
            node.text or "" for node in document_xml.xpath(".//w:t", namespaces=NS)
        )
        if MARKER_RE.search(remaining_markers):
            raise ValueError("unconverted footnote marker remains in document body")

        _ensure_relationship(relationships)
        _ensure_content_type(content_types)
        _ensure_footnote_settings(settings)
        _ensure_styles(styles)
        footnotes = _make_footnotes(note_texts)

        replacements = {
            "word/document.xml": _xml_bytes(document_xml),
            "word/_rels/document.xml.rels": _xml_bytes(relationships),
            "[Content_Types].xml": _xml_bytes(content_types),
            "word/settings.xml": _xml_bytes(settings),
            "word/styles.xml": _xml_bytes(styles),
            "word/footnotes.xml": _xml_bytes(footnotes),
        }
        existing = set(archive.namelist())
        output.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as patched:
            for info in archive.infolist():
                if info.filename in replacements:
                    patched.writestr(info, replacements[info.filename])
                else:
                    patched.writestr(info, archive.read(info.filename))
            for name, data in replacements.items():
                if name not in existing:
                    patched.writestr(name, data)


def build_footnote_report(
    root: Path, output: Path, desktop_output: Path | None
) -> Path:
    root = root.resolve()
    output = output if output.is_absolute() else root / output
    desktop_output = (
        desktop_output
        if desktop_output is None or desktop_output.is_absolute()
        else root / desktop_output
    )
    references: list[dict[str, Any]] = json.loads(
        (root / REFERENCE_PATH).read_text(encoding="utf-8")
    )

    with tempfile.TemporaryDirectory(prefix="maf-paper-footnotes-") as temp_dir:
        temp = Path(temp_dir)
        base = temp / "base.docx"
        marked = temp / "marked.docx"
        build_base_report(root, base, None)
        document = Document(base)
        first_use_order = _replace_body_citations(document)
        _rewrite_bibliography(document, references, first_use_order)
        document.save(marked)
        note_texts = [_short_note(references[old_id - 1]) for old_id in first_use_order]
        _patch_true_footnotes(marked, output, note_texts)

    if desktop_output is not None:
        desktop_output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(output, desktop_output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build the paper report with true sequential Word footnotes."
    )
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(
            "reports/基于MAF的多智能体私有信息披露机制复现与对比研究_脚注版_2026-08-05.docx"
        ),
    )
    parser.add_argument("--desktop-output", type=Path)
    args = parser.parse_args()
    built = build_footnote_report(args.root, args.output, args.desktop_output)
    print(built)


if __name__ == "__main__":
    main()
