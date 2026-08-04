from __future__ import annotations

import json
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from mas_experiment.hiddenbench_atomic_disclosure import (
    build_atomic_disclosure_prompt,
)
from mas_experiment.hiddenbench_confirmatory_study import ConfirmatoryRunRecord


TABLE_WIDTH = 9360
NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(90, 98, 108)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F4F7"

CONDITION_LABELS = {
    "fixed-60": "固定轮转 60",
    "dynamic-60": "MAF 五因子动态 60",
    "fixed-reveal-all": "固定 60 + 机械全披露",
    "dynamic-reveal-all": "动态 60 + 机械全披露",
    "structured-12": "Exchange→Decide 12",
    "fixed-12": "固定轮转 12",
    "single-local": "局部信息单智能体",
}


def _font(run, size=10.5, *, bold=False, color=None, mono=False) -> None:
    western = "Consolas" if mono else "Calibri"
    east = "Microsoft YaHei"
    run.font.name = western
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), western)
    fonts.set(qn("w:hAnsi"), western)
    fonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def _margins(cell) -> None:
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for edge, value in (("top", 70), ("start", 90), ("bottom", 70), ("end", 90)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def _geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH}")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        props.append(width)
    width.set(qn("w:w"), str(TABLE_WIDTH))
    width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            tc_w = props.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                props.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            _margins(cell)


def _table(
    doc: DocumentType,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int],
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _shade(cell, "2E74B5")
        for run in cell.paragraphs[0].runs:
            _font(run, 8.5, bold=True, color=WHITE)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            if row_index % 2 == 0:
                _shade(cells[index], LIGHT_GRAY)
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    _font(run, 8.5)
    _geometry(table, widths)
    doc.add_paragraph()


def _heading(doc: DocumentType, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        _font(run, 15 if level == 1 else 12, bold=True, color=NAVY)


def _paragraph(
    doc: DocumentType,
    text: str,
    *,
    bold_lead: str | None = None,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _font(lead, bold=True, color=NAVY)
        rest = paragraph.add_run(text[len(bold_lead) :])
        _font(rest)
    else:
        run = paragraph.add_run(text)
        _font(run)


def _bullet(doc: DocumentType, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    _font(run)


def _code(doc: DocumentType, title: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(title)
    _font(run, 10, bold=True, color=NAVY)
    for line in text.splitlines() or [""]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line or " ")
        _font(run, 7.5, mono=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F6F8FA")
        paragraph._p.get_or_add_pPr().append(shading)


def _configure(doc: DocumentType) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    header = section.header.paragraphs[0]
    header.text = "HiddenBench × MAF 确认性修订实验｜2026-08-04"
    for run in header.runs:
        _font(run, 8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("MAS 知识治理项目｜可复现确认性报告")
    _font(run, 8.5, color=GRAY)


def _select_records(path: Path) -> dict[str, dict[str, Any]]:
    wanted = {
        "fixed-60",
        "fixed-reveal-all",
        "structured-12",
        "fixed-12",
    }
    result = {}
    with path.open(encoding="utf-8") as stream:
        for line in stream:
            row = json.loads(line)
            key = row["key"]
            if (
                key["task_id"] == 1
                and key["repetition"] == 0
                and key["condition"] in wanted
            ):
                result[key["condition"]] = row
                if set(result) == wanted:
                    break
    if set(result) != wanted:
        raise ValueError("prompt exemplars are incomplete")
    return result


def _pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _full_profile_metrics(runs_path: Path) -> dict[int, dict[str, float]]:
    values: dict[int, list[tuple[bool, float]]] = defaultdict(list)
    with runs_path.open(encoding="utf-8") as stream:
        for line in stream:
            row = json.loads(line)
            if row["key"]["condition"] != "fixed-60":
                continue
            run = row["run"]
            votes = [item["vote"] for item in run["full_profile_votes"]]
            ranked = Counter(votes).most_common()
            majority = (
                ranked[0][0]
                if len(ranked) == 1 or ranked[0][1] > ranked[1][1]
                else None
            )
            correct = run["task"]["correct_answer"]
            values[row["key"]["task_id"]].append(
                (majority == correct, sum(v == correct for v in votes) / 4)
            )
    return {
        task_id: {
            "majority": sum(item[0] for item in rows) / len(rows),
            "vote": sum(item[1] for item in rows) / len(rows),
        }
        for task_id, rows in values.items()
    }


def _title(doc: DocumentType) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(72)
    run = paragraph.add_run("HiddenBench 确认性修订实验报告")
    _font(run, 24, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Microsoft Agent Framework 复现、原子披露与同轨迹早停")
    _font(run, 14, color=BLUE)
    for text in (
        "正式范围：ID1、ID2、ID3｜7 条件 × 3 题 × 10 次 = 210 次",
        "模型：DeepSeek V4 Flash｜温度 0｜Thinking disabled",
        "日期：2026 年 8 月 4 日",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        _font(r, 10.5, color=GRAY)
    doc.add_page_break()


def build_report(
    *,
    root: Path,
    output: Path,
    desktop_output: Path | None,
) -> Path:
    summary = json.loads(
        (root / "reports/data/hiddenbench-confirmatory-summary.json").read_text(
            encoding="utf-8"
        )
    )
    official = json.loads(
        (
            root / "reports/data/hiddenbench-official-gpt41-short-summary.json"
        ).read_text(encoding="utf-8")
    )
    if summary["run_count"] != 210 or summary["task_ids"] != [1, 2, 3]:
        raise ValueError("report requires the locked 210-run ID1/2/3 summary")
    if not summary["gate_passed"]:
        raise ValueError("report generation requires a passed data gate")
    runs_path = root / "artifacts/hiddenbench-confirmatory-20260804.jsonl"
    examples = _select_records(runs_path)
    fixed12 = ConfirmatoryRunRecord.model_validate(examples["fixed-12"])
    atomic_prompt = build_atomic_disclosure_prompt(fixed12.run)
    full_profile = _full_profile_metrics(runs_path)

    doc = Document()
    _configure(doc)
    _title(doc)

    _heading(doc, "摘要")
    _paragraph(
        doc,
        "本轮只研究论文短基准的 ID1、ID2、ID3，不把此前 ID1/5/7/25 的探索性结果混入确认性结论。"
        "固定轮转、动态发言、机械全披露、Structured、12 发言和局部信息单智能体共 7 个条件，"
        "每题重复 10 次，共 210 次。讨论与投票使用 11,971 次 API 请求，原子披露审计使用 152 次。",
    )
    _paragraph(
        doc,
        "最重要的结果是：动态机械全披露在三题均达到 100% 多数正确；固定机械全披露在 ID2 只有 60%，"
        "说明“信息出现过”不等于“信息被正确使用”。固定 60 的 30 次运行都出现候选早停点，29/30 与最终多数一致，"
        "但观察到 1 次翻转，因此不能宣称提前停止普遍安全。",
    )

    _heading(doc, "1. 对老师问题的直接回答")
    answers = [
        ["原论文用什么？", "作者自定义 Python 模拟器；短基准对标模型为 GPT-4.1。不是 Microsoft Agent Framework。"],
        ["原论文有没有源码？", "有。公开仓库含 simulator.py、metrics.py、65 题数据、3 题 short benchmark、提示词和结果目录。"],
        ["我们为什么用 MAF？", "MAF 负责 OpenAI-compatible/DeepSeek Agent 调用与会话封装；题目、轮转、动态选择和门禁由本项目显式控制。"],
        ["为何前几轮一致还继续？", "为了比较同一条 60 发言轨迹。每轮影子投票不回灌讨论；事后评估候选早停。结果确有 1 次翻转。"],
        ["一条私有信息首轮不就披露完？", "不会。普通条件由模型自行决定说多少；审计逐个原子事实核验。只有 Reveal-All 由系统机械附加，保证 100%。"],
        ["披露法则是什么？", "普通条件不随机、不强制；模型自然发言。Reveal-All 在每个 Agent 第一次公开发言后机械附加其私有事实。"],
        ["披露百分比在哪？", "每次运行按 disclosed 原子事实数 / 原子事实总数计算，AI 提供消息 ID 和原文引文，规则程序再验所有者与子串。"],
    ]
    _table(doc, ["问题", "本报告的回答"], answers, [2100, 7260])

    _heading(doc, "2. 原论文、官方对标与本项目")
    _paragraph(
        doc,
        "HiddenBench 论文：Li、Naito、Shirado，arXiv:2505.11556。论文以 Hidden Profile 范式测试分布式信息整合。"
        "原论文运行系统是作者自定义 Python 模拟器，不是 Microsoft Agent Framework。"
        "作者公开仓库为 https://github.com/Yassellee/HiddenBench_ICML；官方逐轮讨论和投票结果位于 "
        "https://huggingface.co/datasets/YuxuanLi1225/HiddenBench-results。",
    )
    _table(
        doc,
        ["项目", "原论文/官方", "本确认性实验"],
        [
            ["框架", "作者自定义 Python 模拟器", "MAF 模型适配 + 本项目显式协议"],
            ["模型", "GPT-4.1", "deepseek-v4-flash"],
            ["题目", "short benchmark：ID1/2/3", "完全相同 ID1/2/3"],
            ["Agent", "4", "4；单智能体条件为 1"],
            ["讨论", "15 轮、每轮 4 人", "60 发言条件相同；另设 12 发言"],
            ["Reveal-All", "首位发言后公开全部隐藏信息", "每位 Agent 首次发言机械附加其私有事实"],
        ],
        [1400, 3780, 4180],
    )

    _heading(doc, "3. 实验设置")
    _paragraph(
        doc,
        "三道验证题分别为 ID1 evacuation_west_city、ID2 evacuation_north_hill、ID3 evacuation_east_town。"
        "同一题同一次重复的七个条件使用相同 pair seed 和相同 assignment fingerprint。DeepSeek 温度为 0，"
        "thinking=disabled；动态发言选择器每轮不调用额外 LLM。",
    )
    condition_rows = [
        ["fixed-60", "60", "固定轮转", "15 轮影子投票"],
        ["dynamic-60", "60", "五因子闭式选择", "每人严格 15 次"],
        ["fixed-reveal-all", "60", "固定轮转", "首轮机械披露"],
        ["dynamic-reveal-all", "60", "五因子闭式选择", "首次发言机械披露"],
        ["structured-12", "12", "2 轮 Exchange + 1 轮 Decide", "结构化短协议"],
        ["fixed-12", "12", "固定轮转", "同预算对照"],
        ["single-local", "0", "单 Agent 直接投票", "只见共享信息+自己一条私有信息"],
    ]
    _table(doc, ["条件", "公开发言", "机制", "说明"], condition_rows, [2100, 1100, 2800, 3360])
    _paragraph(
        doc,
        "关于 b_i：本轮不让 AI 自报一个连续的“信念值”，也不把它伪装成物理量。b_i 在概念上表示 Agent i 当前的信念/立场；"
        "在本实验中以离散投票（West City、East Town、North Hill）和投票是否正确来记录。动态选择器的分歧项读取这些离散立场，"
        "其余四项为未披露信息、相关讨论、待回应程度和等待时间。权重为 0.30、0.30、0.15、0.15、0.10。",
    )

    _heading(doc, "4. 指标与披露率")
    for text in (
        "平均投票正确率：最终正确票数 / 最终票数。",
        "多数正确率：一次运行中严格多数答案是否为标准答案。",
        "错误共识率：四个 Agent 全体一致但答案错误的运行比例。",
        "原子事实披露率：AI 逐项判断 disclosed，必须给出拥有者消息 ID 和原文引文；规则程序验证拥有者、消息存在、原文子串和极性。",
        "披露百分比由程序对 AI 的逐项判断求和得到，而不是让 AI 自己随口报一个百分数。Reveal-All 的精确注入无需 AI 判断即可机械确认。",
    ):
        _bullet(doc, text)

    _heading(doc, "5. 官方 GPT-4.1 逐题对标")
    official_rows = []
    for task_id in (1, 2, 3):
        baseline = official["baseline"]["by_task"][str(task_id)]
        reveal = official["reveal_all"]["by_task"][str(task_id)]
        official_rows.append(
            [
                f"ID{task_id}",
                baseline["runs"],
                _pct(baseline["majority_accuracy"]),
                _pct(baseline["false_consensus_rate"]),
                reveal["runs"],
                _pct(reveal["majority_accuracy"]),
                _pct(reveal["false_consensus_rate"]),
            ]
        )
    _table(
        doc,
        ["题", "Baseline n", "Baseline 多数正确", "Baseline 错误共识", "Reveal n", "Reveal 多数正确", "Reveal 错误共识"],
        official_rows,
        [700, 900, 1500, 1500, 800, 1800, 2160],
    )
    _paragraph(
        doc,
        "官方 baseline 原始文件为 90 次（每题 30 次），Reveal-All 为 30 次（每题 10 次）。"
        "两文件 SHA-256 分别为 b470918c6ef2fa50ca6b237caf4700b37800f30e032416061ade00dedc5cab18 和 "
        "de2ab0d39ebc6fa61c387cd5e63402924c8745df6fe5bb09bc52a938d52079bc。",
    )

    _heading(doc, "6. 本轮 210 次结果")
    current_rows = []
    for condition in CONDITION_LABELS:
        for task_id in (1, 2, 3):
            value = summary["by_condition_task"][f"{condition}:ID{task_id}"]
            current_rows.append(
                [
                    CONDITION_LABELS[condition],
                    f"ID{task_id}",
                    _pct(value["majority_accuracy"]),
                    _pct(value["vote_accuracy"]),
                    _pct(value["atomic_disclosure_rate"]),
                    _pct(value["false_consensus_rate"]),
                    value["runs"],
                ]
            )
    _table(
        doc,
        ["条件", "题", "多数正确", "平均票正确", "原子披露率", "错误共识", "n"],
        current_rows,
        [2100, 600, 1150, 1150, 1250, 1250, 1860],
    )
    _paragraph(
        doc,
        "结果不能简化为“多智能体一定优于单智能体”。动态全披露三题全对，是最稳定的治理条件；"
        "但固定全披露在 ID2 仍只有 60% 多数正确，说明调度与解释过程仍影响结果。Structured-12 在三题为 70%、70%、100%，"
        "用较少公开发言获得了较高的整体表现。",
    )

    _heading(doc, "7. 局部信息、完整信息与多智能体的公平解释")
    full_rows = [
        [
            f"ID{task_id}",
            _pct(summary["by_condition_task"][f"single-local:ID{task_id}"]["majority_accuracy"]),
            _pct(full_profile[task_id]["majority"]),
            _pct(summary["by_condition_task"][f"fixed-60:ID{task_id}"]["majority_accuracy"]),
            _pct(summary["by_condition_task"][f"dynamic-60:ID{task_id}"]["majority_accuracy"]),
        ]
        for task_id in (1, 2, 3)
    ]
    _table(
        doc,
        ["题", "局部信息单智能体", "完整信息单智能体/Full Profile", "固定多智能体", "动态多智能体"],
        full_rows,
        [900, 2000, 2400, 1900, 2160],
    )
    _paragraph(
        doc,
        "完整信息单智能体是能力上界对照，不与局部信息单智能体混为一谈。值得注意的是 DeepSeek 的 Full Profile 在 ID3 多数正确率为 0%，"
        "所以“给全信息”也不保证该模型正确；这正是必须同时报告模型能力与协作治理的原因。",
    )

    _heading(doc, "8. 同轨迹影子投票与早停")
    early = summary["early_stop"]
    _table(
        doc,
        ["固定 60 运行", "有候选早停", "与最终一致", "发生翻转", "候选平均节省公开发言"],
        [[30, early["candidate_runs"], early["exact_match_true"], early["exact_match_false"], f"{early['mean_saved_messages_candidates']:.1f}"]],
        [1700, 1800, 1800, 1500, 2560],
    )
    _paragraph(
        doc,
        "候选规则要求连续两轮、每轮四票全体一致且答案完全相同。30 次均出现候选点，29/30 与最终多数一致，"
        "但有 1 次翻转。因此只能说“本样本中一致率为 96.7%”，不能说早停一般安全。影子投票从未进入后续讨论 prompt，"
        "所以这是同一条 60 发言轨迹上的事后比较。",
    )

    _heading(doc, "9. MAST 失败模式在本实验中的位置")
    _table(
        doc,
        ["MAST 模式", "含义", "本实验的可观测证据", "结论边界"],
        [
            ["FM-2.4", "重要信息未分享", "非 Reveal 条件原子披露率低于 100%，如 fixed-60：ID1/2/3 为 75%/70%/80%。", "可作为直接操作化指标。"],
            ["FM-2.5", "忽略其他 Agent 输入", "fixed-reveal-all 的 ID2 已 100% 披露但错误共识率仍 30%。", "是候选证据；需逐轮人工编码才能断言忽略。"],
            ["FM-2.6", "推理与行动不一致", "需要比较最终 rationale 与最终 vote，当前自动门禁没有正式标注。", "本报告不虚构发生率。"],
        ],
        [1100, 1900, 3900, 2460],
    )
    _paragraph(
        doc,
        "因此，本实验正式测量的是信息是否披露、是否形成错误共识以及早停是否翻转。FM-2.5/2.6 的因果归类仍需要人工盲审逐轮文本，"
        "不能只凭最终错误自动贴标签。对应关系为 FM-2.4 信息未分享、FM-2.5 忽略他人输入、FM-2.6 推理与行动不一致。",
    )

    _heading(doc, "10. 版本、数据门禁与复现")
    _paragraph(
        doc,
        "运行提交：80f621b223c2a7f0f88434a083c5cf5268d52b51；冻结运行代码提交："
        "1aa577992d3abb3c83e5aeb7ce5475de7100e7af；证据重定位审计代码提交："
        "e1b32d3c217960bc331c0c751217011e92dd9d2b。",
    )
    _paragraph(
        doc,
        "项目仓库：https://github.com/sbw51898-beep/mas；固定分支：codex/budget-matched-dynamic；"
        "预定发布页：https://github.com/sbw51898-beep/mas/releases/tag/hiddenbench-confirmatory-20260804。",
    )
    source_rows = [[name, digest] for name, digest in summary["source_files"].items()]
    _table(doc, ["原始文件", "SHA-256"], source_rows, [3000, 6360])
    _paragraph(doc, "门禁 15 项全部通过；讨论/投票 API 请求 11,971，审计 API 请求 152。")

    _heading(doc, "11. 局限")
    for text in (
        "确认性样本只有三道短题，不能外推到全部 65 题。",
        "本地使用 DeepSeek，不是官方 GPT-4.1；框架与模型效应没有被完全正交分离。",
        "温度 0 不能消除服务端非确定性，10 次重复只描述本次样本稳定性。",
        "AI 审计不是人工金标准；规则只保证引文来源真实、所有者正确和极性一致。",
        "审计阶段因模型引文轻微改写触发门禁；修复仅在合法拥有者消息内重定位原文，并重新审计全部 210 次。",
    ):
        _bullet(doc, text)

    doc.add_section(WD_SECTION.NEW_PAGE)
    _heading(doc, "附录 A：本轮实际提示词与机械披露")
    fixed = examples["fixed-60"]["run"]
    _code(doc, "A1. ID1 Agent A 实际 system prompt", fixed["hidden_pre_votes"][0]["system_prompt"])
    _code(doc, "A2. 首位发言 user prompt", fixed["discussion_messages"][0]["user_prompt"])
    _code(doc, "A3. 后续发言 user prompt 示例", fixed["discussion_messages"][1]["user_prompt"])
    reveal = examples["fixed-reveal-all"]["run"]["discussion_messages"][0]
    _code(doc, "A4. Reveal-All 存储后的公开消息（含系统机械块）", reveal["content"])
    structured = examples["structured-12"]["run"]["discussion_messages"]
    _code(doc, "A5. Structured Exchange prompt", structured[0]["user_prompt"])
    _code(doc, "A6. Structured Decide prompt", structured[8]["user_prompt"])
    checkpoint = examples["fixed-60"]["shadow_checkpoints"][0]
    _code(doc, "A7. 第一轮后影子投票 prompt", checkpoint["votes"][0]["user_prompt"])
    _code(doc, "A8. 原子事实 AI 审计完整 prompt", atomic_prompt)

    doc.core_properties.title = "HiddenBench 确认性修订实验报告"
    doc.core_properties.subject = "MAF 多智能体分布式信息整合、披露与早停"
    doc.core_properties.author = "MAS 知识治理项目组"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    check = Document(output)
    if len(check.tables) < 8:
        raise ValueError("report table count is incomplete")
    for table in check.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        if width is None or int(width.get(qn("w:w"))) != TABLE_WIDTH:
            raise ValueError("report table geometry validation failed")
    if desktop_output is not None:
        desktop_output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(output, desktop_output)
    return output


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    output = root / "reports" / "HiddenBench确认性修订实验报告_2026-08-04.docx"
    desktop = Path.home() / "Desktop" / output.name
    print(build_report(root=root, output=output, desktop_output=desktop))


if __name__ == "__main__":
    main()
