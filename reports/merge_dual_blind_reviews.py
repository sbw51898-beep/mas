"""Validate and merge two independently completed human blind-review sheets.

The merger is intentionally conservative: it can calculate a final human
metric only when both reviewers have completed valid forms *and* agree on
every label.  Otherwise it records the agreement rate and the conflict list
but leaves final AI precision/recall and revised disclosure rate pending human
adjudication.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from mas_experiment.hiddenbench_ai_disclosure import DisclosureAudit
from mas_experiment.hiddenbench_blind_review import (
    BlindReviewCase,
    BlindReviewMessage,
    ReviewJudgment,
    compute_weighted_review_metrics,
    make_blind_id,
    parse_human_disclosed,
    validate_review_judgment_evidence,
)
from mas_experiment.hiddenbench_stability_domain import StudyRunRecord
from mas_experiment.hiddenbench_stability_reporting import build_fact_comparisons


ROOT = Path(__file__).resolve().parents[1]
QUEUE_PATH = Path(
    "artifacts/hiddenbench-stability-20260729.blind-review.queue.csv"
)
RUNS_PATH = Path("artifacts/hiddenbench-stability-20260729.jsonl")
AUDITS_PATH = Path(
    "artifacts/hiddenbench-stability-20260729.ai-disclosure.jsonl"
)
BLIND_SEED = 20260730
MODEL_OR_TOOL_NAMES = {
    "ai",
    "assistant",
    "chatgpt",
    "claude",
    "codex",
    "deepseek",
    "gemini",
    "gpt",
    "llm",
    "model",
    "unassigned",
}


def _build_population(root: Path) -> tuple[BlindReviewCase, ...]:
    records = tuple(
        StudyRunRecord.model_validate_json(line)
        for line in (root / RUNS_PATH).read_text(encoding="utf-8").splitlines()
        if line.strip()
    )
    audits = tuple(
        DisclosureAudit.model_validate_json(line)
        for line in (root / AUDITS_PATH).read_text(encoding="utf-8").splitlines()
        if line.strip()
    )
    audit_by_key = {item.study_key.value: item for item in audits}
    population: list[BlindReviewCase] = []
    for record in records:
        audit = audit_by_key[record.key.value]
        comparisons = build_fact_comparisons(record, audit)
        comparison_by_fact = {item.fact_id: item for item in comparisons}
        for owner_agent_id, fact in record.run.assignment.private_information.items():
            fact_id = f"private-fact:{owner_agent_id}"
            comparison = comparison_by_fact[fact_id]
            owner_messages = tuple(
                BlindReviewMessage(
                    message_id=message.message_id,
                    turn_index=message.turn_index,
                    content=message.content,
                )
                for message in record.run.discussion_messages
                if message.agent_id == owner_agent_id
            )
            population.append(
                BlindReviewCase(
                    blind_id=make_blind_id(
                        seed=BLIND_SEED,
                        study_key=record.key.value,
                        fact_id=fact_id,
                    ),
                    study_key=record.key.value,
                    task_id=record.key.task_id,
                    condition=record.key.condition,
                    fact_id=fact_id,
                    owner_agent_id=owner_agent_id,
                    fact=fact,
                    owner_messages=owner_messages,
                    ai_disclosed=comparison.ai_disclosed,
                    rule_disclosed=comparison.rule_disclosed,
                )
            )
    if len(population) != 320:
        raise ValueError(f"expected 320 population cases, found {len(population)}")
    return tuple(population)


def _queue_ids(root: Path) -> tuple[str, ...]:
    with (root / QUEUE_PATH).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    ids = tuple(row["blind_id"].strip() for row in rows)
    if len(ids) != 144 or not all(ids) or len(set(ids)) != len(ids):
        raise ValueError("frozen blind-review queue is not a unique 144-case set")
    return ids


def _normalize_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _validate_human_identity(name: object, review_date: object, declaration: object) -> tuple[str, str]:
    normalized_name = _normalize_name(name)
    if len(re.sub(r"[^A-Za-z\u4e00-\u9fff]", "", normalized_name)) < 2:
        raise ValueError("human reviewer name is missing or too short")
    lowered = normalized_name.casefold()
    if any(token in lowered for token in MODEL_OR_TOOL_NAMES):
        raise ValueError("human reviewer name cannot be a model or tool identifier")
    if isinstance(review_date, datetime):
        normalized_date = review_date.date().isoformat()
    elif isinstance(review_date, date):
        normalized_date = review_date.isoformat()
    else:
        try:
            normalized_date = date.fromisoformat(str(review_date or "").strip()).isoformat()
        except ValueError as error:
            raise ValueError("human reviewer date must use YYYY-MM-DD") from error
    normalized_declaration = _normalize_name(declaration)
    if "本人" not in normalized_declaration or "独立" not in normalized_declaration:
        raise ValueError("human reviewer declaration must state independent completion")
    return normalized_name, normalized_date


def _load_review_sheet(
    path: Path,
    *,
    cases_by_id: dict[str, BlindReviewCase],
    expected_ids: set[str],
) -> tuple[dict[str, Any], tuple[ReviewJudgment, ...]]:
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, data_only=True)
    try:
        guide = workbook["填写说明"]
        sheet = workbook["盲审签核"]
    except KeyError as error:
        raise ValueError(f"{path.name} does not match the dual blind-review form") from error
    reviewer_name, reviewer_date = _validate_human_identity(
        guide["B4"].value,
        guide["B5"].value,
        guide["B6"].value,
    )
    expected_headers = (
        "序号",
        "盲审编号",
        "私有信息包（待判断内容）",
        "信息拥有者",
        "拥有者公开发言（含消息ID）",
        "是否披露（是/否）",
        "证据消息ID（判“是”必填）",
        "证据原文引句（判“是”必填）",
        "理由（可选）",
    )
    actual_headers = tuple(sheet.cell(row=1, column=index).value for index in range(1, 10))
    if actual_headers != expected_headers:
        raise ValueError(f"{path.name} has unexpected review-sheet headers")

    judgments: list[ReviewJudgment] = []
    seen: set[str] = set()
    for row_index in range(2, sheet.max_row + 1):
        raw_id = sheet.cell(row=row_index, column=2).value
        if raw_id is None or not str(raw_id).strip():
            continue
        blind_id = str(raw_id).strip()
        if blind_id not in expected_ids or blind_id not in cases_by_id:
            raise ValueError(f"{path.name} contains an unknown blind ID: {blind_id}")
        if blind_id in seen:
            raise ValueError(f"{path.name} contains duplicate blind ID: {blind_id}")
        seen.add(blind_id)
        raw_decision = sheet.cell(row=row_index, column=6).value
        if raw_decision is None or not str(raw_decision).strip():
            raise ValueError(
                f"{path.name}:{blind_id} is missing an explicit disclosure decision"
            )
        try:
            disclosed = parse_human_disclosed(str(raw_decision))
        except ValueError as error:
            raise ValueError(f"{path.name}:{blind_id} has an invalid disclosure decision") from error
        evidence_ids = tuple(
            value.strip()
            for value in str(sheet.cell(row=row_index, column=7).value or "").split("|")
            if value.strip()
        )
        judgment = ReviewJudgment(
            blind_id=blind_id,
            disclosed=disclosed,
            evidence_message_ids=evidence_ids,
            evidence_quote=str(sheet.cell(row=row_index, column=8).value or "").strip(),
            reason=str(sheet.cell(row=row_index, column=9).value or "").strip(),
            reviewer_id=reviewer_name,
        )
        validate_review_judgment_evidence(cases_by_id[blind_id], judgment)
        judgments.append(judgment)
    if seen != expected_ids:
        missing = sorted(expected_ids - seen)
        raise ValueError(f"{path.name} is missing review rows: {missing[:5]}")
    return (
        {"name": reviewer_name, "date": reviewer_date, "source": str(path)},
        tuple(judgments),
    )


def merge_dual_blind_reviews(
    root: Path,
    reviewer_a_path: Path,
    reviewer_b_path: Path,
    output_dir: Path,
) -> dict[str, Any]:
    """Merge two completed sheets without inventing a human consensus."""
    root = root.resolve()
    output_dir = output_dir.resolve()
    population = _build_population(root)
    cases_by_id = {case.blind_id: case for case in population}
    expected_ids = set(_queue_ids(root))
    if not expected_ids.issubset(cases_by_id):
        raise ValueError("blind-review queue does not match the frozen population")
    metadata_a, judgments_a = _load_review_sheet(
        reviewer_a_path,
        cases_by_id=cases_by_id,
        expected_ids=expected_ids,
    )
    metadata_b, judgments_b = _load_review_sheet(
        reviewer_b_path,
        cases_by_id=cases_by_id,
        expected_ids=expected_ids,
    )
    if metadata_a["name"] == metadata_b["name"]:
        raise ValueError("the two human reviewers must have distinct declared names")

    by_id_a = {item.blind_id: item for item in judgments_a}
    by_id_b = {item.blind_id: item for item in judgments_b}
    disagreements = [
        blind_id
        for blind_id in sorted(expected_ids)
        if by_id_a[blind_id].disclosed != by_id_b[blind_id].disclosed
    ]
    agreement = (len(expected_ids) - len(disagreements)) / len(expected_ids)
    result: dict[str, Any] = {
        "schema_version": "hiddenbench-human-dual-blind-merge-v1",
        "reviewers": {"a": metadata_a, "b": metadata_b},
        "reviewed_size": len(expected_ids),
        "inter_rater_exact_agreement": agreement,
        "inter_rater_agreement_count": f"{len(expected_ids) - len(disagreements)}/{len(expected_ids)}",
        "label_disagreement_count": len(disagreements),
        "label_disagreement_ids": disagreements,
        "identity_boundary": (
            "The merger validates declared names, dates, and independent-completion "
            "statements but cannot independently prove a person's identity."
        ),
    }
    if disagreements:
        result.update(
            {
                "status": "pending_human_adjudication",
                "metrics": None,
                "next_step": (
                    "Reconcile the listed dual-review disagreements with a documented "
                    "human adjudication before reporting AI precision, recall, or a "
                    "revised disclosure rate as final."
                ),
            }
        )
    else:
        metrics = compute_weighted_review_metrics(population, judgments_a)
        result.update(
            {
                "status": "completed_human_double_blind_review",
                "metrics": metrics.model_dump(mode="json"),
                "next_step": "No label disagreement remains; the merged human labels support the reported metrics.",
            }
        )
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / "dual-blind-review-merge-20260817.json"
    output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    if disagreements:
        conflict_output = output_dir / "dual-blind-review-conflicts-20260817.csv"
        with conflict_output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=(
                    "blind_id",
                    "fact",
                    "reviewer_a_decision",
                    "reviewer_b_decision",
                ),
            )
            writer.writeheader()
            for blind_id in disagreements:
                writer.writerow(
                    {
                        "blind_id": blind_id,
                        "fact": cases_by_id[blind_id].fact,
                        "reviewer_a_decision": "是" if by_id_a[blind_id].disclosed else "否",
                        "reviewer_b_decision": "是" if by_id_b[blind_id].disclosed else "否",
                    }
                )
        result["conflict_file"] = str(conflict_output)
    result["output"] = str(output)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Merge two valid human double-blind disclosure review sheets."
    )
    parser.add_argument("reviewer_a", type=Path)
    parser.add_argument("reviewer_b", type=Path)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "reports" / "data")
    args = parser.parse_args()
    result = merge_dual_blind_reviews(
        args.root,
        args.reviewer_a,
        args.reviewer_b,
        args.output_dir,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
