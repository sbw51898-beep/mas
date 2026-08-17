from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from reports.build_paper_style_footnotes import _patch_true_footnotes, _short_note
from reports.build_paper_style_report import (
    ATOMIC_AUDIT_PROMPT,
    DISCUSSION_PROMPT,
    EVIDENCE_PATH,
    OFFICIAL_GLOBAL_REVEAL,
    REFERENCE_PATH,
    VOTE_PROMPT,
    _body,
    _callout,
    _code,
    _configure_paper,
    _figure,
    _font,
    _heading,
    _page_break,
    _reference_entry,
    _table,
    _table_caption,
    _validate_tables,
)


REPORT_NAME = "基于MAF的多智能体正确性治理与错误链路分析_2026-08-17.docx"
GOVERNANCE_FIGURE = Path("reports/figures/mas-correctness-governance-loop.png")
ORIGINAL_PAPER = Path("reports/基于MAF的多智能体私有信息披露机制复现与对比研究_2026-08-05.docx")
FOOTNOTE_PAPER = Path("reports/基于MAF的多智能体私有信息披露机制复现与对比研究_脚注版_2026-08-05.docx")
COMPLETE_REPORT = Path("reports/给彭老师的完整实验报告_2026-08-02至2026-08-04.docx")
PRESERVED_HASHES = {
    ORIGINAL_PAPER: "e1aa4be327c847ddfa3987e408a1a7a5f178f798eb8903928d95ff2a37c5f668",
    FOOTNOTE_PAPER: "4f486bf87a380b0961f69ae1e822609993297ac94c2590f6a4334851b572f575",
    COMPLETE_REPORT: "0d56f74c359e1a7359102de678c8c8e39bd568d381e9ff2ccf37189b83f72765",
}

NAVY = RGBColor(23, 54, 93)
MUTED = RGBColor(91, 101, 115)


def _load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _assert_preserved(root: Path) -> None:
    for relative, expected in PRESERVED_HASHES.items():
        path = root / relative
        if _sha256(path) != expected:
            raise ValueError(f"preserved report hash changed: {relative}")


@dataclass
class CitationTracker:
    references: list[dict[str, Any]]
    ordered_keys: list[str] = field(default_factory=list)
    _by_key: dict[str, dict[str, Any]] = field(init=False)

    def __post_init__(self) -> None:
        self._by_key = {item["key"]: item for item in self.references}

    def cite(self, key: str) -> str:
        if key not in self._by_key:
            raise KeyError(key)
        if key in self.ordered_keys:
            return ""
        self.ordered_keys.append(key)
        return f"[[FN{len(self.ordered_keys):02d}]]"

    def ordered_references(self) -> list[dict[str, Any]]:
        return [self._by_key[key] for key in self.ordered_keys]


def _configure_identity(document: DocumentType) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _font(header.add_run("HiddenBench × MAF | 多智能体正确性治理"), 8.5, color=MUTED)


def _title_and_abstract(document: DocumentType, evidence: dict[str, Any]) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(8)
    _font(title.add_run("基于 MAF 的多智能体正确性治理与错误链路分析"), 22, bold=True, color=NAVY)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(18)
    _font(meta.add_run("实验研究论文 | 2026 年 8 月 17 日"), 10.5, color=MUTED)

    _heading(document, "摘要")
    abstract = (
        "大语言模型多智能体系统通过讨论汇集分散信息，但形成共识并不意味着最终答案正确。"
        "HiddenBench 与多智能体失败研究已揭示分布式信息条件下的错误共识，却较少把错误定位、观测指标和干预机制连成一套治理过程。"
        "本文把私有事实到最终决策划分为公开、读取、解释和行动环节，并以最终正确率作为结果指标，以披露、证据利用、解释一致性和投票指标作为诊断信号。"
        "研究使用 Microsoft Agent Framework 调用 DeepSeek，复现三个具有官方逐题结果的 HiddenBench 短任务，并比较结构化交换、固定轮转、同预算动态发言和不同 Reveal-All 时序。"
        f"现有证据共包含 {evidence['totals']['local_runs']} 次本地运行和 {evidence['totals']['model_requests']:,} 次模型请求，官方兼容第一轮全局 Reveal-All 为 30/30，而逐步 fixed-reveal-all 为 26/30。"
        "固定与动态发言、固定与结构化讨论的配对检验均未达到统计显著，说明当前样本只支持方向性差异。"
        "结果表明，多智能体正确性治理应先定位信息公开、利用、解释或决策失败，再实施针对性干预，而不能用披露率或共识强度替代正确性检验。"
    )
    _body(document, abstract)

    keywords = document.add_paragraph()
    keywords.paragraph_format.space_after = Pt(12)
    _font(keywords.add_run("关键词："), 10.5, bold=True)
    _font(
        keywords.add_run(
            "多智能体系统；Microsoft Agent Framework；HiddenBench；正确性治理；错误共识；信息披露"
        ),
        10.5,
    )
    _callout(
        document,
        "核心观点：",
        "多智能体治理的目标是提高最终决策正确率；披露率、发言顺序和共识程度只是定位错误的中间观测量。",
        gold=True,
    )


def _introduction(document: DocumentType, citations: CitationTracker) -> None:
    _heading(document, "1 引言")
    _body(
        document,
        "大语言模型驱动的多智能体系统把复杂任务交给多个具有不同角色、工具或信息的 Agent，再通过消息交换和投票形成答案。相关综述把角色分工、通信结构、记忆和工作流视为系统能力的重要来源"
        + citations.cite("llm_mas_survey")
        + "；多智能体辩论则表明，相互审查可能改善部分事实性和推理任务"
        + citations.cite("multiagent_debate")
        + "。但参与者更多、讨论更长或共识更强，都不能直接保证答案正确。",
    )
    _body(
        document,
        "信息不对称任务尤其容易暴露这一问题。Hidden Profile 范式让每个成员只掌握一部分决定性信息，正确答案只有在讨论中汇集未共享信息后才会出现"
        + citations.cite("stasser_titus")
        + "；长期研究表明，群体往往重复公共信息而忽视未共享信息"
        + citations.cite("hidden_profile_meta")
        + "。HiddenBench 将该范式迁移到 LLM 多智能体，并公开了分布式信息条件下的集体推理失败"
        + citations.cite("hiddenbench")
        + "。",
    )
    _body(
        document,
        "本文因此把研究主线从“全披露是否更准”调整为“如何保证多智能体最终结果正确”。全局 Reveal-All 仅作为信息可得性的上限对照，用来判断任务错误是否仍可能由信息缺失解释；它不是本文主要贡献。本文关注私有事实从产生答案必须经过的公开、读取、解释和行动四个环节，并把错误现象—观测指标—针对性干预连接成治理闭环。",
    )
    _body(document, "RQ1：分布式私有信息任务中会出现哪些可观测的错误现象？", bold_lead="RQ1：")
    _body(document, "RQ2：哪些信息处理环节导致最终答案错误，分别可以用哪些指标捕捉？", bold_lead="RQ2：")
    _body(document, "RQ3：针对不同错误环节的干预能否改善最终正确率？", bold_lead="RQ3：")
    _body(document, "RQ4：更换执行框架和大模型后，原论文结论是否仍成立，数值差异可能来自哪里？", bold_lead="RQ4：")


def _error_phenomena_and_prior_work(
    document: DocumentType, citations: CitationTracker
) -> None:
    _heading(document, "2 多智能体错误现象与已有研究")
    _body(
        document,
        "本章按照错误发生链路组织相关研究：先定义可见的结果异常，再追溯信息公开、利用、解释和协调环节。MAST 从多种框架与任务中归纳了 14 种失败模式，为系统设计、智能体间错位以及验证终止问题提供了诊断框架"
        + citations.cite("mast")
        + "。本文使用该分类作有限映射，但只有消息级证据充分时才确认某一失败模式。",
    )

    _heading(document, "2.1 结果层：错误答案与错误共识", level=2)
    _body(
        document,
        "最直观的错误是最终多数票或全体投票选择错误答案。更危险的情况是错误共识：所有 Agent 给出一致判断，但一致性来自从众、早期立场锁定或共同误解，而不是充分证据。社会影响可能降低意见多样性，却不一定降低群体误差"
        + citations.cite("social_influence")
        + "；反共识辩论研究也指出，保留完整论证轨迹比只观察最后多数票更有信息"
        + citations.cite("free_mad")
        + "。",
    )

    _heading(document, "2.2 公开层：私有事实未披露或披露过晚", level=2)
    _body(
        document,
        "公开失败指决定性私有事实没有进入公共消息，或在立场已经锁定后才出现。它可以通过原子披露率、事实所有者是否发言、首次披露轮次和投票前证据覆盖率捕捉。该问题对应 Hidden Profile 的未共享信息偏差，也是 HiddenBench 任务的基本风险来源。",
    )

    _heading(document, "2.3 利用与解释层：说出来仍可能用错", level=2)
    _body(
        document,
        "事实出现在公共消息中不代表其他 Agent 已经读取并用于推理。利用失败表现为后续理由不引用关键事实、收到证据后立场不变或投票仍依赖旧信息；解释失败则表现为事实含义被反转、忽略否定关系，或把证据错误绑定到候选项。思维退化研究说明，高置信立场形成后，后续反思未必产生新的推理路径"
        + citations.cite("divergent_debate")
        + "；自动责任归因研究也显示，即使保留完整轨迹，定位责任 Agent 仍然困难"
        + citations.cite("failure_attribution")
        + "。",
    )

    _heading(document, "2.4 协调层：发言机会、过早收敛与虚假协同", level=2)
    _body(
        document,
        "协调失败包括发言机会失衡、异议过早消失以及时间上同步但信息上没有互补。AutoGen 等框架提供集中式选择器组织群聊"
        + citations.cite("autogen")
        + "，DyLAN 则动态优化 Agent 网络"
        + citations.cite("dylan")
        + "；偏信息分解研究进一步区分真实协同与表面时间耦合"
        + citations.cite("emergent_coordination")
        + "。这些工作说明发言结构可以设计，但仍需用最终正确率判断治理是否有效。",
    )

    _heading(document, "2.5 多智能体执行框架", level=2)
    _body(
        document,
        "MetaGPT、ChatDev 和 AgentVerse 分别用标准作业流程、聊天链和任务环境固化多智能体协作"
        + citations.cite("metagpt")
        + citations.cite("chatdev")
        + citations.cite("agentverse")
        + "。Microsoft Agent Framework 面向 Python 与 .NET 提供 Agent 和模型执行能力"
        + citations.cite("maf")
        + "。本文只使用 MAF 的执行层；讨论协议、选择器、Reveal、投票和审计均为项目代码。",
    )


def _governance_system(
    document: DocumentType, evidence: dict[str, Any], figure: Path
) -> None:
    _heading(document, "3 多智能体正确性治理体系")
    _body(
        document,
        "本文治理过程依次为：锁定任务与私有事实分配，执行多智能体讨论，对公共消息做原子事实审计，检查证据是否被利用和正确解释，检测投票与错误共识，最后根据错误所在环节选择干预并重新评估正确率。治理不把所有错误都交给同一种提示词或调度器处理。",
    )
    _figure(document, figure, "图 1  多智能体错误链路、观测指标与针对性干预闭环")

    _heading(document, "3.1 错误现象、观测点与干预矩阵", level=2)
    rows = [
        ["公开", "事实未出现或出现过晚", "原子披露率；首次披露轮次；所有者披露情况", "结构化交换；缺失事实提醒；全局 Reveal-All"],
        ["读取/利用", "事实出现但未影响理由或投票", "后续证据引用；立场变化；投票前证据覆盖", "要求引用证据；提示未处理事实"],
        ["解释", "事实含义或事实-选项关系理解错误", "理由-事实一致性；候选项-证据映射", "反思；冲突核验；重新解释"],
        ["协调", "过早收敛、从众或发言机会失衡", "发言分布；共识形成轮次；异议消失时间", "动态发言；保留异议；延迟表决"],
        ["决策", "理由与投票不一致或错误共识", "个体票；多数答案；理由-票一致性；错误共识率", "异常投票复审；二次投票"],
    ]
    _table_caption(document, "表 1  错误链路的可观测点与对应干预")
    _table(document, ["错误环节", "可观测现象", "主要观测指标", "对应干预"], rows, [1100, 2300, 3000, 2960], font_size=7.1)

    _heading(document, "3.2 结果指标与诊断指标", level=2)
    _body(
        document,
        "最终正确率和多数正确率用于判断系统是否真正完成任务；错误共识率用于识别“稳定但错误”的输出。原子披露率 D=m/n 只回答多少条私有事实通过证据核验进入了公共讨论，不能直接代替正确率。后续引用、立场变化、理由-证据一致性和理由-投票一致性用于进一步判断公开信息是否被正确使用。",
    )
    _body(
        document,
        "本项目已实现原子事实披露审计、多数投票、错误共识率、同预算固定/动态对照和运行级 MAST 映射；后续证据引用与解释一致性的全样本人工双盲编码尚未完成，因此本文只对有明确消息证据的案例作有限判断。",
    )


def _contribution_boundaries(document: DocumentType) -> None:
    _heading(document, "4 已有工作、本文改进与原创边界")
    rows = [
        ["已有研究", "HiddenBench 提供任务、官方 GPT-4.1 基线和 Reveal-All 结果；MAST 提供失败类型；MAF 提供 Agent/模型执行能力。"],
        ["本文实现与改进", "在 MAF 上复现 HiddenBench；实现原子事实审计；比较固定、结构化和同预算动态发言；逐运行定位错误链路；形成观测-干预闭环。"],
        ["非本文原创", "Hidden Profile 范式、Reveal-All 思路、MAST 分类以及 Microsoft Agent Framework 本身。"],
    ]
    _table_caption(document, "表 2  已有工作、本文实现与原创性边界")
    _table(document, ["类别", "内容"], rows, [1800, 7560], font_size=8.2)
    for label, text in rows:
        _body(document, f"{label}：{text}", bold_lead=f"{label}：")
    _body(
        document,
        "五因子动态选择器是项目代码中的 external centralized scheduler，可以读取完整私有事实分配，并非 MAF 原生群聊能力。因此它目前是机制研究工具，不是可直接部署的去中心化治理器。Codex 对 144 项样本的独立盲审属于模型辅助复核，不是人工金标准。",
    )


def _experiment_design(document: DocumentType, evidence: dict[str, Any]) -> None:
    _heading(document, "5 实验设置与对标方法")
    _heading(document, "5.1 三道案例的选择依据", level=2)
    _body(
        document,
        "选择 ID1、ID2 和 ID3，是因为官方公开了逐题 GPT-4.1 基线和 Reveal-All 结果，可以直接和本地复现逐题对表。三题共享同一背景和候选项 West City、North Hill 与 East Town，只改变决定正确答案的私有事实链，因此能够在背景基本固定时观察错误位置差异。本文不把三题结论外推到全部 65 题。",
    )
    task_rows = [
        [f"ID{task['task_id']}", task["correct_answer"], "4 条公共信息 + 4 条私有事实", task["evidence_chain"]]
        for task in evidence["task_details"]
    ]
    _table_caption(document, "表 3  三道短任务及决定性证据链")
    _table(document, ["任务", "正确答案", "信息结构", "决定性证据链"], task_rows, [700, 1200, 2000, 5460], font_size=7.0)

    _heading(document, "5.2 对标层次与实验控制", level=2)
    _body(
        document,
        "第一层是跨实现复现：官方 GPT-4.1 与作者自写模拟器，对比本地 DeepSeek 与 MAF，只判断结论方向是否仍然存在。第二层是本地机制比较：保持任务、模型、Prompt、事实分配和发言预算一致，再比较讨论协议，避免把模型、框架或预算差异误认为治理效果。pair seed 只控制事实归属，generation seed 未传给 DeepSeek，因此不能保证逐字复现。",
    )
    scale_rows = [
        ["本地运行", f"{evidence['totals']['local_runs']} 次", "各阶段冻结证据合计"],
        ["讨论与投票请求", f"{evidence['totals']['discussion_vote_requests']:,} 次", "DeepSeek Agent 发言与最终投票"],
        ["披露审计请求", f"{evidence['totals']['audit_requests']} 次", "AI 原子事实判定"],
        ["模型请求总数", f"{evidence['totals']['model_requests']:,} 次", "讨论/投票与审计合计"],
    ]
    _table_caption(document, "表 4  冻结证据规模")
    _table(document, ["项目", "规模", "说明"], scale_rows, [2500, 1700, 5160], font_size=8.0)


def _results_and_attribution(document: DocumentType, evidence: dict[str, Any]) -> None:
    _heading(document, "6 实验结果与归因分析")
    _heading(document, "6.1 原论文与本地复现是否得到相同结论", level=2)
    comparison_rows = [
        ["ID1", "11/30", "10/10", "10/10", "10/10"],
        ["ID2", "4/30", "0/10", "9/10", "10/10"],
        ["ID3", "6/30", "7/10", "10/10", "10/10"],
    ]
    _table_caption(document, "表 5  官方 GPT-4.1 与本地 DeepSeek 的逐题结果")
    _table(
        document,
        ["任务", "官方基线", "本地 fixed-60", "官方 Reveal-All", "本地官方兼容 Reveal-All"],
        comparison_rows,
        [700, 1500, 1700, 1700, 3760],
        font_size=7.5,
    )
    _body(
        document,
        "官方部分信息基线 ID1、ID2、ID3 分别为 11/30、4/30、6/30；本地 fixed-60 分别为 10/10、0/10、7/10。两边方向上都观察到分布式信息条件下的失败，且全局 Reveal-All 后接近或达到满分，但具体数值并不一致。ID2 在本地更差，ID1 和 ID3 则更好。",
    )
    _body(
        document,
        "可能原因包括 DeepSeek 与 GPT-4.1 的模型差异、MAF 与作者模拟器的消息组织差异、讨论协议和预算差异，以及未受控的生成随机性。由于模型与框架在跨实现比较中同时改变，现有实验不能把差异单独归因给模型或框架；上述因素只是候选解释，不是已经建立的因果结论。",
    )

    _heading(document, "6.2 信息完整性之外：披露时序与正确使用", level=2)
    timing_rows = [
        ["官方兼容第一轮全局 Reveal-All", "第一轮每条响应注入全部四条事实", "30/30", "信息可得性上限"],
        ["逐步 fixed-reveal-all", "事实随所有者轮流进入讨论", "26/30", "ID2 为 6/10，最终披露率仍为 100%"],
    ]
    _table_caption(document, "表 6  两种全披露时序的结果差异")
    _table(document, ["条件", "注入方式", "总体正确", "解释"], timing_rows, [2700, 3100, 1300, 2260], font_size=7.6)
    _body(
        document,
        "第一轮直接提供全景信息时为 30/30；逐步披露虽然最终同样达到 100% 原子披露率，总体却只有 26/30。这说明“最后都说出来”不等于关键事实被及时读取、正确解释并落实到投票。全局 Reveal-All 因此只用于诊断信息时序和使用链路，而不是证明一个显然的“信息越多越好”。",
    )

    _heading(document, "6.3 针对性干预是否改善最终正确率", level=2)
    intervention_rows = [
        ["发言机会", "fixed-60", "dynamic-60", "17/30", "21/30", "p=0.125", "未达到统计显著"],
        ["信息交换结构", "fixed-12", "structured-12", "20/30", "24/30", "p=0.388", "未达到统计显著"],
        ["披露时序", "fixed-reveal-all", "dynamic-reveal-all", "26/30", "30/30", "p=0.125", "未达到统计显著"],
    ]
    _table_caption(document, "表 7  本地干预条件的配对比较")
    _table(
        document,
        ["干预对象", "条件 A", "条件 B", "A 正确", "B 正确", "精确检验", "结论"],
        intervention_rows,
        [1350, 1500, 1700, 1000, 1000, 1200, 1610],
        font_size=6.9,
    )
    _body(
        document,
        "动态发言和结构化讨论的点估计均高于对应固定条件，但 p=0.125 和 p=0.388 均大于 0.05。当前证据只能说明这些机制值得继续检验，没有证据证明动态调度或结构化协议具有统计显著优势。",
    )

    _heading(document, "6.4 错误链路案例与 MAST 映射", level=2)
    mast_rows = [
        ["FM-2.4 已确认", "383ab32a67ba98e7 / fixed-60 / ID1", "至少一条所有者事实没有进入公共消息；该次仍答对，说明这是风险机制而非必然失败。"],
        ["FM-2.5 待人工因果编码", "e46521e6691a10ca / fixed-reveal-all / ID2", "四条事实均公开，但 Agent 一致错误解释 East Town 隧道状态；需要人工复核因果链。"],
        ["FM-2.6 未确认", "无可靠运行", "没有找到理由与最终行动直接冲突的充分证据。"],
    ]
    _table_caption(document, "表 8  本地运行中的失败模式证据边界")
    _table(document, ["模式判断", "证据位置", "解释"], mast_rows, [2000, 2700, 4660], font_size=7.4)
    _body(
        document,
        "现有案例支持错误可以发生在公开和解释环节，但尚不足以统计分解各环节对错误率的独立贡献。尤其是“信息出现但被忽略”与“信息被错误解释”需要人类审阅者依据完整上下文编码，不能只靠关键词或披露百分比判断。",
    )

    _heading(document, "6.5 对四个研究问题的回答", level=2)
    _body(document, "RQ1：可观测错误包括未披露、披露过晚、证据未被利用、事实被误解、过早收敛、理由与投票不一致以及错误共识。", bold_lead="RQ1：")
    _body(document, "RQ2：原子披露率、首次轮次、后续引用、立场变化、理由-证据一致性、发言分布和投票一致性可以分别定位公开、利用、解释、协调和决策问题。", bold_lead="RQ2：")
    _body(document, "RQ3：全局信息可得性明显改善三题表现；动态和结构化干预有方向性改善，但当前样本没有统计显著证据。", bold_lead="RQ3：")
    _body(document, "RQ4：原论文“分布式信息会导致失败、Reveal-All 明显改善”的方向在本地仍存在，但逐题数值不同，且现有设计不能完成单因素归因。", bold_lead="RQ4：")


def _conclusion(document: DocumentType) -> None:
    _heading(document, "7 总结与展望")
    _body(
        document,
        "本文把多智能体系统的首要目标定义为最终决策正确，并建立了从私有事实公开、读取利用、解释协调到投票决策的错误链路。各环节分别对应可审计观测点和干预方式，使治理从“统一增加讨论轮数”转向“先定位、再干预、后验证”。",
    )
    _body(
        document,
        "实验表明，完整性、披露时序和正确使用共同影响结果。第一轮全局 Reveal-All 为 30/30，而逐步 fixed-reveal-all 为 26/30；动态发言和结构化讨论的点估计更高，但配对检验未达到统计显著。因此本文不能宣称某个调度器已经稳定解决问题，只能确认单看最终披露率或共识强度是不够的。",
    )
    _body(
        document,
        "下一步最重要的是做解耦实验：在 MAF 中使用 GPT-4.1，并让作者模拟器运行 DeepSeek，以分别估计模型和框架效应；同时扩展到更多具有官方结果的任务，由两名不知道 AI 标签的人类审阅者完成利用与解释环节的双盲编码。只有在更多任务、模型和人工证据上重复出现，治理机制才能从实验性工具变成可部署方案。",
    )
    _callout(
        document,
        "最终结论：",
        "正确性治理不是让所有 Agent 更快达成一致，而是保证决定性证据及时公开、被正确理解，并真正改变最终行动。",
        gold=True,
    )


def _references(
    document: DocumentType, citations: CitationTracker
) -> None:
    _heading(document, "参考文献")
    for index, item in enumerate(citations.ordered_references(), start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.28)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.keep_together = True
        _font(paragraph.add_run(_reference_entry(index, item)), 9.2)


def _appendices(document: DocumentType, evidence: dict[str, Any]) -> None:
    _page_break(document)
    _heading(document, "附录")
    _heading(document, "附录 A  三题选择与决定性证据链", level=2)
    _body(
        document,
        "ID1-ID3 是官方结果中可直接获得逐题 GPT-4.1 基线与 Reveal-All 对照的三个短任务。三题使用同一撤离背景和候选项，但由不同私有事实链决定唯一可行路线。",
    )
    rows = [
        [f"ID{task['task_id']}", task["name"], task["correct_answer"], task["evidence_chain"]]
        for task in evidence["task_details"]
    ]
    _table(document, ["任务", "场景", "正确答案", "证据链"], rows, [700, 2000, 1200, 5460], font_size=7.2)

    _heading(document, "附录 B  实验 Prompt 与信息注入模板", level=2)
    _code(document, "B.1  讨论 Prompt", DISCUSSION_PROMPT)
    _code(document, "B.2  最终投票 Prompt", VOTE_PROMPT)
    _code(document, "B.3  原子事实披露审计 Prompt", ATOMIC_AUDIT_PROMPT)
    _code(document, "B.4  官方兼容全局 Reveal-All 模板", OFFICIAL_GLOBAL_REVEAL)

    _heading(document, "附录 C  代码、数据与复现位置", level=2)
    publication = evidence["publication"]
    official = evidence["official_code"]
    source_rows = [
        ["HiddenBench 论文", official["paper"]],
        ["HiddenBench 官方代码", official["repository"]],
        ["官方代码 commit", official["verified_commit"]],
        ["官方逐题结果", official["results"]],
        ["本项目仓库", publication["repository"]],
        ["实验分支", publication["branch"]],
        ["Pull Request", publication["pull_request"]],
        ["证据 Release", publication["release"]],
    ]
    _table(document, ["项目", "位置"], source_rows, [2400, 6960], font_size=7.2)
    _body(
        document,
        "冻结证据文件的 SHA-256 已保存在 reports/data/hiddenbench-paper-evidence-20260805.json；正文所有数量和比例均从该快照读取。",
    )

    _heading(document, "附录 D  术语与结论边界", level=2)
    term_rows = [
        ["MAF", "Microsoft Agent Framework；本文只使用 Agent/模型执行能力。"],
        ["披露率 D", "通过证据核验的已披露原子事实数 m 除以事实总数 n。"],
        ["错误共识", "多数或全体 Agent 对错误选项形成一致判断。"],
        ["全局 Reveal-All", "第一轮每条响应都获得全部私有事实，用作信息可得性上限。"],
        ["逐步 reveal-all", "事实随所有者发言逐条进入讨论，最终完整但时序不同。"],
        ["模型辅助复核", "Codex 独立标签复核，不是人工金标准。"],
        ["统计不显著", "当前样本不足以拒绝成功概率相同的零假设，不代表两条件完全相同。"],
    ]
    _table(document, ["术语", "本文含义"], term_rows, [2500, 6860], font_size=7.8)


def _build_marked_document(
    root: Path,
    evidence: dict[str, Any],
    references: list[dict[str, Any]],
    marked_output: Path,
) -> CitationTracker:
    figure = root / GOVERNANCE_FIGURE
    if not figure.exists():
        raise FileNotFoundError(figure)

    citations = CitationTracker(references)
    document = Document()
    _configure_paper(document)
    _configure_identity(document)
    _title_and_abstract(document, evidence)
    _introduction(document, citations)
    _error_phenomena_and_prior_work(document, citations)
    _governance_system(document, evidence, figure)
    _contribution_boundaries(document)
    _experiment_design(document, evidence)
    _results_and_attribution(document, evidence)
    _conclusion(document)
    _references(document, citations)
    _appendices(document, evidence)
    _validate_tables(document)
    marked_output.parent.mkdir(parents=True, exist_ok=True)
    document.save(marked_output)
    return citations


def build_correctness_governance_report(
    root: Path, output: Path, desktop_output: Path | None
) -> Path:
    root = root.resolve()
    output = output if output.is_absolute() else root / output
    desktop_output = (
        desktop_output
        if desktop_output is None or desktop_output.is_absolute()
        else root / desktop_output
    )
    _assert_preserved(root)
    evidence = _load_json(root / EVIDENCE_PATH)
    references: list[dict[str, Any]] = _load_json(root / REFERENCE_PATH)

    with tempfile.TemporaryDirectory(prefix="maf-correctness-paper-") as temp_dir:
        marked = Path(temp_dir) / "marked.docx"
        citations = _build_marked_document(root, evidence, references, marked)
        ordered = citations.ordered_references()
        if len(ordered) < 15:
            raise ValueError(f"expected at least 15 cited references, got {len(ordered)}")
        note_texts = [_short_note(item) for item in ordered]
        _patch_true_footnotes(marked, output, note_texts)

    if desktop_output is not None:
        desktop_output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(output, desktop_output)
    _assert_preserved(root)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build the MAS correctness-governance paper with true Word footnotes."
    )
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, default=Path("reports") / REPORT_NAME)
    parser.add_argument("--desktop-output", type=Path)
    args = parser.parse_args()
    print(
        build_correctness_governance_report(
            args.root, args.output, args.desktop_output
        )
    )


if __name__ == "__main__":
    main()
