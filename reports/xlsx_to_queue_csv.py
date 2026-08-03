"""Convert a completed sign-off Excel sheet back into the queue CSV format.

Usage:
    python reports/xlsx_to_queue_csv.py <filled.xlsx> <output.csv>

The output CSV carries the same columns as the blinded queue, with the
human_* columns filled, ready for:
    python reports/run_hiddenbench_blind_review.py --from-human-queue
(after copying it to artifacts/hiddenbench-stability-20260729.blind-review.queue.csv)
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
QUEUE_CSV = (
    ROOT
    / "artifacts"
    / "hiddenbench-stability-20260729.blind-review.queue.csv"
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("output", type=Path, nargs="?", default=None)
    parser.add_argument(
        "--reviewer-id",
        default="",
        help="Reviewer identifier recorded in human_reviewer_id.",
    )
    args = parser.parse_args()

    workbook = load_workbook(args.excel, data_only=True)
    sheet = workbook["盲审签核"]
    rows = list(csv.DictReader(QUEUE_CSV.open(encoding="utf-8-sig")))
    by_id = {row["blind_id"]: row for row in rows}

    filled = 0
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        if excel_row[1] is None:
            continue
        blind_id = str(excel_row[1]).strip()
        if blind_id not in by_id:
            raise ValueError(f"unknown blind ID in Excel: {blind_id}")
        raw = str(excel_row[5] or "").strip()
        if not raw:
            continue
        disclosed = raw in {"是", "true", "1", "yes", "y"}
        if raw not in {"是", "否", "true", "false", "1", "0", "yes", "no", "y", "n"}:
            raise ValueError(f"{blind_id}: cannot parse disclosure '{raw}'")
        row = by_id[blind_id]
        row["human_disclosed"] = "true" if disclosed else "false"
        row["human_evidence_message_ids"] = str(
            excel_row[6] or ""
        ).strip().replace("|", "|")
        row["human_evidence_quote"] = str(excel_row[7] or "").strip()
        reason = str(excel_row[8] or "").strip()
        row["human_reason"] = reason
        row["human_reviewer_id"] = args.reviewer_id
        filled += 1

    output = args.output or (ROOT / "artifacts" / "signoff-from-excel.queue.csv")
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    print(f"filled {filled}/{len(rows)} rows -> {output}")


if __name__ == "__main__":
    main()
