from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "reports"
    / "给彭老师的HiddenBench_MAF复现实验最终报告_2026-07-28.docx"
)
OUTPUT = (
    ROOT
    / "reports"
    / "给彭老师的HiddenBench_MAF复现实验最终报告_2026-07-30_修订版.docx"
)
DESKTOP_OUTPUT = (
    Path.home()
    / "Desktop"
    / "给彭老师的HiddenBench_MAF复现实验最终报告_2026-07-30_修订版.docx"
)
OFFICIAL_COMPARISON = (
    ROOT / "reports" / "HiddenBench_官方GPT4.1四题对照_2026-07-30.csv"
)
STABILITY_SUMMARY = (
    ROOT / "artifacts" / "hiddenbench-stability-20260729.summary.csv"
)
DISAGREEMENTS = (
    ROOT / "artifacts" / "hiddenbench-stability-20260729.disagreements.csv"
)
MANIFEST = (
    ROOT / "artifacts" / "hiddenbench-stability-20260729.manifest.json"
)

BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EDF4FA"
GRAY = "F2F2F2"
WHITE = "FFFFFF"


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_width(cell, width_cm: float) -> None:
    width_twips = int(Cm(width_cm).twips)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_twips))


def _set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(Cm(sum(widths_cm)).twips)))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_cm in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(Cm(width_cm).twips)))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width_cm in zip(row.cells, widths_cm, strict=True):
            _set_cell_width(cell, width_cm)


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _format_table(table, widths_cm: list[float]) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_widths(table, widths_cm)
    _repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        _prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(
                        qn("w:eastAsia"), "Microsoft YaHei"
                    )
                    run.font.size = Pt(8.5)
        if row_index == 0:
            for cell in row.cells:
                _set_cell_shading(cell, BLUE)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
        elif row_index % 2 == 0:
            for cell in row.cells:
                _set_cell_shading(cell, PALE_BLUE)


def _add_table(
    document: DocumentType,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    widths_cm: list[float],
):
    materialized = [list(row) for row in rows]
    table = document.add_table(rows=1 + len(materialized), cols=len(headers))
    for index, value in enumerate(headers):
        table.cell(0, index).text = value
    for row_index, values in enumerate(materialized, start=1):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = str(value)
    _format_table(table, widths_cm)
    document.add_paragraph()
    return table


def _load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _set_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _replace_existing_text(document: DocumentType) -> None:
    replacements = {
        "复现实验最终报告": "复现实验最终报告（2026-07-30 修订版）",
        "每题只运行一个固定种子，所以这里的结果叫“十题单种子筛选”，不等同于 原论文的 65 题多 session 主实验。": (
            "第一阶段仍保留十题单种子筛选；第二阶段另行冻结 ID 1、5、7、25，"
            "在固定轮转与动态发言两种条件下各重复 10 次，共 80 个正式 run。"
        ),
        "每题只运行一个固定种子。": (
            "下表仍是第一阶段十题单种子筛选；后文另报四题十次重复稳定性实验，"
            "两套结果不混算。"
        ),
        "全套测试：149 项通过；": (
            "新增稳定性与官方对照模块后，全套测试 197 项通过；"
        ),
        "原论文每题有多次 session，本次每题只有一个固定种子；": (
            "第一阶段十题筛选每题一个固定种子；第二阶段四道代表题已扩展为"
            "固定轮转 10 次与动态发言 10 次。"
        ),
        "下一步不宜马上加入更多自定义参数。建议先做两件事：": (
            "在第一阶段基础上，下面两项后续工作已经完成："
        ),
        "对 ID 1、5、7 三个错误共识案例进行逐轮人工编码，判断是信息未披露、": (
            "已对 ID 1、5、7 的 180 条动态发言逐条编码，区分信息未披露、"
        ),
        "信息被忽略，还是正确证据被错误解释；": (
            "信息被忽略和证据被错误解释，并保留 message ID 与模型原话；"
        ),
        "冻结当前固定轮转基线后，再用相同任务、相同模型和相同 60 次发言": (
            "已冻结固定轮转基线，并以相同任务、相同模型、成对种子和相同"
            " 60 次发言"
        ),
        "预算比较内容感知动态发言机制，避免把预算差异误认为治理效果。": (
            "预算完成内容感知动态发言比较；选择器不调用 LLM，因此两种条件的"
            "生成预算严格相同。"
        ),
        "上述七个实验附件在本次提交中被明确纳入版本控制，不再只是本地 .gitignore 下的临时文件。老师从 GitHub 下载仓库即可复核。": (
            "第一阶段附件与第二阶段 80 个 run、4800 条发言、AI 披露审计、"
            "差异清单、汇总表、闸门和 manifest 均已纳入版本控制。"
        ),
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            _set_paragraph_text(paragraph, replacements[text])

    if len(document.tables) > 1:
        document.tables[0].cell(3, 1).text = "2026 年 7 月 30 日"
        document.tables[1].cell(0, 0).text = (
            "阶段结论：已完成十题单种子筛选，以及 ID 1、5、7、25 的"
            "固定轮转/动态发言各 10 次重复实验。动态机制在等额 60 次发言"
            "预算下未优于固定轮转；官方 GPT-4.1 逐题结果与本实验已按公开"
            "session 对齐，但因模型和实现不同，只作描述性比较。"
        )
        _set_cell_shading(document.tables[1].cell(0, 0), LIGHT_BLUE)


def _ensure_revision_styles(document: DocumentType) -> None:
    if "Revision Note" not in document.styles:
        style = document.styles.add_style(
            "Revision Note", WD_STYLE_TYPE.PARAGRAPH
        )
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(9.5)
        style.font.color.rgb = RGBColor(70, 70, 70)
        style.paragraph_format.left_indent = Cm(0.7)
        style.paragraph_format.right_indent = Cm(0.7)
        style.paragraph_format.space_after = Pt(6)


def _add_revision(document: DocumentType) -> None:
    official = {
        int(row["task_id"]): row for row in _load_csv(OFFICIAL_COMPARISON)
    }
    stability = _load_csv(STABILITY_SUMMARY)
    disagreements = _load_csv(DISAGREEMENTS)
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    manifest_lookup = {
        item["path"]: item for item in manifest["files"]
    }

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("十、2026-07-30 修订：老师问题逐项落实", level=1)
    note = document.add_paragraph(style="Revision Note")
    note.add_run(
        "本章不改写第一阶段原始实验，而是在其后加入官方逐题对照、"
        "AI 披露审计、四题十次重复、等预算动态机制及逐轮失败分析。"
        "因此，前文“十题单种子筛选”与本章“4 题×2 条件×10 次”"
        "是两套边界明确、互不混算的结果。"
    )

    document.add_heading("1. 原论文使用什么框架、模型，是否公开代码", level=2)
    document.add_paragraph(
        "HiddenBench 官方实现不是 Microsoft Agent Framework、AutoGen "
        "或其他现成群聊框架，而是论文作者自己编写的 Python 模拟器。"
        "公开仓库的核心讨论编排位于 src/hiddenbench/simulator.py，模型访问"
        "通过自定义 ModelClient 抽象完成。论文正文称评测 15 个模型；"
        "作者当前公开的 Figure 3 校验表实际列出 17 个模型配置，完整名单如下。"
    )
    _add_table(
        document,
        ["模型家族", "作者公开结果中的模型配置"],
        [
            [
                "OpenAI GPT（8）",
                "GPT-4o；GPT-4.1-Nano；GPT-4.1-Mini；GPT-4.1；"
                "GPT-5-Nano-Minimal；GPT-5-Mini-Minimal；"
                "GPT-5-Minimal；GPT-5-Medium",
            ],
            [
                "Google Gemini（3）",
                "Gemini-2.5-Flash-Lite；Gemini-2.5-Flash；Gemini-2.5-Pro",
            ],
            [
                "Alibaba Qwen（4）",
                "Qwen3-8B；Qwen3-14B；Qwen3-32B；Qwen3-235B-A22B",
            ],
            [
                "Meta Llama（2）",
                "Llama-4-Scout；Llama-4-Maverick（结果文件名 llama-4）",
            ],
        ],
        [3.4, 12.6],
    )
    document.add_paragraph(
        "因此，本报告不把“15”误写成一个不完整名单，而是保留论文正文口径，"
        "同时披露当前官方结果包的 17 项实际配置。本实验只选取其中 GPT-4.1，"
        "作为四道题的逐题外部参照。"
    )
    document.add_paragraph(
        "原论文公开了源代码、65 题数据、Prompt、评测脚本和结果包。"
        "代码仓库：https://github.com/Yassellee/HiddenBench_ICML；"
        "本次核验提交：3be6ca16973e4fb751ffc0dfb7eb11f2d28335d1。"
        "结果包：https://huggingface.co/datasets/"
        "YuxuanLi1225/HiddenBench-results。"
    )
    document.add_paragraph(
        "本实验使用 Microsoft Agent Framework + DeepSeek V4 Flash，"
        "属于“同任务、同信息结构和同轮数协议下的框架迁移”。它不是同模型"
        "严格复现，任何数值差异都不能只归因于 Microsoft 框架。"
    )

    document.add_heading("2. 没有跑 65 题，四题的官方结果是什么", level=2)
    document.add_paragraph(
        "论文正文主要报告总体结果，没有把 ID 1、5、7、25 的逐题数字逐行"
        "印在主表中；但作者公开的原始结果包包含每题 session。本次直接读取"
        "官方 GPT-4.1 文件，按官方 metrics.py 重新计算，每题 Hidden 与"
        " Full Profile 均为 10 个 session。"
    )
    official_rows = []
    for task_id in (1, 5, 7, 25):
        row = official[task_id]
        official_rows.append(
            [
                str(task_id),
                row["scenario"],
                (
                    f"{_pct(float(row['official_hidden_pre_average_accuracy']))}"
                    " → "
                    f"{_pct(float(row['official_hidden_post_average_accuracy']))}"
                ),
                (
                    f"{_pct(float(row['official_hidden_pre_majority_accuracy']))}"
                    " → "
                    f"{_pct(float(row['official_hidden_post_majority_accuracy']))}"
                ),
                (
                    f"{_pct(float(row['official_full_pre_average_accuracy']))}"
                    " / "
                    f"{_pct(float(row['official_full_pre_majority_accuracy']))}"
                ),
                (
                    f"{_pct(float(row['deepseek_fixed_post_majority_accuracy']))}"
                    " / "
                    f"{_pct(float(row['deepseek_dynamic_post_majority_accuracy']))}"
                ),
            ]
        )
    _add_table(
        document,
        [
            "ID",
            "任务",
            "官方 Hidden\n平均正确率 前→后",
            "官方 Hidden\n多数正确率 前→后",
            "官方 Full\n平均/多数",
            "本实验 后多数\n固定/动态",
        ],
        official_rows,
        [0.7, 3.2, 3.1, 3.1, 2.6, 3.3],
    )
    document.add_paragraph(
        "最稳定的复现现象是 ID 5：官方 GPT-4.1 的讨论后多数正确率为"
        " 0%，本实验固定轮转和动态机制也都是 0%，完整信息条件却是"
        " 100%。这说明失败主要与分布式信息整合有关。ID 1、7、25 的"
        "数值则随模型和编排而明显变化：固定轮转分别为 100%、20%、80%，"
        "动态机制为 60%、0%、80%，官方 GPT-4.1 为 50%、50%、50%。"
    )
    document.add_paragraph(
        "四题平均的官方 GPT-4.1 Hidden 讨论后多数正确率为 37.5%；"
        "本实验固定轮转为 50.0%，动态机制为 35.0%。这只是四道代表题的"
        "描述性对照，不代表 65 题总体，也不做显著性推断。"
    )

    document.add_heading("3. 披露法则不是随机，AI 如何计算百分比", level=2)
    document.add_paragraph(
        "私有信息分配由基准题目固定：四条 private_information 按成对种子"
        "一一分给四名 Agent，不在讨论过程中随机追加。固定轮转始终按"
        " A-B-C-D；动态机制也不是随机抽人，而是用五项闭式分数选择下一位"
        "发言者：分歧 0.30、尚未披露 0.30、与当前讨论相关 0.15、需要回应"
        " 0.15、等待时间 0.10。选择器不调用 LLM，两种条件都是 60 次"
        "正式发言。"
    )
    document.add_paragraph(
        "AI 披露审计按“每个 run 的每条私有信息”逐项判断。审计模型同时"
        "看到该条私有信息和完整 60 条公开消息，只能输出：是否披露、证据"
        "所在 message ID、原文引句和简短理由。程序随后检查引句确实存在于"
        "对应消息，证据错误时最多修复一次。单个 run 的 AI 披露率 = 被判定"
        "已披露的私有信息条数 ÷ 4；报告百分比是 10 次 run 的均值。"
    )
    document.add_paragraph(
        "80 个 run 共完成 80 份有效审计，实际发送 84 次审计请求，其中"
        " 4 次用于结构或证据修复。AI 与透明规则在 320 个“run×私有信息”"
        "判断上有 84 次分歧，一致率 73.75%。55 次分歧集中于 ID 7，原因是"
        "该题一条“私有信息”本身包含多项 RFP 条款：规则只要命中部分原子"
        "条款就会计为披露，而 AI 更倾向要求整条信息的决策含义被完整表达。"
        "所以两种口径并列报告，不能挑一个对结论更有利的数。"
    )

    document.add_heading("4. 四题各做十次后，结果稳不稳定", level=2)
    stability_rows = []
    for row in stability:
        condition = "固定" if row["condition"] == "fixed" else "动态"
        repetitions = int(row["repetitions"])
        stability_rows.append(
            [
                f"{row['task_id']}/{condition}",
                (
                    f"{int(row['post_majority_correct_count'])}/"
                    f"{repetitions}"
                ),
                f"{int(row['wrong_consensus_count'])}/{repetitions}",
                _pct(float(row["ai_disclosure_mean"])),
                f"{int(row['stable_consensus_count'])}/{repetitions}",
                (
                    f"{float(row['mean_first_stable_consensus_turn']):.1f}"
                    if row["mean_first_stable_consensus_turn"]
                    else "—"
                ),
            ]
        )
    _add_table(
        document,
        [
            "ID/机制",
            "多数正确",
            "错误共识",
            "AI 披露率",
            "稳定共识",
            "首次稳定共识\n平均发言序号",
        ],
        stability_rows,
        [2.0, 2.2, 2.2, 2.4, 2.2, 5.0],
    )
    document.add_paragraph(
        "等额预算下，动态机制没有优于固定轮转：40 个固定 run 有 20 个多数"
        "正确、20 个错误共识；40 个动态 run 有 14 个多数正确、24 个错误"
        "共识，另有个别非全体一致结局。ID 1 的动态正确率从固定的 100%"
        "下降到 60%，ID 7 从 20%下降到 0%；ID 25 两者同为 80%。因此"
        "当前证据不支持“内容感知动态发言必然改善治理”，反而提示发言顺序"
        "会改变哪些证据先成为群体锚点。"
    )
    document.add_paragraph(
        "披露也不是充分条件。ID 1 的 AI 披露率较高（固定 90.0%、动态"
        " 85.0%），动态条件仍出现 4/10 错误共识；ID 5 披露接近于零，"
        "两种机制均 10/10 错误共识。ID 7 的 AI 口径为 0%，规则口径却"
        "较高，说明该题更需要按复合条款逐项人工核验，而不能用一个百分比"
        "替代机制分析。"
    )

    document.add_heading("5. 已经达成一致，为什么仍继续到 15 轮", level=2)
    document.add_paragraph(
        "原因首先是实验可比性：HiddenBench 官方 Hidden 结果按 15 轮运行，"
        "本实验也固定为每题 15 轮、每轮 4 条。若某些 run 一致后提前结束，"
        "不同机制就会得到不同 token 和发言预算，无法判断差异来自治理策略"
        "还是预算。其次，早期一致不等于稳定，也不等于正确；继续观察可以"
        "测量共识翻转、错误锁定和重复确认。"
    )
    document.add_paragraph(
        "本次已补充“首次共识”“首次稳定共识”“共识翻转次数”“稳定后仍"
        "继续的消息数”。例如 ID 5 固定轮转平均在第 7.2 条发言形成稳定"
        "共识，之后平均仍有 48.8 条消息，但 10 次全部是错误共识；这不是"
        "有效推理的证据，而是错误锚点被反复确认。未来可以把“达到稳定共识"
        "即停止”作为新的治理干预单独实验，但不能事后截断本轮基线。"
    )

    document.add_heading("6. MAST 失败模式在本实验哪里出现", level=2)
    document.add_paragraph(
        "下表来自此前 ID 1、5、7 动态单次 run 的 180 条逐发言人工编码。"
        "它给出可定位的 message ID 和模型原话，不把自动关键词命中直接"
        "当成 MAST 金标签。由于目前只有一名编码者，以下写为"
        "“MAST 对应候选”，不宣称已完成正式双人标注。"
    )
    _add_table(
        document,
        ["ID", "MAST 对应候选", "定位证据", "机制判断"],
        [
            [
                "1",
                "FM-2.6 式证据—结论错位（短暂）",
                "5243372b5c7b360d",
                "已说出补给车堵住东城隧道，却将其解释成“有风险但仍可通行”并建议东城；后续被纠正，因此不是最终锁死。",
            ],
            [
                "5",
                "FM-2.4 + FM-2.6 候选",
                "8ff56ce5223ba2c9；27244f00c26241b6",
                "关键候选人事实没有进入公共讨论；随后把“参与董事会/资本项目”反复改写成“募资成功”，错误共识持续。",
            ],
            [
                "7",
                "FM-2.5 + FM-2.6 候选",
                "f840375b2e2b177a → 06d313ed02d2f705",
                "agent-b 明确披露 Cape 在 b/f/g/h 失败并提出正确的 Starlight；下一条立即忽略该反证，后续群体持续复读 Cape。",
            ],
        ],
        [0.8, 3.1, 4.2, 7.9],
    )

    document.add_heading(
        "7. 每人只有一条私有信息，第一轮是否已经全景", level=2
    )
    document.add_paragraph(
        "“每人拥有一条”只描述 system prompt 的分配单位，不等于第一轮"
        "一定原样公开。Agent 可能只给建议、不提依据，可能只说复合信息中的"
        "一个子条款，也可能披露后被其他人忽略或错误解释。第一轮的可见消息"
        "因此只是“有四次发言机会”，不是自动 Reveal-All。真正的 Full "
        "Profile 条件是在投票前把全部四条私有信息直接放进每名 Agent 的"
        " system prompt；这与依赖 Agent 自主说出的 Hidden 讨论不同。"
    )
    document.add_paragraph(
        "正式重复结果也直接否定了“第一轮必然全景”：ID 5 的 AI 披露率"
        "固定仅 2.5%、动态为 0%，ID 7 的复合信息在 AI 口径下为 0%。"
        "即使某题第一轮已经披露全部关键信息，后续是否正确利用仍是另一个"
        "问题，ID 7 的人工证据就是“已披露但被忽略”的典型。"
    )

    document.add_heading("8. 审计边界与可复核附件", level=2)
    document.add_paragraph(
        "本次正式运行共 80 个 run、4800 条公开发言、5760 次正式模型请求；"
        "动态选择器 LLM 调用为 0。完整性闸门检查 80 个唯一 run、80 份"
        "唯一审计、成对种子与信息分配、每次 60 条发言、证据引句、凭据泄露"
        "和文件哈希，全部通过。AI 审计模型与生成模型同为 DeepSeek V4 "
        "Flash，可能存在同模型判断偏差，因此同时保留透明规则与全部 84 条"
        "分歧记录，供人工抽查。"
    )
    key_artifact = manifest_lookup["hiddenbench-stability-20260729.jsonl"]
    trace_artifact = manifest_lookup[
        "hiddenbench-stability-20260729.trace.jsonl"
    ]
    _add_table(
        document,
        ["附件", "规模", "SHA-256"],
        [
            [
                "hiddenbench-stability-20260729.jsonl",
                f"{key_artifact['bytes']:,} bytes",
                key_artifact["sha256"],
            ],
            [
                "hiddenbench-stability-20260729.trace.jsonl",
                f"{trace_artifact['bytes']:,} bytes",
                trace_artifact["sha256"],
            ],
            [
                "hiddenbench-stability-20260729.ai-disclosure.jsonl",
                f"{manifest_lookup['hiddenbench-stability-20260729.ai-disclosure.jsonl']['bytes']:,} bytes",
                manifest_lookup[
                    "hiddenbench-stability-20260729.ai-disclosure.jsonl"
                ]["sha256"],
            ],
            [
                "hiddenbench-stability-20260729.disagreements.csv",
                f"{len(disagreements)} 条 AI/规则分歧",
                manifest_lookup[
                    "hiddenbench-stability-20260729.disagreements.csv"
                ]["sha256"],
            ],
        ],
        [6.0, 3.0, 7.0],
    )
    document.add_paragraph(
        "官方逐题对照另见 "
        "reports/HiddenBench_官方GPT4.1四题对照_2026-07-30.csv；"
        "复现命令、配置、输出结构和校验步骤见 "
        "docs/hiddenbench-reproduction.md。"
    )


def _polish_new_content(document: DocumentType) -> None:
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.startswith("十、2026-07-30"):
            paragraph.paragraph_format.space_before = Pt(0)
    document.core_properties.title = (
        "HiddenBench 在 Microsoft Agent Framework 上的复现实验最终报告"
    )
    document.core_properties.subject = (
        "官方 GPT-4.1 逐题对照、AI 披露审计与十次重复稳定性修订"
    )
    document.core_properties.comments = (
        "2026-07-30 revised with official task-level results and stability study"
    )


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    for required in (
        OFFICIAL_COMPARISON,
        STABILITY_SUMMARY,
        DISAGREEMENTS,
        MANIFEST,
    ):
        if not required.is_file():
            raise FileNotFoundError(required)
    document = Document(SOURCE)
    _ensure_revision_styles(document)
    _replace_existing_text(document)
    _add_revision(document)
    _polish_new_content(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    DESKTOP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, DESKTOP_OUTPUT)
    print(OUTPUT)
    print(DESKTOP_OUTPUT)


if __name__ == "__main__":
    main()
