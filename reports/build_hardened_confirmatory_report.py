from __future__ import annotations

import hashlib
import json
import shutil
import sys
from pathlib import Path
from typing import Any

# Support both ``python -m reports.build_hardened_confirmatory_report`` and
# direct execution as ``python reports/build_hardened_confirmatory_report.py``.
if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from mas_experiment.hiddenbench_atomic_disclosure import (
    build_atomic_disclosure_prompt,
)
from mas_experiment.hiddenbench_confirmatory_study import ConfirmatoryRunRecord
from reports.build_hardened_confirmatory_summaries import wilson_interval


TABLE_WIDTH = 9360
TABLE_INDENT = 120
NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 98, 108)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"


CONDITION_LABELS = {
    "fixed-60": "固定轮转 60",
    "dynamic-60": "五因子顺序调度 60",
    "fixed-reveal-all": "固定 60 + 逐人机械披露",
    "dynamic-reveal-all": "动态 60 + 逐人机械披露",
    "structured-12": "Exchange→Decide 12",
    "fixed-12": "固定轮转 12",
    "single-local": "局部信息单智能体",
    "official-global-reveal": "官方同款全局 Reveal-All",
}


def _font(
    run,
    size: float = 11,
    *,
    bold: bool = False,
    color: RGBColor | None = None,
    mono: bool = False,
    italic: bool = False,
) -> None:
    western = "Consolas" if mono else "Calibri"
    east = "Microsoft YaHei"
    run.font.name = western
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), western)
    fonts.set(qn("w:hAnsi"), western)
    fonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def _cell_margins(cell) -> None:
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for edge, value in (
        ("top", 80),
        ("start", 120),
        ("bottom", 80),
        ("end", 120),
    ):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError("table widths must sum to 9360 DXA")
    table.autofit = False
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        props.append(width)
    width.set(qn("w:w"), str(TABLE_WIDTH))
    width.set(qn("w:type"), "dxa")
    indent = props.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT))
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            _cell_margins(cell)


def _table(
    doc: DocumentType,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int],
    *,
    font_size: float = 8.5,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _shade(cell, LIGHT_GRAY)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.10
            for run in paragraph.runs:
                _font(run, font_size, bold=True, color=DARK)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.10
                for run in paragraph.runs:
                    _font(run, font_size)
    _geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _paragraph(
    doc: DocumentType,
    text: str,
    *,
    bold_lead: str | None = None,
    italic: bool = False,
    color: RGBColor | None = None,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _font(lead, bold=True, color=DARK)
        rest = paragraph.add_run(text[len(bold_lead) :])
        _font(rest, italic=italic, color=color)
    else:
        run = paragraph.add_run(text)
        _font(run, italic=italic, color=color)


def _bullet(doc: DocumentType, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    _font(run)


def _heading(doc: DocumentType, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        _font(
            run,
            16 if level == 1 else 13 if level == 2 else 12,
            bold=True,
            color=BLUE if level < 3 else DARK,
        )


def _callout(doc: DocumentType, lead: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    _shade(cell, CALLOUT)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.10
    first = paragraph.add_run(lead)
    _font(first, 10.5, bold=True, color=DARK)
    rest = paragraph.add_run(text)
    _font(rest, 10.5)
    _geometry(table, [9360])
    doc.add_paragraph()


def _code(doc: DocumentType, title: str, text: str) -> None:
    _heading(doc, title, level=3)
    for line in text.splitlines() or [""]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F6F8FA")
        paragraph._p.get_or_add_pPr().append(shading)
        run = paragraph.add_run(line or " ")
        _font(run, 7.3, mono=True)


def _page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))
    _font(run, 8.5, color=GRAY)


def _configure(doc: DocumentType) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    bullet = doc.styles["List Bullet"]
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(8)
    bullet.paragraph_format.line_spacing = 1.167
    header = section.header.paragraphs[0]
    header.text = "HiddenBench × MAF｜确认性完善实验报告"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        _font(run, 8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("MAS 知识治理项目｜第 ")
    _font(label, 8.5, color=GRAY)
    _page_field(footer)
    suffix = footer.add_run(" 页")
    _font(suffix, 8.5, color=GRAY)


def _masthead(doc: DocumentType) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(14)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("TECHNICAL RESEARCH BRIEF")
    _font(run, 10, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    run = title.add_run("HiddenBench 确认性完善实验报告")
    _font(run, 23, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run(
        "官方同款 Reveal-All 补充复现、统计校正与 MAST 案例定位"
    )
    _font(run, 13.5, color=GRAY)
    metadata = (
        ("研究范围", "官方 short benchmark ID1/ID2/ID3"),
        ("模型与框架", "deepseek-v4-flash；Microsoft Agent Framework 调用层"),
        ("实验规模", "原确认性 210 次 + 官方同款补充 30 次"),
        ("日期", "2026 年 8 月 4 日"),
        ("状态", "全部数据门禁通过；结论按统计边界表述"),
    )
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        first = paragraph.add_run(f"{label}：")
        _font(first, 10.5, bold=True)
        second = paragraph.add_run(value)
        _font(second, 10.5)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(12)
    rule.paragraph_format.space_after = Pt(12)
    ppr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    ppr.append(borders)


def _pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _ci(successes: int, runs: int) -> str:
    low, high = wilson_interval(successes, runs)
    return f"{low * 100:.1f}%–{high * 100:.1f}%"


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def _find_confirmatory(
    rows: list[dict[str, Any]],
    study_key: dict[str, Any],
) -> dict[str, Any]:
    return next(row for row in rows if row["key"] == study_key)


def _teacher_answers(doc: DocumentType) -> None:
    _heading(doc, "1. 对老师问题的直接回答")
    _table(
        doc,
        ["老师可能追问", "本报告的准确回答"],
        [
            [
                "原论文用什么框架和模型？",
                "作者使用自定义 Python 模拟器；官方 short 结果模型为 GPT-4.1。作者公开了代码与逐轮结果。",
            ],
            [
                "我们是否真的使用了 MAF？",
                "MAF 负责 OpenAI-compatible Agent 创建与模型调用；轮转、调度、Reveal、投票、审计和门禁均为本项目显式代码。",
            ],
            [
                "原来的 Reveal-All 是否复现准确？",
                "旧条件是 owner-by-owner mechanical reveal，不等同官方协议。本轮新增 official-compatible global Reveal-All：第一轮每条响应均附加全部四条私有事实。",
            ],
            [
                "披露率数值在哪里？",
                "AI 对四个原子事实逐条输出 disclosed、消息 ID 与原文引文；规则程序验真后计算比例。正文给出完整 3/4=75% 实例。",
            ],
            [
                "为什么一致后还继续？",
                "影子投票不回灌讨论，用于同轨迹比较。唯一异常不是答案翻转，而是早期全体一致后来变成 2-2 平票。",
            ],
            [
                "多智能体是不是一定更好？",
                "不是。报告同时给出局部单智能体、完整信息单智能体、固定与动态多智能体，并保留所有失败结果。",
            ],
        ],
        [2200, 7160],
        font_size=9,
    )


def _protocol_section(doc: DocumentType) -> None:
    _heading(doc, "2. 协议边界与本次修正")
    _paragraph(
        doc,
        "原论文的 Reveal-All 在第一位参与者发言后，其他参与者已经能够看到全部隐藏事实；官方结果文件还在第一轮四条响应后重复附加完整隐藏事实块。旧实验只在每个 Agent 第一次发言时附加该 Agent 自己的事实，因此属于 owner-by-owner mechanical reveal，而不是严格复现。",
    )
    _table(
        doc,
        ["协议", "第一位发言后可见信息", "第一轮注入", "定位"],
        [
            [
                "官方 GPT-4.1 Reveal-All",
                "全部四条隐藏事实",
                "四条响应均带全局事实块",
                "作者官方条件",
            ],
            [
                "旧 fixed/dynamic-reveal-all",
                "只出现第一位发言者自己的事实",
                "每人只附加自己的事实",
                "逐人机械披露",
            ],
            [
                "本轮 official-compatible global Reveal-All",
                "全部四条隐藏事实",
                "四条响应均附加全部事实",
                "官方同款补充条件",
            ],
        ],
        [1600, 2600, 2580, 2580],
        font_size=9,
    )
    _callout(
        doc,
        "关键结果：",
        "新增 30 supplementary runs 共使用 2,160 次真实模型请求，ID1、ID2、ID3 均为 10/10 多数正确。每题样本仍只有 10 次，因此 Wilson 95% 区间均约为 72.2%–100.0%，不能把观察到的 100% 当成总体真值。",
    )
    _paragraph(
        doc,
        "MAF 边界：MAF 是模型-Agent 执行层，不是本实验的完整编排器。动态五因子机制由本项目代码执行，是一个 external centralized scheduler，并且读取完整 assignment 中的私有事实清单。它具有特权信息，因此不能称为完全去中心化。",
    )
    _paragraph(
        doc,
        "Seed 边界：pair seed 只固定私有事实分配和条件配对；generation seed was not transmitted to DeepSeek。温度 0 也不能消除服务端非确定性。",
    )
    _paragraph(
        doc,
        "指标边界：选择器运行时的 undisclosed 项依靠词法规则判断；报告中的原子披露率则由 AI 逐事实审计并由规则验证证据。两者不是同一个指标。五个选择器权重 0.30/0.30/0.15/0.15/0.10 是预设设计选择，尚未进行敏感性分析。",
    )


def _task_section(doc: DocumentType, analysis: dict[str, Any]) -> None:
    _heading(doc, "3. 三道官方短题到底是什么")
    _paragraph(
        doc,
        "三题共享同一个山村撤离场景和 West City、East Town、North Hill 三个选项，但公共信息与四条分散私有信息不同。以下完整列出本轮使用的题目事实和标准答案依据。",
    )
    for task in analysis["task_details"]:
        _heading(
            doc,
            f"ID{task['task_id']}｜{task['name']}｜正确答案：{task['correct_answer']}",
            level=2,
        )
        rows: list[list[Any]] = []
        for index, fact in enumerate(task["shared_information"], 1):
            rows.append(["公共信息", index, fact])
        for index, fact in enumerate(task["private_facts"], 1):
            rows.append(["私有信息", index, fact])
        _table(
            doc,
            ["类型", "序号", "内容"],
            rows,
            [1400, 800, 7160],
            font_size=8.5,
        )
        _paragraph(
            doc,
            f"Evidence chain：{task['evidence_chain']}",
            bold_lead="Evidence chain：",
        )


def _method_section(doc: DocumentType) -> None:
    _heading(doc, "4. 实验设置与指标")
    _table(
        doc,
        ["条件 ID", "公开发言", "机制", "解释"],
        [
            ["fixed-60", 60, "固定轮转", "原确认性基线；含15轮影子投票"],
            ["dynamic-60", 60, "五因子顺序调度", "每人严格15次；选择器0次LLM调用"],
            ["fixed-reveal-all", 60, "逐人机械披露", "旧 owner-by-owner mechanical reveal"],
            ["dynamic-reveal-all", 60, "逐人机械披露+调度", "旧逐人披露动态条件"],
            ["official-global-reveal", 60, "固定轮转+全局披露", "新增官方同款补充条件"],
            ["structured-12", 12, "Exchange→Decide", "结构化短协议"],
            ["fixed-12", 12, "固定轮转", "同12发言预算对照"],
            ["single-local", 0, "单Agent投票", "只见共享信息和自己一条私有信息"],
        ],
        [1850, 900, 2500, 4110],
        font_size=8.5,
    )
    for text in (
        "多数正确率：一次运行中，标准答案获得严格多数票；2-2 平票记为未获得正确多数。",
        "错误共识率：四个 Agent 最终全体一致，但一致答案错误。",
        "原子披露率：AI 判断每条拥有者事实是否在拥有者公开消息中被忠实表达；规则再核验消息 ID、所有者、原文子串和极性。",
        "Wilson 95%：对每个 10 次重复的正确率给出 Wilson 95% 置信区间，避免只展示0%、60%或100%的点估计。",
        "McNemar：对相同题目、重复编号和私有信息分配的本地条件进行双侧精确配对检验；官方 GPT-4.1 与本地 DeepSeek 不做配对显著性检验。",
    ):
        _bullet(doc, text)


def _disclosure_section(
    doc: DocumentType,
    analysis: dict[str, Any],
    audit_prompt: str,
) -> None:
    _heading(doc, "5. 披露率的 AI 输出、规则验证与百分比")
    example = analysis["disclosure_example"]
    key = example["study_key"]
    _paragraph(
        doc,
        "披露百分比不是让模型随口报数。审计模型只对每条事实返回 disclosed=true/false、消息 ID、精确引文、理由和置信度；程序验真后再求和。",
    )
    _callout(
        doc,
        "完整计算实例：",
        f"study key = task-{key['task_id']}:{key['condition']}:rep-{key['repetition']}；run ID = {example['run_id']}；{example['arithmetic']}。",
    )
    fact_by_id = {fact["fact_id"]: fact for fact in example["facts"]}
    rows = []
    for judgment in example["judgments"]:
        rows.append(
            [
                judgment["fact_id"][-8:],
                fact_by_id[judgment["fact_id"]]["owner_agent_id"],
                "是" if judgment["disclosed"] else "否",
                ", ".join(judgment["evidence_message_ids"]) or "—",
                judgment["evidence_quote"] or judgment["reason"],
            ]
        )
    _table(
        doc,
        ["事实ID", "拥有者", "披露", "证据消息", "引文或未披露理由"],
        rows,
        [900, 1000, 700, 1950, 4810],
        font_size=7.8,
    )
    _paragraph(
        doc,
        "当前210次原子审计由 deepseek-v4-flash 判断 DeepSeek 生成的对话，虽然证据经过确定性规则验证，但仍不是人工金标准，存在同模型自审偏差。机械披露条件的100%由系统注入记录直接确认，不调用审计模型。",
    )
    _code(doc, "5.1 原子事实审计完整 Prompt", audit_prompt)


def _results_section(
    doc: DocumentType,
    analysis: dict[str, Any],
    confirmatory: dict[str, Any],
) -> None:
    _heading(doc, "6. 结果：官方、旧条件与官方同款补充")
    official = analysis["official_gpt41"]
    official_rows = []
    for task_id in (1, 2, 3):
        base = official["baseline"]["by_task"][str(task_id)]
        reveal = official["reveal_all"]["by_task"][str(task_id)]
        official_rows.append(
            [
                f"ID{task_id}",
                f"{round(base['majority_accuracy'] * base['runs'])}/{base['runs']}",
                _pct(base["majority_accuracy"]),
                _ci(round(base["majority_accuracy"] * base["runs"]), base["runs"]),
                f"{round(reveal['majority_accuracy'] * reveal['runs'])}/{reveal['runs']}",
                _pct(reveal["majority_accuracy"]),
            ]
        )
    _table(
        doc,
        ["题", "官方Baseline", "正确率", "Wilson 95%", "官方Reveal-All", "正确率"],
        official_rows,
        [700, 1300, 1100, 2100, 1850, 2310],
        font_size=8.5,
    )
    _paragraph(
        doc,
        "官方数据与本地数据是跨模型、跨实现的描述性对标：官方为 GPT-4.1，本地为 DeepSeek；不能把差异归因为 MAF、调度机制或单一因素。",
    )
    result_rows = []
    ordered_conditions = (
        "fixed-60",
        "dynamic-60",
        "fixed-reveal-all",
        "dynamic-reveal-all",
        "official-global-reveal",
        "structured-12",
        "fixed-12",
        "single-local",
    )
    for condition in ordered_conditions:
        for task_id in (1, 2, 3):
            stats = analysis["condition_statistics"][f"{condition}:ID{task_id}"]
            if condition == "official-global-reveal":
                disclosure = "100.0%（机械）"
            else:
                old = confirmatory["by_condition_task"][f"{condition}:ID{task_id}"]
                disclosure = _pct(old["atomic_disclosure_rate"])
            result_rows.append(
                [
                    CONDITION_LABELS[condition],
                    f"ID{task_id}",
                    f"{stats['successes']}/{stats['runs']}",
                    _pct(stats["rate"]),
                    f"{stats['wilson_95_low'] * 100:.1f}%–{stats['wilson_95_high'] * 100:.1f}%",
                    disclosure,
                ]
            )
    _table(
        doc,
        ["条件", "题", "多数正确", "正确率", "Wilson 95%", "原子披露率"],
        result_rows,
        [2400, 650, 1100, 1000, 2210, 2000],
        font_size=7.8,
    )
    _callout(
        doc,
        "解释限制：",
        "官方同款全局 Reveal-All 在三题均10/10正确，解决了旧逐人披露的协议不一致；但每题只有10次，且仍使用DeepSeek，因此只能报告本样本观察值，不能宣称机制已经得到普遍证明。",
    )


def _statistics_section(doc: DocumentType, analysis: dict[str, Any]) -> None:
    _heading(doc, "7. 配对统计检验")
    rows = []
    labels = {
        "fixed-60_vs_dynamic-60": "固定60 vs 动态60",
        "fixed-reveal-all_vs_dynamic-reveal-all": "固定逐人披露 vs 动态逐人披露",
        "fixed-12_vs_structured-12": "固定12 vs Structured-12",
    }
    for key, label in labels.items():
        value = analysis["paired_exact_tests"][key]
        rows.append(
            [
                label,
                value["left_successes"],
                value["right_successes"],
                value["a_only"],
                value["b_only"],
                f"{value['p_value']:.3f}",
            ]
        )
    _table(
        doc,
        ["配对比较", "左条件正确", "右条件正确", "仅左正确", "仅右正确", "精确 p"],
        rows,
        [2600, 1250, 1250, 1100, 1100, 2060],
        font_size=8.5,
    )
    _paragraph(
        doc,
        "三个预设比较的双侧精确 McNemar p 值分别为0.125、0.125和0.388，均未达到0.05。因此，动态条件或Structured条件的点估计较高不能表述为统计上显著优于对应固定条件。",
    )


def _early_stop_section(doc: DocumentType, analysis: dict[str, Any]) -> None:
    _heading(doc, "8. 同轨迹早停：不是答案翻转，而是最终平票")
    early = analysis["early_stop"]
    exceptional = early["exceptional_cases"][0]
    _table(
        doc,
        ["固定60运行", "同一多数", "最终平票", "候选正确", "最终正确"],
        [[30, early["classifications"].get("same-majority", 0), early["classifications"].get("final-tie", 0), f"{early['candidate_correct']}/30", f"{early['final_correct']}/30"]],
        [1800, 1700, 1700, 2000, 2160],
        font_size=9,
    )
    _callout(
        doc,
        "唯一异常：",
        f"task-{exceptional['study_key']['task_id']} / repetition-{exceptional['study_key']['repetition']} / run {exceptional['run_id']}：第2轮候选为 North Hill，最终四票为 North Hill、West City、North Hill、West City，即 2-2 final tie。候选与最终均错误。",
    )
    _paragraph(
        doc,
        "因此，本样本中使用候选点停止不会改变总正确次数（候选17/30，最终17/30），但有1次早期全体一致没有保持到最终。该结果不能外推为早停普遍安全。",
    )


def _mast_section(doc: DocumentType, analysis: dict[str, Any]) -> None:
    _heading(doc, "9. MAST 失败模式的具体位置")
    rows = []
    for case in analysis["mast_cases"]:
        if case["study_key"]:
            key = case["study_key"]
            location = (
                f"ID{key['task_id']} / {key['condition']} / rep{key['repetition']}\n"
                f"run={case['run_id']}"
            )
        else:
            location = "未确认具体运行"
        rows.append(
            [
                case["mast_mode"],
                case["boundary"],
                location,
                case["interpretation"],
            ]
        )
    _table(
        doc,
        ["模式", "判断边界", "具体位置", "证据解释"],
        rows,
        [900, 2100, 2200, 4160],
        font_size=8.3,
    )
    _paragraph(
        doc,
        "FM-2.4 已由拥有者事实未出现在拥有者公开消息中直接确认。FM-2.5 目前是强候选：信息已100%公开，但最终四个理由仍把被阻塞的East Town隧道解释为畅通；要断言“忽略”仍需人工因果编码。FM-2.6 未发现经验证的理由-行动冲突，不虚构发生率。",
    )


def _limitations_section(doc: DocumentType) -> None:
    _heading(doc, "10. 必须保留的局限")
    for text in (
        "只研究官方short的3题，不能外推至全部65题或其他任务类型。",
        "官方使用GPT-4.1，本地使用DeepSeek；框架、模型和提示词效应仍未完全正交分离。",
        "本地条件每题10次，置信区间较宽，配对比较均未达到0.05显著性水平。",
        "MAF只负责Agent与模型调用；完整讨论协议是项目自写，不应称为MAF原生GroupChat复现。",
        "动态调度器读取全局私有事实清单，属于带特权信息的中心化外部调度器。",
        "DeepSeek生成随机性没有服务端seed控制；pair seed只控制事实分配。",
        "AI披露审计不是人工金标准；当前证据规则能够验真引文，但不能消除同模型判断偏差。",
        "本轮新增官方同款Reveal-All是独立补充研究，不修改原210次实验及其发布哈希。",
    ):
        _bullet(doc, text)


def _artifact_section(
    doc: DocumentType,
    analysis: dict[str, Any],
    supplement_gate: dict[str, Any],
    global_message: str,
    supplement_package: tuple[str, str],
) -> None:
    _heading(doc, "11. 复现证据、门禁与实际提示内容")
    _paragraph(
        doc,
        f"原确认性实验：210次运行、210次原子审计；新增补充实验：30次运行、{supplement_gate['api_requests']:,}次API请求。补充门禁10项全部通过。",
    )
    _table(
        doc,
        ["文件", "SHA-256"],
        [[name, digest] for name, digest in analysis["source_sha256"].items()],
        [3500, 5860],
        font_size=7.8,
    )
    _paragraph(
        doc,
        "代码分支：https://github.com/sbw51898-beep/mas/tree/codex/budget-matched-dynamic；Pull Request：https://github.com/sbw51898-beep/mas/pull/5。原210次证据仍位于 hiddenbench-confirmatory-20260804 Release；新增补充证据另行打包并记录SHA-256。",
    )
    _table(
        doc,
        ["补充证据包", "SHA-256"],
        [[supplement_package[0], supplement_package[1]]],
        [4300, 5060],
        font_size=7.8,
    )
    _code(doc, "11.1 官方同款全局 Reveal-All 第一条公开消息", global_message)


def build_report(
    *,
    root: Path,
    output: Path,
    desktop_output: Path | None,
) -> Path:
    analysis = json.loads(
        (root / "reports/data/hiddenbench-hardened-analysis-20260804.json").read_text(
            encoding="utf-8"
        )
    )
    confirmatory = json.loads(
        (root / "reports/data/hiddenbench-confirmatory-summary.json").read_text(
            encoding="utf-8"
        )
    )
    supplement_gate = json.loads(
        (
            root
            / "artifacts/hiddenbench-official-reveal-supplement-20260804.gate.json"
        ).read_text(encoding="utf-8")
    )
    if analysis["scope"] != {
        "confirmatory_runs": 210,
        "atomic_audits": 210,
        "official_reveal_supplement_runs": 30,
        "task_ids": [1, 2, 3],
    }:
        raise ValueError("hardened report requires the locked 210+30 scope")
    if not confirmatory["gate_passed"] or not supplement_gate["passed"]:
        raise ValueError("all formal data gates must pass before report build")
    confirmatory_rows = _read_jsonl(
        root / "artifacts/hiddenbench-confirmatory-20260804.jsonl"
    )
    supplement_rows = _read_jsonl(
        root / "artifacts/hiddenbench-official-reveal-supplement-20260804.jsonl"
    )
    example_row = _find_confirmatory(
        confirmatory_rows,
        analysis["disclosure_example"]["study_key"],
    )
    audit_prompt = build_atomic_disclosure_prompt(
        ConfirmatoryRunRecord.model_validate(example_row).run
    )
    global_message = supplement_rows[0]["run"]["discussion_messages"][0][
        "content"
    ]
    package_sha_path = (
        root
        / "release"
        / "hiddenbench-official-reveal-supplement-20260804-raw.sha256"
    )
    package_sha, package_name = package_sha_path.read_text(
        encoding="utf-8"
    ).strip().split(maxsplit=1)
    old_report = root / "reports/HiddenBench确认性修订实验报告_2026-08-04.docx"
    old_hash_before = hashlib.sha256(old_report.read_bytes()).hexdigest()

    doc = Document()
    _configure(doc)
    _masthead(doc)
    _callout(
        doc,
        "结论先行：",
        "新增官方同款全局Reveal-All后，DeepSeek在ID1/ID2/ID3均10/10正确；但样本量小、跨模型不可归因、配对机制比较不显著。新版报告据此收紧了全部结论。",
    )
    _teacher_answers(doc)
    _protocol_section(doc)
    _task_section(doc, analysis)
    _method_section(doc)
    _disclosure_section(doc, analysis, audit_prompt)
    _results_section(doc, analysis, confirmatory)
    _statistics_section(doc, analysis)
    _early_stop_section(doc, analysis)
    _mast_section(doc, analysis)
    _limitations_section(doc)
    _artifact_section(
        doc,
        analysis,
        supplement_gate,
        global_message,
        (package_name, package_sha),
    )

    doc.core_properties.title = "HiddenBench 确认性完善实验报告"
    doc.core_properties.subject = "官方同款 Reveal-All、统计校正与 MAST 案例"
    doc.core_properties.author = "MAS 知识治理项目组"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    if hashlib.sha256(old_report.read_bytes()).hexdigest() != old_hash_before:
        raise ValueError("the preserved original report changed during build")
    check = Document(output)
    if len(check.tables) < 10:
        raise ValueError("hardened report table count is incomplete")
    for table in check.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        if width is None or width.get(qn("w:w")) != "9360":
            raise ValueError("report table width validation failed")
        if indent is None or indent.get(qn("w:w")) != "120":
            raise ValueError("report table indent validation failed")
    if desktop_output is not None:
        desktop_output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(output, desktop_output)
    return output


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    output = root / "reports" / "HiddenBench确认性完善实验报告_2026-08-04.docx"
    desktop = Path.home() / "Desktop" / output.name
    print(build_report(root=root, output=output, desktop_output=desktop))


if __name__ == "__main__":
    main()
