"""Build a human-friendly Excel sign-off sheet for the 144-item blind review.

The sheet stays blind: it shows only blind_id, the private fact, the owner,
and the owner's messages. The AI/rule labels are NOT included. Each row has
drop-downs for the disclosure decision (是/否) and free-text columns for
evidence message IDs, quotes, and reasons.
"""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
QUEUE_CSV = (
    ROOT
    / "artifacts"
    / "hiddenbench-stability-20260729.blind-review.queue.csv"
)
OUTPUT = ROOT / "reports" / "盲审签核表_144项_2026-08-03.xlsx"
DESKTOP_OUTPUT = Path.home() / "Desktop" / "盲审签核表_144项_2026-08-03.xlsx"


def main() -> None:
    rows = list(csv.DictReader(QUEUE_CSV.open(encoding="utf-8-sig")))
    if len(rows) != 144:
        raise ValueError(f"expected 144 queue rows, found {len(rows)}")

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "填写说明"
    instructions.column_dimensions["A"].width = 100
    instruction_lines = [
        ("盲审签核表（144 项）填写说明", True),
        ("", False),
        ("1. 本表用于确认 AI 披露审计的结果：对每一行判断该私有信息包是否被其拥有者在公开发言中披露。", False),
        ("2. 判定规则（与审计 prompt 一致）：", False),
        ("   a. 只有信息拥有者本人说出原事实、或给出保留决策关键含义的忠实转述，才算披露。", False),
        ("   b. 极性颠倒、与事实矛盾、或把关键断言明显弱化，不算披露。", False),
        ("   c. 只说出部分内容时，只有保留了影响决策的核心部分才算披露。", False),
        ("   d. 非拥有者的猜测或重复，不替代拥有者披露其私有信息。", False),
        ("   e. 判“是”时，必须填写拥有者发言的证据消息 ID 和原文引句（程序会校验）。", False),
        ("   f. 判“否”时，证据列留空。", False),
        ("3. 操作：在“是否披露”列选择 是/否；证据消息 ID 可从拥有者发言列中复制；引句必须是原文中的连续片段。", False),
        ("4. 填完后把本文件发回，或导出为 CSV；程序会校验证据并计算人工一致率、AI 精确率、召回率和修订披露率。", False),
        ("", False),
        ("说明：本表不含 AI 原标签与规则标签，保证盲审。每条 2-3 分钟，全部约 5-7 小时。", False),
    ]
    for index, (text, bold) in enumerate(instruction_lines, start=1):
        cell = instructions.cell(row=index, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True, size=14)

    sheet = workbook.create_sheet("盲审签核")
    headers = [
        "序号",
        "盲审编号",
        "私有信息包（待判断内容）",
        "信息拥有者",
        "拥有者公开发言（含消息ID）",
        "是否披露（是/否）",
        "证据消息ID（判“是”必填）",
        "证据原文引句（判“是”必填）",
        "理由（可选）",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="17365D")
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"

    for index, row in enumerate(rows, start=1):
        messages = json.loads(row["owner_messages_json"])
        message_text = "\n".join(
            f"[{item['message_id']}] {item['content']}"
            for item in messages
        )
        sheet.append(
            [
                index,
                row["blind_id"],
                row["fact"],
                row["owner_agent_id"],
                message_text,
                "",
                "",
                "",
                "",
            ]
        )

    widths = {
        1: 6,
        2: 18,
        3: 50,
        4: 12,
        5: 90,
        6: 14,
        7: 28,
        8: 50,
        9: 40,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=145):
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )
        sheet.row_dimensions[row[0].row].height = 90

    validation = DataValidation(
        type="list",
        formula1='"是,否"',
        allow_blank=True,
        showDropDown=False,
    )
    validation.error = "只能选择 是 或 否"
    validation.errorTitle = "无效输入"
    sheet.add_data_validation(validation)
    validation.add(f"F2:F{len(rows) + 1}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    DESKTOP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, DESKTOP_OUTPUT)
    print(OUTPUT)
    print(DESKTOP_OUTPUT)


if __name__ == "__main__":
    main()
