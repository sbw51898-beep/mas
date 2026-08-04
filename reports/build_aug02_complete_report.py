from __future__ import annotations

import hashlib
import json
import shutil
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from reports.build_hardened_confirmatory_report import (  # noqa: E402
    BLACK,
    BLUE,
    DARK,
    GRAY,
    _callout,
    _code,
    _configure,
    _font,
    _heading,
    _page_field,
    _paragraph,
    _table,
)


REPORT_NAME = "给彭老师的完整实验报告_2026-08-02至2026-08-04.docx"
PRESERVED_REPORTS = (
    "对标与披露治理改进实验报告_2026-08-02.docx",
    "整合治理（Structured协议）实验报告_2026-08-02.docx",
    "对照实验报告（单智能体与同预算基线）_2026-08-03.docx",
    "实验总报告（自老师提要求以来）_2026-08-03.docx",
    "HiddenBench确认性完善实验报告_2026-08-04.docx",
)
CONDITION_LABELS = {
    "single-local": "局部信息单智能体",
    "fixed-12": "固定轮转 12",
    "structured-12": "Exchange→Decide 12",
    "fixed-60": "固定轮转 60",
    "dynamic-60": "五因子调度 60",
    "fixed-reveal-all": "固定 60 + 逐人机械披露",
    "dynamic-reveal-all": "动态 60 + 逐人机械披露",
    "official-global-reveal": "官方同款全局 Reveal-All",
}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _fmt_pct(value: float, decimals: int = 1) -> str:
    return f"{value * 100:.{decimals}f}%"


def _fmt_success(stat: dict[str, Any]) -> str:
    return f"{stat['successes']}/{stat['runs']} ({_fmt_pct(stat['rate'])})"


def _fmt_ci(stat: dict[str, Any]) -> str:
    return (
        f"{_fmt_pct(stat['wilson_95_low'])}–"
        f"{_fmt_pct(stat['wilson_95_high'])}"
    )


def _configure_complete(doc: DocumentType) -> None:
    _configure(doc)
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "HiddenBench × MAF｜2026年8月2日至8月4日完整实验报告"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        _font(run, 8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("MAS 知识治理项目｜第 ")
    _font(label, 8.5, color=GRAY)
    _page_field(footer)
    suffix = footer.add_run(" 页")
    _font(suffix, 8.5, color=GRAY)


def _bottom_rule(doc: DocumentType) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(12)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _masthead(doc: DocumentType, summary: dict[str, Any]) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(4)
    _font(kicker.add_run("TECHNICAL EXPERIMENT REPORT"), 10, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    _font(
        title.add_run("HiddenBench × MAF 完整实验报告"),
        23,
        bold=True,
        color=BLACK,
    )
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    _font(
        subtitle.add_run("2026年8月2日至8月4日：从披露治理到官方三题确认性复现"),
        13.5,
        color=GRAY,
    )
    scope = summary["scope"]
    metadata = (
        ("汇报对象", "彭老师"),
        ("实验期间", "2026年8月2日—8月4日；报告生成于2026年8月5日"),
        ("框架与模型", "Microsoft Agent Framework 调用层；deepseek-v4-flash"),
        ("实验规模", f"{scope['local_runs']} 次本地运行；{scope['total_model_requests']:,} 次模型请求"),
        ("当前状态", "六组数据门禁全部通过；结论按探索性/确认性边界表述"),
    )
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        _font(paragraph.add_run(f"{label}："), 10.5, bold=True)
        _font(paragraph.add_run(value), 10.5)
    _bottom_rule(doc)


def _page_break(doc: DocumentType) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _bullets(doc: DocumentType, items: list[str]) -> None:
    for text in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167
        _font(paragraph.add_run(text))


def _technical_summary(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "技术摘要：多智能体的收益取决于信息使用方式，而非是否达成共识")
    _callout(
        doc,
        "核心结论：",
        "从8月2日开始的560次本地运行表明，多智能体并不天然优于单智能体；"
        "真正影响结果的是私有信息能否被公开、是否被正确解释，以及发言协议如何组织信息。"
        "8月4日新增官方同款全局Reveal-All后，DeepSeek在官方三题均为10/10正确，"
        "但每题仅10次，且与官方GPT-4.1不是同模型实验。",
    )
    _bullets(
        doc,
        [
            "披露不等于正确。ID2的旧固定逐人机械披露达到100%，多数正确率仍只有6/10；一个FM-2.5候选案例在信息全部公开后仍形成错误共识。",
            "动态和Structured条件的点估计有改善，但三组配对精确McNemar检验p值为0.125、0.125和0.388，均未达到0.05。",
            "单智能体不是无条件更差。官方三题上的局部信息单智能体仅2/10、2/10、0/10，但固定60在ID2也为0/10，说明多智能体同样会系统性失败。",
            "同轨迹早停评估中，29/30个候选多数与最终多数一致；唯一异常是2-2 final tie。候选与最终正确数均为17/30，不能据此宣称早停普遍安全。",
            "144项复核由Codex模型辅助完成，不是人工金标准；当前披露指标仍存在同模型审计偏差，需要后续人工盲审。",
        ],
    )
    _callout(
        doc,
        "本报告不宣称：",
        "不把跨模型差异归因于MAF，不把10/10当作总体真值，不把共识当作正确，"
        "也不把探索性四题结果推广到HiddenBench全部65题。",
    )


def _teacher_answers(doc: DocumentType) -> None:
    _heading(doc, "1. 对老师问题的直接回答")
    rows = [
        ["原论文用什么框架和模型？", "HiddenBench使用作者自写Python模拟器，不是MAF。本文直接对标的官方short结果模型为GPT-4.1。"],
        ["原论文有没有源码？", "有。仓库为https://github.com/Yassellee/HiddenBench_ICML；本项目核验提交3be6ca16973e4fb751ffc0dfb7eb11f2d28335d1。"],
        ["我们真的用了MAF吗？", "用了，但MAF只负责Agent/模型调用。轮转、动态选择、Reveal、投票、审计和门禁均为项目代码。"],
        ["为什么不用贺同学代码？", "本项目从官方任务和公开结果出发独立实现，便于逐项核验协议；贺同学工作只作为研究思路参考。"],
        ["披露法则是什么？", "普通条件由模型自然决定是否披露；旧Reveal条件逐Agent附加自己的事实；官方同款条件在第一轮四条响应中都附加全部四条事实。"],
        ["一人只有一条私有信息，首轮不是说完了吗？", "自然条件不保证首轮忠实说出；有时只说结论、弱化关键信息或完全省略。只有机械Reveal保证公开。"],
        ["披露百分比在哪里？", "审计模型逐条返回disclosed、消息ID和原文引文；规则程序验真后再以已披露原子事实数/原子事实总数计算。正文给出3/4=75%完整案例。"],
        ["前几轮一致后为什么还继续？", "为保持60发言预算和同轨迹比较。影子投票不回灌讨论，事后判断候选早停；唯一异常最终形成2-2平票。"],
        ["多智能体一定比单智能体好吗？", "不一定。报告同时保留单智能体、固定/动态多智能体、Structured和Reveal条件，错误结果也全部报告。"],
    ]
    _table(doc, ["老师的问题", "本报告的回答"], rows, [2300, 7060], font_size=8.4)


def _timeline(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "2. 三天工作时间线：从探索性治理到官方任务确认")
    counts = summary["stage_counts"]
    run_requests = summary["stage_api_requests"]
    audit_counts = summary["audit_counts"]
    audit_requests = summary["audit_api_requests"]
    rows = [
        ["8月2日", "强制披露治理", f"{counts['governance_20260802']}次", f"{run_requests['governance_20260802']:,}+{audit_requests['governance_20260802']}次", "固定/动态各60发言；ID1/5/7/25"],
        ["8月2日", "Structured协议", f"{counts['structured_20260802']}次", f"{run_requests['structured_20260802']:,}+{audit_requests['structured_20260802']}次", "Exchange→Decide；每次12发言"],
        ["8月2日", "分层复核", "144/320项", "不计DeepSeek运行", "Codex模型辅助复核，不是人工金标准"],
        ["8月3日", "单智能体/同预算对照", f"{counts['contrast_20260803']}次", f"{run_requests['contrast_20260803']:,}+{audit_requests['contrast_20260803']}次", "single-direct、single-reflect、fixed-12"],
        ["8月3日", "4/8发言早停探索", f"{counts['earlystop_20260803']}次", f"{run_requests['earlystop_20260803']:,}+{audit_requests['earlystop_20260803']}次", "独立短轨迹；不能替代同轨迹验证"],
        ["8月4日", "官方三题确认性实验", f"{counts['confirmatory_20260804']}次", f"{run_requests['confirmatory_20260804']:,}+{audit_requests['confirmatory_20260804']}次", "ID1/2/3；7条件；210项原子审计"],
        ["8月4日", "官方同款Reveal补充", f"{counts['official_global_reveal_20260804']}次", f"{run_requests['official_global_reveal_20260804']:,}次", "第一轮全局注入；披露100%由系统记录确认"],
    ]
    _table(doc, ["日期", "工作", "规模", "模型请求", "定位"], rows, [1000, 2100, 1100, 1700, 3460], font_size=7.7)
    scope = summary["scope"]
    _callout(
        doc,
        "累计规模：",
        f"{scope['local_runs']} 次本地运行，{scope['discussion_vote_api_requests']:,} 次讨论/投票请求，"
        f"{scope['audit_api_requests']} 次披露审计请求，共 {scope['total_model_requests']:,} 次模型请求。"
        "官方GPT-4.1公开结果只做外部参照，不计入本地运行数。",
    )


def _paper_boundary(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "3. 原论文、官方结果与本项目的边界")
    official = summary["official_code"]
    _paragraph(
        doc,
        f"HiddenBench论文为{official['paper']}。作者使用自写Python模拟器组织四名参与者的"
        "Hidden Profile讨论，并公开65题数据、Prompt、评测代码和逐轮结果。"
        f"代码仓库：{official['repository']}；官方结果：{official['results']}。",
    )
    rows = [
        ["实现层", "作者自写Python simulator/ModelClient", "MAF Agent与模型适配；协议为项目代码"],
        ["对标模型", "GPT-4.1 short benchmark", "deepseek-v4-flash"],
        ["核心题目", "官方short ID1/2/3", "8月4日使用完全相同ID1/2/3"],
        ["Agent数量", "4", "多智能体条件4；单智能体条件1"],
        ["讨论预算", "15轮×4人=60条响应", "60发言条件一致；另设12/8/4发言"],
        ["Reveal-All", "首位发言后所有人可见全部隐藏事实；第一轮响应附加完整事实", "新增official-compatible global Reveal-All严格按第一轮全局注入"],
        ["可归因性", "官方系统内部比较", "跨框架且跨模型，只能描述性对照"],
    ]
    _table(doc, ["维度", "HiddenBench官方", "本项目"], rows, [1400, 3780, 4180], font_size=8.0)
    _callout(
        doc,
        "MAF边界：",
        "Microsoft Agent Framework不是完整实验编排器。五因子动态机制是一个external centralized scheduler，"
        "它读取完整assignment中的私有事实清单，属于带特权信息的中心化外部调度器。",
    )
    _paragraph(
        doc,
        "Seed边界：pair seed只固定事实分配和配对；generation seed was not transmitted to DeepSeek。"
        "温度设为0仍不能消除服务端非确定性，因此同一配置需要重复运行。",
    )
    _paragraph(
        doc,
        "协议边界：旧fixed/dynamic reveal-all属于owner-by-owner mechanical reveal，"
        "即每名Agent首次发言后只附加自己的事实；新增official-compatible global Reveal-All"
        "才在第一轮四条响应中都附加全部四条事实。",
    )


def _method_and_metrics(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "4. 实验设置：任务、信息结构、协议与指标")
    configs = summary["configs"]
    selector = configs["governance"]["selector"]
    rows = [
        ["模型", "deepseek-v4-flash；temperature=0；thinking=disabled"],
        ["正式题目", "探索性：ID1/5/7/25；确认性：官方short ID1/2/3"],
        ["信息分配", "4名Agent共享公共信息；每名Agent各持有1条私有事实；事实顺序随机打乱"],
        ["固定60", "4名Agent固定轮转15轮，共60条公开发言"],
        ["动态60", "同为60条发言且每人严格15条；每次选择不增加LLM调用"],
        ["Structured-12", "两轮Exchange后一次Decide，共12条发言"],
        ["单智能体", "只看到公共信息和自己的一条私有事实，不参与公开讨论"],
        ["动态权重", f"分歧{selector['disagreement']:.2f}、未披露{selector['undisclosed']:.2f}、相关讨论{selector['related_discussion']:.2f}、应答{selector['response_due']:.2f}、等待{selector['waiting']:.2f}"],
    ]
    _table(doc, ["项目", "锁定设置"], rows, [1900, 7460], font_size=8.3)
    metric_rows = [
        ["多数正确率", "每次运行中，正确答案是否获得严格多数票；2-2记为未获得正确多数"],
        ["错误共识率", "四名Agent最终一致，但一致答案错误的运行比例"],
        ["原子披露率", "被拥有者公开且通过证据校验的原子私有事实数/原子私有事实总数"],
        ["Wilson 95%", "对每题每条件10次重复的二项正确率给区间，不只展示0%、60%或100%点估计"],
        ["精确McNemar", "在相同题目、重复编号和事实分配上比较成对条件；双侧精确检验"],
        ["同轨迹早停", "影子投票不写回讨论；候选点与同一条60发言轨迹的最终投票比较"],
    ]
    _table(doc, ["指标", "定义与分母"], metric_rows, [1900, 7460], font_size=8.3)


def _aug02_results(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "5. 8月2日：强制披露解决了“披露率为0”的测量问题，但没有消除错误共识")
    exploration = summary["exploratory"]
    correct = exploration["correct_runs_by_condition"]
    _table(
        doc,
        ["条件", "正确运行", "总体正确率", "解释"],
        [
            ["固定60+disclosure-first", f"{correct['fixed-disc']}/40", _fmt_pct(correct['fixed-disc']/40), "固定轮转、每人15条"],
            ["动态60+disclosure-first", f"{correct['dynamic-disc']}/40", _fmt_pct(correct['dynamic-disc']/40), "等预算动态排序"],
            ["Structured-12", f"{correct['structured']}/40", _fmt_pct(correct['structured']/40), "Exchange→Decide"],
        ],
        [2600, 1500, 1600, 3660],
    )
    gov_rows = []
    for row in exploration["governance_rows"]:
        gov_rows.append(
            [
                f"ID{row['task_id']}",
                "固定" if row["condition"] == "fixed-disc" else "动态",
                f"{row['majority_correct']}/{row['runs']}",
                f"{row['wrong_consensus']}/{row['runs']}",
                _fmt_pct(row["ai_disclosure_rate"]),
                _fmt_pct(row["rule_disclosure_rate"]),
            ]
        )
    _table(
        doc,
        ["题", "条件", "多数正确", "错误共识", "AI披露", "规则披露"],
        gov_rows,
        [850, 950, 1500, 1500, 1800, 2760],
        font_size=7.8,
    )
    _paragraph(
        doc,
        "结果解释：ID1和ID25较稳定，ID5和ID7在两种60发言机制下仍大量错误。"
        "这说明披露规则可以避免审计全为0，却不能保证关键信息被正确理解。"
        "AI披露率与规则披露率在部分题上差异很大，促成了后续分层复核和原子化审计。",
    )
    structured_rows = [
        [
            f"ID{row['task_id']}",
            f"{row['majority_correct']}/{row['runs']}",
            f"{row['wrong_consensus']}/{row['runs']}",
            _fmt_pct(row['ai_disclosure_rate']),
        ]
        for row in exploration["structured_rows"]
    ]
    _heading(doc, "5.1 Structured协议压缩到12条发言，整体正确20/40", level=2)
    _table(
        doc,
        ["题", "多数正确", "错误共识", "AI披露率"],
        structured_rows,
        [1200, 1900, 1900, 4360],
    )
    blind = summary["blind_review"]
    _heading(doc, "5.2 144项分层复核揭示AI审计漏检", level=2)
    _table(
        doc,
        ["项目", "结果", "正确解释"],
        [
            ["复核总体", f"{blind['reviewed_size']}/{blind['population_size']}项", "覆盖全部分歧项并分层抽样一致项"],
            ["原AI与Codex复核一致率", _fmt_pct(blind['ai_human_label_agreement_estimate']), "字段沿用旧summary命名，复核者不是人类"],
            ["估计精确率", _fmt_pct(blind['estimated_precision']), "被AI判为已披露的判断多数可靠"],
            ["估计召回率", _fmt_pct(blind['estimated_recall']), "AI漏掉约三分之一实际披露"],
            ["修订披露率", _fmt_pct(blind['revised_disclosure_rate']), "基于抽样估计，不是最终人工金标准"],
        ],
        [2300, 1900, 5160],
        font_size=8.2,
    )
    _callout(
        doc,
        "审计边界：",
        "Codex 模型辅助复核，不是人工金标准。8月3日生成的144项签核表同样是模型预填，"
        "用途是方便后续人工逐项确认，不能写成教师或独立人工已经完成签核。",
    )


def _aug03_results(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "6. 8月3日：单智能体基线较弱；短预算没有降低四题总体正确数")
    correct = summary["exploratory"]["correct_runs_by_condition"]
    _table(
        doc,
        ["条件", "正确运行", "正确率", "结论"],
        [
            ["single-direct", f"{correct['single-direct']}/40", _fmt_pct(correct['single-direct']/40), "一次直接作答"],
            ["single-reflect", f"{correct['single-reflect']}/40", _fmt_pct(correct['single-reflect']/40), "15轮自我反思仍提升有限"],
            ["fixed-12", f"{correct['fixed-12']}/40", _fmt_pct(correct['fixed-12']/40), "同12发言预算多智能体"],
            ["fixed-60历史基线", f"{correct['fixed']}/40", _fmt_pct(correct['fixed']/40), "四题60发言固定轮转"],
            ["dynamic-60历史基线", f"{correct['dynamic']}/40", _fmt_pct(correct['dynamic']/40), "四题60发言动态调度"],
        ],
        [2200, 1500, 1500, 4160],
    )
    _paragraph(
        doc,
        "在这四道探索题上，多智能体固定12发言明显高于两个局部单智能体基线；"
        "但该比较同时改变了Agent数量和可交换的信息量，不能单独归因为“多智能体推理能力”。",
    )
    early = summary["exploratory"]["earlystop_by_task"]
    rows = []
    for task in (1, 5, 7, 25):
        rows.append(
            [
                f"ID{task}",
                f"{early['fixed'][str(task)]}/10",
                f"{early['fixed-4'][str(task)]}/10",
                f"{early['fixed-8'][str(task)]}/10",
            ]
        )
    rows.append(["合计", "20/40", "20/40", "20/40"])
    _heading(doc, "6.1 4/8发言是独立短轨迹探索，不是同轨迹早停证明", level=2)
    _table(
        doc,
        ["题", "固定60", "固定4", "固定8"],
        rows,
        [1500, 2400, 2400, 3060],
    )
    _paragraph(
        doc,
        "三种预算的总体正确数均为20/40，但每个预算重新调用模型，不能判断同一条长轨迹在第4或第8条截断后会怎样。"
        "因此8月4日增加了不回灌讨论的影子投票，专门进行同轨迹比较。",
    )


def _confirmatory_matrix(summary: dict[str, Any]) -> list[list[str]]:
    statistics = summary["confirmatory"]["condition_statistics"]
    by_condition = summary["confirmatory"]["by_condition_task"]
    rows: list[list[str]] = []
    order = (
        "single-local",
        "fixed-12",
        "structured-12",
        "fixed-60",
        "dynamic-60",
        "fixed-reveal-all",
        "dynamic-reveal-all",
        "official-global-reveal",
    )
    for condition in order:
        correct = []
        disclosure = []
        for task_id in (1, 2, 3):
            stat = statistics[f"{condition}:ID{task_id}"]
            correct.append(f"{stat['successes']}/{stat['runs']}")
            if condition == "official-global-reveal":
                disclosure.append("100%机械")
            else:
                rate = by_condition[f"{condition}:ID{task_id}"][
                    "atomic_disclosure_rate"
                ]
                disclosure.append(_fmt_pct(rate))
        rows.append(
            [
                CONDITION_LABELS[condition],
                *correct,
                " / ".join(disclosure),
            ]
        )
    return rows


def _aug04_results(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "7. 8月4日：官方三题确认性实验把题目、预算和配对关系锁定")
    _paragraph(
        doc,
        "确认性实验只使用官方short benchmark ID1/ID2/ID3，每个条件每题10次，"
        "同一重复编号使用相同事实分配，共7×3×10=210次。另以独立30次运行严格补齐官方同款全局Reveal-All。",
    )
    _table(
        doc,
        ["条件", "ID1", "ID2", "ID3", "ID1/ID2/ID3披露率"],
        _confirmatory_matrix(summary),
        [2900, 950, 950, 950, 3610],
        font_size=7.4,
    )
    _callout(
        doc,
        "最清楚的反例：",
        "ID2中，固定60为0/10，动态60为2/10，Structured-12为7/10，旧固定逐人机械披露为6/10，"
        "动态逐人机械披露为10/10，官方同款全局Reveal-All为10/10。信息公开的时机和解释方式都会改变结果。",
    )
    statistics = summary["confirmatory"]["condition_statistics"]
    _paragraph(
        doc,
        "区间解释：本地每题每条件只有10次。任何10/10的Wilson 95%区间均为"
        f"{_fmt_ci(statistics['official-global-reveal:ID1'])}，不能把观察到的100%当作总体真值。",
    )
    _heading(doc, "7.1 官方GPT-4.1与本地DeepSeek仅做描述性对照", level=2)
    official = summary["confirmatory"]["official_gpt41"]
    rows = []
    for task_id in (1, 2, 3):
        base = official["baseline"]["by_task"][str(task_id)]
        reveal = official["reveal_all"]["by_task"][str(task_id)]
        local = statistics[f"official-global-reveal:ID{task_id}"]
        rows.append(
            [
                f"ID{task_id}",
                f"{round(base['majority_accuracy']*base['runs'])}/{base['runs']}",
                f"{round(reveal['majority_accuracy']*reveal['runs'])}/{reveal['runs']}",
                f"{local['successes']}/{local['runs']}",
                base["correct_answer"],
            ]
        )
    _table(
        doc,
        ["题", "官方GPT-4.1 Baseline", "官方GPT-4.1 Reveal-All", "本地DeepSeek官方同款", "正确答案"],
        rows,
        [850, 2350, 2350, 2350, 1460],
        font_size=7.7,
    )
    _paragraph(
        doc,
        "官方Baseline分别为11/30、4/30、6/30；官方Reveal-All为10/10、9/10、10/10；"
        "本地DeepSeek官方同款为10/10、10/10、10/10。由于模型、实现层和采样次数不同，不能把差异归因于MAF。",
    )


def _disclosure_section(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "8. 披露率如何由AI判断并由规则程序计算")
    _paragraph(
        doc,
        "审计模型不直接报百分比。它对每个原子事实返回一次判断、拥有者、证据消息ID、原文引文、理由和置信度；"
        "程序验证事实ID完整性、拥有者一致性、消息归属和引文子串后，再求和得到百分比。",
    )
    prompt = (
        "Determine whether every atomic private claim was disclosed in the public discussion.\n\n"
        "Rules:\n"
        "1. Only an owner-authored message can disclose that owner's fact.\n"
        "2. A faithful paraphrase counts; a polarity reversal, contradiction, or decision-relevant omission does not.\n"
        "3. For disclosed=true, cite at least one owner-authored message ID and copy one exact non-empty substring from it.\n"
        "4. For disclosed=false, return an empty message-ID list and quote.\n"
        "5. Return exactly one judgment for every fact ID. Do not calculate the percentage.\n\n"
        "ATOMIC PRIVATE FACTS:\n"
        "- <fact_id> | owner=<agent_id> | claim=<private fact>\n\n"
        "PUBLIC MESSAGES:\n"
        "<message_id> | <agent_id> | <content>\n\n"
        "Return strict JSON only: facts[{fact_id, owner_agent_id, disclosed, evidence_message_ids, evidence_quote, reason, confidence}]"
    )
    _code(doc, "8.1 原子披露审计Prompt（固定规则）", prompt)
    example = summary["disclosure_example"]
    _callout(
        doc,
        "完整计算案例：",
        f"study key=task-1/dynamic-60/rep-2；run ID={example['run_id']}；{example['arithmetic']}。",
    )
    judgments = {item["fact_id"]: item for item in example["judgments"]}
    rows = []
    for fact in example["facts"]:
        judgment = judgments[fact["fact_id"]]
        rows.append(
            [
                fact["owner_agent_id"],
                "是" if judgment["disclosed"] else "否",
                judgment["evidence_message_ids"][0]
                if judgment["evidence_message_ids"]
                else "—",
                judgment["evidence_quote"] or judgment["reason"],
            ]
        )
    _table(
        doc,
        ["拥有者", "披露", "证据消息", "引文或未披露理由"],
        rows,
        [1300, 1000, 2200, 4860],
        font_size=7.8,
    )
    _paragraph(
        doc,
        "机械Reveal条件不调用审计模型：系统注入记录直接证明事实公开，因此披露率按构造记为100%。"
        "这解决的是“信息是否公开”，不代表信息已经被正确理解。",
    )


def _statistics_and_earlystop(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "9. 统计校正与同轨迹早停：点估计改善尚未达到显著水平")
    paired = summary["paired_tests"]
    rows = []
    labels = {
        "fixed-60_vs_dynamic-60": "固定60 vs 动态60",
        "fixed-reveal-all_vs_dynamic-reveal-all": "固定逐人披露 vs 动态逐人披露",
        "fixed-12_vs_structured-12": "固定12 vs Structured-12",
    }
    for key in labels:
        item = paired[key]
        rows.append(
            [
                labels[key],
                f"{item['left_successes']}/{item['pairs']}",
                f"{item['right_successes']}/{item['pairs']}",
                str(item["a_only"]),
                str(item["b_only"]),
                f"{item['p_value']:.3f}",
            ]
        )
    _table(
        doc,
        ["配对比较", "左条件正确", "右条件正确", "仅左正确", "仅右正确", "精确p"],
        rows,
        [2700, 1400, 1400, 1150, 1150, 1560],
        font_size=7.8,
    )
    _paragraph(
        doc,
        "三个双侧精确McNemar p值为0.125、0.125和0.388，均未达到0.05。"
        "因此只能说动态或Structured条件在本样本中的正确次数更高，不能写成统计上显著优越。",
    )
    early = summary["early_stop"]
    _heading(doc, "9.1 同一条60发言轨迹上的候选早停", level=2)
    _table(
        doc,
        ["固定60运行", "同一多数", "最终平票", "候选正确", "最终正确"],
        [[str(early["runs"]), "29", "1", f"{early['candidate_correct']}/30", f"{early['final_correct']}/30"]],
        [1800, 1800, 1800, 1980, 1980],
    )
    exceptional = early["exceptional_cases"][0]
    _callout(
        doc,
        "唯一异常：",
        f"task-{exceptional['study_key']['task_id']} / fixed-60 / repetition-{exceptional['study_key']['repetition']} / "
        f"run {exceptional['run_id']}。第{exceptional['candidate_round']}轮候选为{exceptional['candidate_answer']}，"
        f"最终投票为{', '.join(exceptional['final_votes'])}，即2-2 final tie。候选和最终均错误。",
    )
    _paragraph(
        doc,
        "结论：本样本中早停候选没有改变总正确数，但存在一致状态不能保持到最终的情况。"
        "该结果支持继续扩大同轨迹样本，而不是立即把早停写成安全优化。",
    )


def _mast_cases(doc: DocumentType, summary: dict[str, Any]) -> None:
    _heading(doc, "10. MAST失败模式在实验中的具体位置")
    rows = []
    for case in summary["mast_cases"]:
        key = case["study_key"]
        location = (
            f"ID{key['task_id']} / {key['condition']} / rep{key['repetition']}\nrun={case['run_id']}"
            if key
            else "未确认具体运行"
        )
        rows.append(
            [
                case["mast_mode"],
                case["boundary"],
                location,
                case["interpretation"],
            ]
        )
    _table(
        doc,
        ["模式", "判断边界", "具体位置", "证据解释"],
        rows,
        [1000, 2050, 2600, 3710],
        font_size=7.6,
    )
    _bullets(
        doc,
        [
            "FM-2.4已确认：ID1固定60 repetition-0中披露率仅25%，至少一个拥有者事实未出现在其公开消息中。",
            "FM-2.5是强候选：ID2固定逐人披露 repetition-1中信息已100%公开，但四个最终理由仍把被阻塞的East Town隧道解释为畅通，并一致选错；要断言因果上的“忽略”，仍需人工编码。",
            "FM-2.6未确认：没有直接验证到理由与最终投票相互冲突的运行，因此不虚构发生率。",
        ],
    )


def _limitations_and_next(doc: DocumentType) -> None:
    _heading(doc, "11. 局限、稳健性与下一步")
    _heading(doc, "11.1 会改变结论解释的局限", level=2)
    _bullets(
        doc,
        [
            "确认性范围只有官方short的3题，不能外推到65题或其他任务类型；8月2日至3日的四题探索结果也不能与官方三题混池。",
            "每题每条件只有10次本地重复，Wilson区间较宽；三组预设配对比较均未达到0.05。",
            "官方使用GPT-4.1，本地使用DeepSeek；框架、模型和Prompt效应未完全正交分离。",
            "动态调度器读取全局私有事实清单，属于中心化特权机制；五个权重尚未做敏感性分析。",
            "DeepSeek生成没有服务端seed控制；temperature=0仍不等于严格确定。",
            "披露审计主要由DeepSeek判断DeepSeek文本；确定性证据规则能验证引文，不能消除同模型偏差。",
            "Codex的144项复核和预填签核表不是独立人工金标准。",
            "机械Reveal的100%披露由系统构造，因此只能检验“完整公开”条件，不能代表自然交流质量。",
        ],
    )
    _heading(doc, "11.2 下一轮建议", level=2)
    _bullets(
        doc,
        [
            "请两名不知道AI标签的人类审阅者独立盲审全部AI—规则分歧，并对一致判断分层抽样，报告一致率、精确率、召回率和修订披露率。",
            "从官方公开逐题结果中再选有明确对标数字的任务，扩展到至少10题，并保持同模型、同Prompt、同预算。",
            "把Reveal机制与调度机制做2×2因子设计，避免把“完整信息”和“动态发言”混成一个效果。",
            "对五因子权重和调度器可见信息做敏感性分析，并增加不读取全局私有事实的可部署版本。",
            "在更多条件上加入不回灌的影子投票，以同轨迹方式评估早停，而不是重新运行短预算轨迹。",
            "把FM-2.5候选案例交给人工做因果编码，区分信息未披露、信息被忽略和正确证据被错误解释。",
        ],
    )
    _heading(doc, "11.3 仍需回答的问题", level=2)
    _bullets(
        doc,
        [
            "官方全局Reveal的提升主要来自更早看到信息，还是来自四次重复注入提高了注意力？",
            "如果动态调度器看不到私有事实清单，当前点估计优势是否仍能保持？",
            "在同一模型下，MAF执行层与作者模拟器是否会产生可复现的系统性差异？",
            "披露、正确解释和最终行动之间，哪一步是ID2错误共识的主要瓶颈？",
        ],
    )


def _appendix_tasks(doc: DocumentType, summary: dict[str, Any]) -> None:
    _page_break(doc)
    _heading(doc, "附录A：官方三题的完整信息结构")
    _paragraph(
        doc,
        "三题共享同一个山村撤离背景和West City、East Town、North Hill三个选项，但公共信息和四条私有事实不同。"
        "以下事实直接来自锁定的官方65题快照。",
    )
    for task in summary["task_details"]:
        _heading(
            doc,
            f"A.{task['task_id']} ID{task['task_id']}｜{task['name']}｜正确答案：{task['correct_answer']}",
            level=2,
        )
        rows = []
        for index, fact in enumerate(task["shared_information"], start=1):
            rows.append(["公共信息", str(index), fact])
        for index, fact in enumerate(task["private_facts"], start=1):
            rows.append(["私有事实", str(index), fact])
        _table(doc, ["类型", "序号", "内容"], rows, [1500, 800, 7060], font_size=7.6)
        _paragraph(doc, f"证据链：{task['evidence_chain']}", bold_lead="证据链：")


def _appendix_reproducibility(
    doc: DocumentType,
    summary: dict[str, Any],
    global_message: str,
) -> None:
    _heading(doc, "附录B：复现证据、代码提交和官方同款公开消息")
    publication = summary["publication"]
    _table(
        doc,
        ["项目", "位置"],
        [
            ["项目仓库", publication["repository"]],
            ["实验分支", publication["branch"]],
            ["Pull Request", publication["pull_request"]],
            ["截至8月4日报告提交", publication["report_commit"]],
            ["原始证据Release", publication["release"]],
            ["官方代码", summary["official_code"]["repository"]],
            ["官方结果", summary["official_code"]["results"]],
        ],
        [2500, 6860],
        font_size=7.8,
    )
    selected_names = [
        "hiddenbench-governance-20260802.jsonl",
        "hiddenbench-structured-20260802.jsonl",
        "hiddenbench-contrast-20260803.jsonl",
        "hiddenbench-earlystop-20260803.jsonl",
        "hiddenbench-confirmatory-20260804.jsonl",
        "hiddenbench-official-reveal-supplement-20260804.jsonl",
        "hiddenbench-confirmatory-20260804.audits.jsonl",
    ]
    _table(
        doc,
        ["证据文件", "SHA-256"],
        [[name, summary["source_sha256"][name]] for name in selected_names],
        [3900, 5460],
        font_size=7.1,
    )
    _code(
        doc,
        "B.1 official-compatible global Reveal-All 第一条公开消息",
        global_message,
    )


def _first_jsonl(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                value = json.loads(line)
                if not isinstance(value, dict):
                    raise ValueError(f"first JSONL row is not an object: {path}")
                return value
    raise ValueError(f"empty JSONL file: {path}")


def _validate_tables(doc: DocumentType) -> None:
    if len(doc.tables) < 15:
        raise ValueError("complete report table count is incomplete")
    for table in doc.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        if width is None or width.get(qn("w:w")) != "9360":
            raise ValueError("report table width validation failed")
        if indent is None or indent.get(qn("w:w")) != "120":
            raise ValueError("report table indent validation failed")


def build_report(
    *,
    root: Path,
    output: Path,
    desktop_output: Path | None,
) -> Path:
    root = root.resolve()
    summary = json.loads(
        (root / "reports/data/hiddenbench-aug02-complete-summary.json").read_text(
            encoding="utf-8"
        )
    )
    expected_scope = {
        "local_runs": 560,
        "discussion_vote_api_requests": 25531,
        "audit_api_requests": 397,
        "total_model_requests": 25928,
        "confirmatory_runs": 210,
        "official_global_reveal_runs": 30,
    }
    if summary["scope"] != expected_scope or not summary["all_gates_passed"]:
        raise ValueError("complete report requires the locked 560-run scope")
    preserved = {
        root / "reports" / name: _sha256(root / "reports" / name)
        for name in PRESERVED_REPORTS
    }
    supplement_first = _first_jsonl(
        root / "artifacts/hiddenbench-official-reveal-supplement-20260804.jsonl"
    )
    global_message = supplement_first["run"]["discussion_messages"][0][
        "content"
    ]

    doc = Document()
    _configure_complete(doc)
    _masthead(doc, summary)
    _technical_summary(doc, summary)
    _page_break(doc)
    _teacher_answers(doc)
    _timeline(doc, summary)
    _paper_boundary(doc, summary)
    _method_and_metrics(doc, summary)
    _aug02_results(doc, summary)
    _aug03_results(doc, summary)
    _aug04_results(doc, summary)
    _disclosure_section(doc, summary)
    _statistics_and_earlystop(doc, summary)
    _mast_cases(doc, summary)
    _limitations_and_next(doc)
    _appendix_tasks(doc, summary)
    _appendix_reproducibility(doc, summary, global_message)

    doc.core_properties.title = (
        "HiddenBench × MAF 完整实验报告（2026年8月2日至8月4日）"
    )
    doc.core_properties.subject = "披露治理、对照实验、确认性复现与失败模式定位"
    doc.core_properties.author = "MAS 知识治理项目组"
    doc.core_properties.keywords = (
        "HiddenBench, Microsoft Agent Framework, DeepSeek, MAST, "
        "disclosure, multi-agent"
    )
    _validate_tables(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    for path, digest in preserved.items():
        if _sha256(path) != digest:
            raise ValueError(f"preserved report changed during build: {path.name}")
    check = Document(output)
    _validate_tables(check)
    if desktop_output is not None:
        desktop_output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(output, desktop_output)
        if _sha256(output) != _sha256(desktop_output):
            raise ValueError("repository and Desktop report copies differ")
    return output


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    output = root / "reports" / REPORT_NAME
    desktop = Path.home() / "Desktop" / REPORT_NAME
    print(build_report(root=root, output=output, desktop_output=desktop))


if __name__ == "__main__":
    main()
