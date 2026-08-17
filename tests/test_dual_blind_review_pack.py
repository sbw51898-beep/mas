from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from reports.build_dual_blind_review_pack import build_dual_blind_review_pack
from reports.merge_dual_blind_reviews import merge_dual_blind_reviews


ROOT = Path(__file__).resolve().parents[1]


def _review_ids(path: Path) -> list[str]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["盲审签核"]
    return [
        str(sheet.cell(row=index, column=2).value)
        for index in range(2, sheet.max_row + 1)
        if sheet.cell(row=index, column=2).value
    ]


def _fill_metadata_and_decisions(path: Path, *, reviewer: str) -> None:
    workbook = load_workbook(path)
    guide = workbook["填写说明"]
    guide["B4"] = reviewer
    guide["B5"] = "2026-08-17"
    guide["B6"] = "我确认以上为本人独立判断"
    sheet = workbook["盲审签核"]
    for index in range(2, sheet.max_row + 1):
        if sheet.cell(row=index, column=2).value:
            sheet.cell(row=index, column=6).value = "否"
    workbook.save(path)


def test_dual_blind_pack_has_identical_cases_in_distinct_blind_orders(
    tmp_path: Path,
) -> None:
    pack = build_dual_blind_review_pack(ROOT, tmp_path)
    reviewer_a = pack["reviewer_a"]
    reviewer_b = pack["reviewer_b"]

    a_ids = _review_ids(reviewer_a)
    b_ids = _review_ids(reviewer_b)
    assert len(a_ids) == len(b_ids) == 144
    assert set(a_ids) == set(b_ids)
    assert a_ids != b_ids

    workbook = load_workbook(reviewer_a, data_only=True)
    headers = [
        str(workbook["盲审签核"].cell(row=1, column=index).value)
        for index in range(1, 10)
    ]
    lower_headers = " ".join(headers).casefold()
    assert "ai" not in lower_headers
    assert "rule" not in lower_headers
    assert "模型标签" not in lower_headers


def test_dual_blind_pack_uses_plain_cells_for_visible_reviewer_instructions(
    tmp_path: Path,
) -> None:
    pack = build_dual_blind_review_pack(ROOT, tmp_path)
    workbook = load_workbook(pack["reviewer_a"], data_only=True)
    guide = workbook["填写说明"]

    merged = {str(cell_range) for cell_range in guide.merged_cells.ranges}
    assert not merged
    assert "双人盲审签核表" in str(guide["A1"].value)
    assert "仅当信息拥有者本人" in str(guide["A9"].value)


def test_merge_rejects_non_human_identity_and_missing_decisions(
    tmp_path: Path,
) -> None:
    pack = build_dual_blind_review_pack(ROOT, tmp_path)
    reviewer_a = pack["reviewer_a"]
    reviewer_b = pack["reviewer_b"]

    _fill_metadata_and_decisions(reviewer_a, reviewer="Codex")
    _fill_metadata_and_decisions(reviewer_b, reviewer="李四")
    with pytest.raises(ValueError, match="human reviewer"):
        merge_dual_blind_reviews(ROOT, reviewer_a, reviewer_b, tmp_path / "merged")

    _fill_metadata_and_decisions(reviewer_a, reviewer="张三")
    workbook = load_workbook(reviewer_b)
    workbook["盲审签核"]["F2"] = ""
    workbook.save(reviewer_b)
    with pytest.raises(ValueError, match="missing an explicit disclosure decision"):
        merge_dual_blind_reviews(ROOT, reviewer_a, reviewer_b, tmp_path / "merged")
