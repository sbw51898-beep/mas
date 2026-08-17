import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from reports.build_correctness_governance_report import (
    build_correctness_governance_report,
)


ROOT = Path(__file__).resolve().parents[1]
ORIGINAL_PAPER = ROOT / "reports" / "基于MAF的多智能体私有信息披露机制复现与对比研究_2026-08-05.docx"
FOOTNOTE_PAPER = ROOT / "reports" / "基于MAF的多智能体私有信息披露机制复现与对比研究_脚注版_2026-08-05.docx"
COMPLETE_REPORT = ROOT / "reports" / "给彭老师的完整实验报告_2026-08-02至2026-08-04.docx"
ORIGINAL_PAPER_SHA256 = "e1aa4be327c847ddfa3987e408a1a7a5f178f798eb8903928d95ff2a37c5f668"
FOOTNOTE_PAPER_SHA256 = "4f486bf87a380b0961f69ae1e822609993297ac94c2590f6a4334851b572f575"
COMPLETE_REPORT_SHA256 = "0d56f74c359e1a7359102de678c8c8e39bd568d381e9ff2ccf37189b83f72765"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + cells)


def xml_text(element: etree._Element) -> str:
    return "".join(node.text or "" for node in element.xpath(".//w:t", namespaces=NS))


def test_report_follows_correctness_governance_thesis(tmp_path: Path) -> None:
    output = tmp_path / "correctness.docx"

    build_correctness_governance_report(ROOT, output, None)

    document = Document(output)
    text = document_text(document)
    for heading in [
        "1 引言",
        "2 多智能体错误现象与已有研究",
        "3 多智能体正确性治理体系",
        "4 已有工作、本文改进与原创边界",
        "5 实验设置与对标方法",
        "6 实验结果与归因分析",
        "7 总结与展望",
        "参考文献",
        "附录",
    ]:
        assert heading in text
    assert "公开、读取、解释和行动" in text
    assert "全局 Reveal-All 仅作为" in text
    assert "不能把差异单独归因" in text
    assert "错误现象—观测指标—针对性干预" in text
    assert len(document.inline_shapes) >= 1


def test_report_uses_locked_results_and_bounded_claims(tmp_path: Path) -> None:
    output = tmp_path / "correctness.docx"

    build_correctness_governance_report(ROOT, output, None)

    text = document_text(Document(output))
    for claim in [
        "11/30、4/30、6/30",
        "10/10、0/10、7/10",
        "30/30",
        "26/30",
        "17/30",
        "21/30",
        "20/30",
        "24/30",
        "p=0.125",
        "p=0.388",
        "未达到统计显著",
    ]:
        assert claim in text
    assert "FM-2.4 已确认" in text
    assert "FM-2.5 待人工因果编码" in text
    assert "FM-2.6 未确认" in text
    assert "动态发言显著优于" not in text
    assert "结构化讨论显著优于" not in text
    assert "覆盖全部 65 题" not in text


def test_report_explains_task_selection_and_originality_boundaries(tmp_path: Path) -> None:
    output = tmp_path / "correctness.docx"

    build_correctness_governance_report(ROOT, output, None)

    text = document_text(Document(output))
    assert "官方公开了逐题 GPT-4.1 基线和 Reveal-All 结果" in text
    assert "三题共享同一背景和候选项" in text
    assert "本文不把三题结论外推到全部 65 题" in text
    assert "已有研究：" in text
    assert "本文实现与改进：" in text
    assert "非本文原创：" in text
    assert "external centralized scheduler" in text
    assert "不是人工金标准" in text


def test_report_has_true_sequential_footnotes_and_ordered_bibliography(tmp_path: Path) -> None:
    output = tmp_path / "correctness.docx"

    build_correctness_governance_report(ROOT, output, None)

    with zipfile.ZipFile(output) as archive:
        assert "word/footnotes.xml" in archive.namelist()
        document = etree.fromstring(archive.read("word/document.xml"))
        footnotes = etree.fromstring(archive.read("word/footnotes.xml"))
        body_ids = [
            int(node.get(f"{{{W_NS}}}id"))
            for node in document.xpath(".//w:footnoteReference", namespaces=NS)
        ]
        note_ids = [
            int(node.get(f"{{{W_NS}}}id"))
            for node in footnotes.xpath(
                ".//w:footnote[number(@w:id) >= 1]", namespaces=NS
            )
        ]
        reserved_types = {
            int(node.get(f"{{{W_NS}}}id")): node.get(f"{{{W_NS}}}type")
            for node in footnotes.xpath(
                ".//w:footnote[number(@w:id) < 1]", namespaces=NS
            )
        }
        assert body_ids == list(range(1, len(body_ids) + 1))
        assert note_ids == body_ids
        assert len(body_ids) >= 15
        assert reserved_types == {-1: "separator", 0: "continuationSeparator"}
        for node in footnotes.xpath(
            ".//w:footnote[number(@w:id) >= 1]", namespaces=NS
        ):
            assert xml_text(node).strip()

        paragraphs = document.xpath(".//w:body/w:p", namespaces=NS)
        texts = [xml_text(paragraph) for paragraph in paragraphs]
        reference_index = texts.index("参考文献")
        appendix_index = texts.index("附录")
        body_text = "\n".join(texts[:reference_index])
        bibliography = [
            paragraph
            for paragraph in paragraphs[reference_index + 1 : appendix_index]
            if xml_text(paragraph).strip()
        ]
        assert len(bibliography) == len(body_ids)
        for index, paragraph in enumerate(bibliography, start=1):
            assert xml_text(paragraph).startswith(f"[{index}] ")
            assert paragraph.xpath("./w:pPr/w:keepLines", namespaces=NS)
        assert "[[FN" not in body_text
        assert not re.search(r"\[\d+(?:-\d+)?\]", body_text)


def test_report_has_visual_page_break_guards(tmp_path: Path) -> None:
    output = tmp_path / "correctness.docx"

    build_correctness_governance_report(ROOT, output, None)

    document = Document(output)
    title = document.paragraphs[0]
    title_sizes = [run.font.size.pt for run in title.runs if run.font.size]
    assert title_sizes and max(title_sizes) <= 20

    appendix_d = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "附录 D  术语与结论边界"
    )
    assert appendix_d.paragraph_format.page_break_before


def test_report_uses_fixed_geometry_and_preserves_existing_reports(tmp_path: Path) -> None:
    output = tmp_path / "correctness.docx"
    desktop = tmp_path / "desktop-correctness.docx"

    build_correctness_governance_report(ROOT, output, desktop)

    document = Document(output)
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(
        round(value.inches, 2) == 1.0
        for value in [
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        ]
    )
    assert len(document.tables) >= 7
    for table in document.tables:
        properties = table._tbl.tblPr
        width = properties.find(qn("w:tblW"))
        indent = properties.find(qn("w:tblInd"))
        assert width is not None and width.get(qn("w:w")) == "9360"
        assert width.get(qn("w:type")) == "dxa"
        assert indent is not None and indent.get(qn("w:w")) == "120"
        assert indent.get(qn("w:type")) == "dxa"

    assert output.read_bytes() == desktop.read_bytes()
    assert hashlib.sha256(ORIGINAL_PAPER.read_bytes()).hexdigest() == ORIGINAL_PAPER_SHA256
    assert hashlib.sha256(FOOTNOTE_PAPER.read_bytes()).hexdigest() == FOOTNOTE_PAPER_SHA256
    assert hashlib.sha256(COMPLETE_REPORT.read_bytes()).hexdigest() == COMPLETE_REPORT_SHA256
