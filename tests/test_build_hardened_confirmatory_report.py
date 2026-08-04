from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from reports.build_hardened_confirmatory_report import build_report


ROOT = Path(__file__).parents[1]
OLD_REPORT = ROOT / "reports" / "HiddenBench确认性修订实验报告_2026-08-04.docx"
OLD_SHA256 = "bb4028b43ae9f7d8b9957e484b03af17ce6d2dc07ee60cc16b35e18234ead211"


def _text(document: Document) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(values)


def test_hardened_report_contains_required_methodological_corrections(
    tmp_path: Path,
) -> None:
    output = tmp_path / "hardened.docx"

    build_report(root=ROOT, output=output, desktop_output=None)

    document = Document(output)
    text = _text(document)
    for required in (
        "official-compatible global Reveal-All",
        "owner-by-owner mechanical reveal",
        "generation seed was not transmitted",
        "external centralized scheduler",
        "Wilson 95%",
        "McNemar",
        "2-2 final tie",
        "3 disclosed / 4 total = 75.0%",
        "FM-2.4",
        "FM-2.5",
        "FM-2.6",
        "evacuation_west_city",
        "evacuation_north_hill",
        "evacuation_east_town",
    ):
        assert required in text
    assert "answer reversal" not in text
    assert "one reversal" not in text
    assert "30 supplementary runs" in text
    assert "2,160" in text
    package_sha, package_name = (
        ROOT
        / "release"
        / "hiddenbench-official-reveal-supplement-20260804-raw.sha256"
    ).read_text(encoding="utf-8").strip().split(maxsplit=1)
    assert package_name in text
    assert package_sha in text


def test_hardened_report_uses_standard_business_geometry_and_preserves_old_report(
    tmp_path: Path,
) -> None:
    output = tmp_path / "hardened.docx"

    build_report(root=ROOT, output=output, desktop_output=None)

    document = Document(output)
    section = document.sections[0]
    assert section.top_margin.inches == 1.0
    assert section.bottom_margin.inches == 1.0
    assert section.left_margin.inches == 1.0
    assert section.right_margin.inches == 1.0
    assert len(document.tables) >= 10
    for table in document.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        assert width is not None
        assert width.get(qn("w:w")) == "9360"
        assert width.get(qn("w:type")) == "dxa"
        assert indent is not None
        assert indent.get(qn("w:w")) == "120"
        assert indent.get(qn("w:type")) == "dxa"
    assert hashlib.sha256(OLD_REPORT.read_bytes()).hexdigest() == OLD_SHA256
