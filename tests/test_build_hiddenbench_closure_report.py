from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document

from reports.build_hiddenbench_closure_report import (
    build_hiddenbench_closure_report,
)


ROOT = Path(__file__).resolve().parents[1]
PRESERVED_REPORT = (
    ROOT / "reports" / "基于MAF的多智能体正确性治理与错误链路分析_2026-08-17.docx"
)
PRESERVED_SHA256 = "21ac7bcbde08b17eb5f1978af35192e03af72642ae9d34fdf4c8c35e790a048f"


def _document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + cells)


def test_closure_report_keeps_only_completed_experiment_content(
    tmp_path: Path,
) -> None:
    output = tmp_path / "closure.docx"

    build_hiddenbench_closure_report(ROOT, output, None)

    text = _document_text(Document(output))
    for required in [
        "基于 Microsoft Agent Framework 的 HiddenBench 多智能体正确性诊断与错误链路分析",
        "6.6 闭环补充：单智能体对照与错误链路",
        "4/30",
        "20/30",
        "17/30",
        "普通条件下不随机、不强制披露",
        "普通 60 次发言条件下的私有事实包披露率",
        "90/120（75.0%）",
        "87/120（72.5%）",
        "模型辅助的敏感性指标而非人工真值",
        "补充的独立审计模型自报披露审计覆盖 60 个普通运行",
        "冻结主审计的 disclosure_rate 位于 artifacts/hiddenbench-confirmatory-20260804.audits.jsonl",
        "B.3 Prompt 要求模型只返回逐条 disclosed=true/false、证据消息 ID 和原文片段",
        "普通 fixed-60 与 dynamic-60 的主口径 D_owner 分别为 90/120（75.0%）和 87/120（72.5%）",
        "模型自报计数、分母和百分比均与其逐条标签一致",
        "ID2 是最直接的警示：全信息单 Agent 为 10/10",
        "0/4=0.0% 是有效边界结果，不是漏检、缺失值或分母为零",
        "它不是“任意 Agent 在公共消息中提到该事实”的群体知晓率",
        "第 1 轮披露为 85/120=70.8%",
        "第 1 轮披露为 72/120=60.0%",
        "按最终多数投票是否正确分层",
        "治理原型",
        "不是等计算预算实验",
        "21/30 对 17/30 的方向不能被解释为纯粹的调度器因果效果",
        "S_i = 0.30 d_i + 0.30 u_i + 0.15 r_i + 0.15 q_i + 0.10 w_i",
        "FM-2.5 还需要把",
        "FM-2.6 还需要证明",
        "披露率的证据边界也需要单独说明",
        "15 轮 × 4 个 Agent",
        "0d078a3c85f64262",
        "e46521e6691a10ca",
        "bfba1db7a056ebea",
    ]:
        assert required in text
    for forbidden in [
        "人工复核入口",
        "盲审",
        "真人双盲审阅",
        "双人真人盲审",
        "模型与框架解耦",
        "解耦实验",
        "GPT-4.1（或同等可验证接口）尚不可用",
        "尚未完成",
        "仍待真实外部输入",
        "下一步最重要的是做解耦实验",
        "待人工因果编码",
        "需要人工复核",
    ]:
        assert forbidden not in text


def test_closure_report_closes_teacher_review_ambiguities(
    tmp_path: Path,
) -> None:
    output = tmp_path / "closure.docx"

    build_hiddenbench_closure_report(ROOT, output, None)

    text = _document_text(Document(output))
    for required in [
        "独立审计模型自报披露率",
        "参与讨论的 Agent 没有在原始公开讨论中输出披露百分比",
        "60/60 只表示同一审计调用内的算术一致性",
        "官方基线每题 n=30，本地条件每题 n=10",
        "百分比仅用于方向性比较，不能作同样本量下的数值复现",
        "MAF Agent/Agent.run 执行层",
        "讨论协议、选择器和审计由项目代码实现",
        "不等计算预算且选择器可见完整私有事实",
        "三题共享同一撤离背景",
        "McNemar p 值仅作探索性描述",
        "没有进行多重比较校正",
        "atomic-fact 只是内部事实包 ID",
        "D_owner 也不等于 D_group",
        "hiddenbench-teacher-review-20260820-v1",
        "旧版 Release 仅保存历史冻结轨迹",
        "OPENAI_BASE_URL=https://api.deepseek.com",
    ]:
        assert required in text


def test_closure_report_marks_frozen_b3_prompt_and_manifest_commit(
    tmp_path: Path,
) -> None:
    output = tmp_path / "closure.docx"

    build_hiddenbench_closure_report(ROOT, output, None)

    text = _document_text(Document(output))
    assert "B.3 来源事实包披露审计 Prompt（冻结原始模板；原文保留）" in text
    assert "ATOMIC PRIVATE FACTS 是冻结 Prompt 的历史字段名" in text
    assert "不等于 clause-level 原子命题" in text
    assert "B.3  原子事实披露审计 Prompt（真实模板）" not in text

    manifest = json.loads(
        (ROOT / "reports/data/hiddenbench-teacher-review-20260820-v2.manifest.json")
        .read_text(encoding="utf-8")
    )
    assert re.fullmatch(r"[0-9a-f]{40}", manifest["commit"])


def test_closure_report_keeps_true_footnotes_and_preserves_original(
    tmp_path: Path,
) -> None:
    output = tmp_path / "closure.docx"

    build_hiddenbench_closure_report(ROOT, output, None)

    with zipfile.ZipFile(output) as archive:
        assert "word/footnotes.xml" in archive.namelist()
    assert hashlib.sha256(PRESERVED_REPORT.read_bytes()).hexdigest() == PRESERVED_SHA256


def test_closure_report_direct_script_entrypoint_loads() -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "reports" / "build_hiddenbench_closure_report.py"),
            "--help",
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert result.returncode == 0, result.stderr


def test_closure_report_has_no_blank_page_break_before_appendix(tmp_path: Path) -> None:
    output = tmp_path / "closure.docx"

    build_hiddenbench_closure_report(ROOT, output, None)

    document = Document(output)
    appendix_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "附录"
    )
    previous = document.paragraphs[appendix_index - 1]
    assert previous.text.strip()
    assert 'w:type="page"' not in previous._p.xml


def test_closure_report_states_actual_prompts_and_comparison_boundaries(
    tmp_path: Path,
) -> None:
    output = tmp_path / "closure.docx"

    build_hiddenbench_closure_report(ROOT, output, None)

    text = _document_text(Document(output))
    for required in [
        "跨实现的方向性验证，不是严格逐字、逐协议复现",
        "同公开发言预算",
        "fixed-60 的 132 个逻辑模型响应槽",
        "dynamic-60 的 72 个逻辑模型响应槽",
        "同模型单智能体参考组（非等计算预算）",
        "不是理论上限",
        "ID2 是最直接的警示：全信息单 Agent 为 10/10",
        "ID3 中全信息单 Agent 为 0/10，而 fixed-60 为 7/10",
        "隐藏不是主张现实决策应故意瞒信息",
        "中心化、内容感知的诊断调度器",
        "不读取正确答案",
        "不把原始私有事实直接注入 Agent prompt",
        "FM-2.4 候选（输出级证据）",
        "dynamic-60 的披露率略低但正确数更高",
        "pair seed 只随机化事实归属和事实顺序",
        "发言顺序由 fixed/dynamic 协议决定",
        "逐题重复一致性（探索性）",
        "fixed-60：ID1 10/10、ID2 0/10、ID3 7/10",
        "dynamic-60：ID1 10/10、ID2 2/10、ID3 9/10",
        "独立调用不等于独立模型",
        "配对精确 McNemar 检验",
        "n=30 个匹配运行",
        "相同 pair_seed",
        "A 独有正确",
        "双侧精确 McNemar 检验",
        "来源事实包",
        "来源包级覆盖率",
        "0/4=0.0% 是有效边界结果，不是漏检、缺失值或分母为零",
        "冻结主审计的 disclosure_rate 位于 artifacts/hiddenbench-confirmatory-20260804.audits.jsonl",
        "Do not calculate the percentage",
        "百分比由程序确定性汇总",
        "第 1 轮/全部事实",
        "普通 Prompt 没有要求 Agent 必须公开私有事实",
        "oracle-like 调度器",
        "冻结原始证据的 SHA-256",
        "ID2 / fixed-60 / rep-0 的真实填充示例",
        "atomic-fact:fdb95995b31ce89325cc",
        '"disclosed": false',
        "真实 Agent 讨论 Prompt 与输出",
        "实际 SYSTEM prompt",
        "实际 USER prompt",
        "模型实际公开输出",
        "You have received the following information, notice the order of this information is",
        "You are the first to speak.",
        "Previous messages from other people:",
        "It’s your turn to speak.",
        '"vote": "<one exact answer from: <option 1> | <option 2> | <option 3>"',
        "Return strict JSON only in this shape:",
        "[SYSTEM MECHANICAL REVEAL-ALL]",
        "官方 Reveal-All 时序兼容的本地条件（非逐字复现）",
        "0d078a3c85f64262",
    ]:
        assert required in text

    for forbidden in [
        "You are one of four community leaders.",
        "State only conclusions supported by the information available to you.",
        "Do not reveal that you are an AI model.",
        "Based on the complete public discussion, choose exactly one final option:",
        "Return the option and a short evidence-based rationale.",
        "同预算动态发言",
        "信息可得性上限",
        "信息可得性的上限",
        "383ab32a67ba98e7",
        "FM-2.5 待人工因果编码",
        "FM-2.6 未确认",
    ]:
        assert forbidden not in text
