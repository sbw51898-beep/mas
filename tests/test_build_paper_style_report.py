import hashlib
import json
from pathlib import Path
import subprocess

from docx import Document
from docx.oxml.ns import qn

from reports.build_paper_style_report import build_report


ROOT = Path(__file__).resolve().parents[1]
REFERENCE_PATH = ROOT / "reports" / "data" / "hiddenbench-paper-references-20260805.json"
PRESERVED_REPORT = ROOT / "reports" / "给彭老师的完整实验报告_2026-08-02至2026-08-04.docx"
PRESERVED_SHA256 = "0d56f74c359e1a7359102de678c8c8e39bd568d381e9ff2ccf37189b83f72765"


def _document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def test_paper_report_has_required_structure_and_locked_claims(tmp_path: Path) -> None:
    output = tmp_path / "paper.docx"

    build_report(ROOT, output, None)

    document = Document(output)
    text = _document_text(document)
    assert "基于 MAF 的多智能体私有信息披露机制复现与对比研究" in text
    for heading in [
        "摘要",
        "1 引言",
        "2 相关研究",
        "3 基于 MAF 的信息披露治理方法",
        "4 实验结果及分析",
        "5 总结与展望",
        "参考文献",
        "附录",
    ]:
        assert heading in text
    assert "560 次本地运行" in text
    assert "25,928 次模型请求" in text
    assert "ID1、ID2、ID3 均为 10/10" in text
    assert "0.125、0.125 和 0.388" in text
    assert "未达到统计显著优势" in text
    assert "external centralized scheduler" in text
    assert "generation seed was not transmitted" in text
    assert "3 disclosed / 4 total = 75.0%" in text
    assert "2-2 final tie" in text


def test_paper_report_abstract_has_seven_sentences_and_no_citations(tmp_path: Path) -> None:
    output = tmp_path / "paper.docx"

    build_report(ROOT, output, None)

    document = Document(output)
    paragraphs = document.paragraphs
    abstract_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "摘要")
    abstract = paragraphs[abstract_index + 1].text.strip()
    assert abstract.count("。") == 7
    assert "[" not in abstract and "]" not in abstract
    assert len(abstract) <= 900


def test_every_reference_is_cited_and_claims_are_bounded(tmp_path: Path) -> None:
    output = tmp_path / "paper.docx"

    build_report(ROOT, output, None)

    document = Document(output)
    text = _document_text(document)
    references = json.loads(REFERENCE_PATH.read_text(encoding="utf-8"))
    assert len(references) >= 15
    for index in range(1, len(references) + 1):
        assert text.count(f"[{index}]") >= 2
    for forbidden in [
        "答案翻转",
        "人工金标准已完成",
        "证明了普遍有效",
        "显著提升动态发言",
        "浅谈",
        "浅见",
        "粗糙了解",
    ]:
        assert forbidden not in text


def test_paper_report_uses_academic_geometry_figures_and_preserves_source(tmp_path: Path) -> None:
    output = tmp_path / "paper.docx"

    build_report(ROOT, output, None)

    document = Document(output)
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(
        round(value.inches, 2) == 1.0
        for value in [section.top_margin, section.right_margin, section.bottom_margin, section.left_margin]
    )
    assert len(document.inline_shapes) >= 2
    assert len(document.tables) >= 7
    for table in document.tables:
        properties = table._tbl.tblPr
        width = properties.find(qn("w:tblW"))
        indent = properties.find(qn("w:tblInd"))
        assert width is not None and width.get(qn("w:w")) == "9360"
        assert width.get(qn("w:type")) == "dxa"
        assert indent is not None and indent.get(qn("w:w")) == "120"
        assert indent.get(qn("w:type")) == "dxa"
    assert hashlib.sha256(PRESERVED_REPORT.read_bytes()).hexdigest() == PRESERVED_SHA256


def test_document_runtime_can_load_paper_builder() -> None:
    runtime = Path(
        "C:/Users/liuli/.cache/codex-runtimes/codex-primary-runtime/"
        "dependencies/python/python.exe"
    )
    if not runtime.exists():
        return

    result = subprocess.run(
        [
            str(runtime),
            "-c",
            (
                "import runpy; "
                f"runpy.run_path({str(ROOT / 'reports/build_paper_style_report.py')!r})"
            ),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    assert result.returncode == 0, result.stderr


def test_reference_entries_use_correct_document_types(tmp_path: Path) -> None:
    output = tmp_path / "paper.docx"

    build_report(ROOT, output, None)

    text = _document_text(Document(output))
    assert "[2] Garold Stasser and William Titus." in text
    assert "Sampling during Discussion [J]." in text
    assert "[3] Li Lu, Y. Connie Yuan, and Poppy Lauretta McLeod." in text
    assert "A Meta-Analysis [J]." in text
    assert "[11] Mert Cemri" in text
    assert "Why Do Multi-Agent LLM Systems Fail? [C]." in text
