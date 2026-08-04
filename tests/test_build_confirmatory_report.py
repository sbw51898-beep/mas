from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from reports.build_confirmatory_report import build_report


ROOT = Path(__file__).parents[1]


def test_report_contains_required_answers_prompts_and_exact_totals(
    tmp_path: Path,
) -> None:
    output = tmp_path / "report.docx"
    copy = tmp_path / "copy.docx"

    build_report(root=ROOT, output=output, desktop_output=copy)
    document = Document(output)
    text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )

    for required in (
        "ID1、ID2、ID3",
        "7 条件 × 3 题 × 10 次 = 210 次",
        "原子事实披露率",
        "披露百分比",
        "Microsoft Agent Framework",
        "作者自定义 Python 模拟器",
        "局部信息单智能体",
        "完整信息单智能体",
        "29/30",
        "1 次翻转",
        "11,971",
        "152",
        "b_i",
        "SYSTEM MECHANICAL REVEAL-ALL",
        "Determine whether every atomic private claim",
        "FM-2.4",
        "FM-2.5",
        "FM-2.6",
    ):
        assert required in text
    assert "early stopping is safe" not in text
    assert "ID1、ID5、ID7 是论文三道人工题" not in text
    assert len(document.tables) >= 8
    for table in document.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        assert width is not None
        assert width.get(qn("w:type")) == "dxa"
        assert int(width.get(qn("w:w"))) == 9360
    assert hashlib.sha256(output.read_bytes()).digest() == hashlib.sha256(
        copy.read_bytes()
    ).digest()


def test_report_preserves_evidence_urls_and_commits(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    build_report(root=ROOT, output=output, desktop_output=None)
    text = "\n".join(p.text for p in Document(output).paragraphs)

    assert "https://github.com/Yassellee/HiddenBench_ICML" in text
    assert "https://huggingface.co/datasets/YuxuanLi1225/HiddenBench-results" in text
    assert "80f621b223c2a7f0f88434a083c5cf5268d52b51" in text
    assert "1aa577992d3abb3c83e5aeb7ce5475de7100e7af" in text
    assert "e1b32d3c217960bc331c0c751217011e92dd9d2b" in text
