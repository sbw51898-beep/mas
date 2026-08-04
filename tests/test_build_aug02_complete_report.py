from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from reports.build_aug02_complete_report import build_report


ROOT = Path(__file__).parents[1]
PRESERVED_REPORTS = (
    ROOT / "reports/对标与披露治理改进实验报告_2026-08-02.docx",
    ROOT / "reports/整合治理（Structured协议）实验报告_2026-08-02.docx",
    ROOT / "reports/对照实验报告（单智能体与同预算基线）_2026-08-03.docx",
    ROOT / "reports/实验总报告（自老师提要求以来）_2026-08-03.docx",
    ROOT / "reports/HiddenBench确认性完善实验报告_2026-08-04.docx",
)
PRESERVED_HASHES = {
    path: hashlib.sha256(path.read_bytes()).hexdigest()
    for path in PRESERVED_REPORTS
}


def _text(document: Document) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(values)


def test_complete_report_contains_scope_results_prompts_and_boundaries(
    tmp_path: Path,
) -> None:
    output = tmp_path / "complete.docx"

    build_report(root=ROOT, output=output, desktop_output=None)

    document = Document(output)
    text = _text(document)
    for phrase in (
        "560 次本地运行",
        "25,928 次模型请求",
        "3 disclosed / 4 total = 75.0%",
        "2-2 final tie",
        "generation seed was not transmitted",
        "FM-2.4",
        "FM-2.5",
        "FM-2.6",
        "0.125",
        "0.388",
        "Codex 模型辅助复核，不是人工金标准",
        "official-compatible global Reveal-All",
        "owner-by-owner mechanical reveal",
        "external centralized scheduler",
        "Determine whether every atomic private claim was disclosed",
        "evacuation_west_city",
        "evacuation_north_hill",
        "evacuation_east_town",
        "https://github.com/Yassellee/HiddenBench_ICML",
    ):
        assert phrase in text
    assert "答案翻转" not in text
    assert "one reversal" not in text
    assert "人工金标准已完成" not in text
    assert document.core_properties.title == (
        "HiddenBench × MAF 完整实验报告（2026年8月2日至8月4日）"
    )


def test_complete_report_uses_fixed_geometry_and_preserves_existing_reports(
    tmp_path: Path,
) -> None:
    output = tmp_path / "complete.docx"

    build_report(root=ROOT, output=output, desktop_output=None)

    document = Document(output)
    section = document.sections[0]
    assert section.page_width.inches == 8.5
    assert section.page_height.inches == 11.0
    assert section.top_margin.inches == 1.0
    assert section.bottom_margin.inches == 1.0
    assert section.left_margin.inches == 1.0
    assert section.right_margin.inches == 1.0
    assert len(document.tables) >= 15
    for table in document.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        assert width is not None
        assert width.get(qn("w:w")) == "9360"
        assert width.get(qn("w:type")) == "dxa"
        assert indent is not None
        assert indent.get(qn("w:w")) == "120"
        assert indent.get(qn("w:type")) == "dxa"
    for path, expected_hash in PRESERVED_HASHES.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == expected_hash


def test_complete_report_uses_only_the_appendix_page_break(tmp_path: Path) -> None:
    output = tmp_path / "complete.docx"

    build_report(root=ROOT, output=output, desktop_output=None)

    document = Document(output)
    page_breaks = document._element.xpath("//w:br[@w:type='page']")
    assert len(page_breaks) == 1
